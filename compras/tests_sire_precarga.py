"""Precarga de compras desde la propuesta SIRE (RCE) con columnas propias del Excel: Orden de compra, Clasificación (por
número) y Cuenta contable. Son opcionales: lo que falta se ve en rojo (pantalla y Excel); un valor que no se reconoce
no deja crear esa compra."""
import io
from datetime import date
from decimal import Decimal as D
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from contabilidad.models import CentroCosto, CuentaContable
from core.models import PropuestaSIRE, Tercero

from . import sire
from .models import Compra, OrdenCompra

CAB = ['RUC', 'Apellidos y Nombres o Razón social', 'Periodo', 'CAR SUNAT', 'Fecha de emisión', 'Fecha Vcto/Pago',
       'Tipo CP/Doc.', 'Serie del CDP', 'Año', 'Nro CP o Doc. Nro Inicial (Rango)', 'Nro Final (Rango)',
       'Tipo Doc Identidad', 'Nro Doc Identidad', 'Apellidos Nombres/ Razón Social', 'BI Gravado DG', 'IGV / IPM DG',
       'Total CP', 'Moneda']


def xlsx(filas):
    from openpyxl import Workbook
    libro = Workbook()
    for f in filas:
        libro.active.append(f)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


class PrecargaConColumnasPropiasTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        with patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450'))):
            call_command('seed_demo', verbosity=0)
        cls.admin = User.objects.create_superuser('admin_sire', 'a@a.pe', 'Clave-segura-9')
        cls.prov = Tercero.objects.create(numero_doc='20600000022', nombre='ANDINA SAC', tipo='PROVEEDOR')
        otro = Tercero.objects.create(numero_doc='20600000033', nombre='OTRO PROVEEDOR SAC', tipo='PROVEEDOR')
        cls.cc = CentroCosto.objects.create(codigo='MANT', nombre='Mantenimiento', tipo='PRODUCCION')
        cls.oc = OrdenCompra.objects.create(numero='OC01-00000501', tercero=cls.prov, estado='APROBADO',
                                            centro_costo=cls.cc)
        OrdenCompra.objects.create(numero='OC01-00000502', tercero=otro, estado='APROBADO')
        cls.cuenta = CuentaContable.objects.filter(codigo__startswith='63', imputable=True).order_by('codigo').first()
        cls.agrupadora = CuentaContable.objects.filter(imputable=False, codigo__startswith='6').order_by('codigo').first()

    def setUp(self):
        p = patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450')))
        p.start()
        self.addCleanup(p.stop)
        self.client.force_login(self.admin)

    def fila(self, numero, oc='', clasificacion='', cuenta=''):
        return ['20100000001', 'MI EMPRESA', '202609', 'X', '10/09/2026', '', '01', 'F001', '', numero, '', '6',
                self.prov.numero_doc, 'ANDINA SAC', 1000, 180, 1180, 'PEN', oc, clasificacion, cuenta]

    def subir(self):
        archivo = xlsx([CAB + ['Orden de compra', 'Clasificación', 'Cuenta contable'],
                        self.fila('101', oc='501', clasificacion=7, cuenta=int(self.cuenta.codigo)),
                        self.fila('102'),
                        self.fila('103', oc='OC01-00000502', clasificacion='5'),
                        self.fila('104', clasificacion='1 - Mercadería', cuenta=self.agrupadora.codigo)])
        self.client.post(reverse('sire_propuestas'), {
            'libro': 'RCE', 'periodo': '2026-09', 'archivo': SimpleUploadedFile('propuesta.xlsx', archivo)})
        return PropuestaSIRE.objects.get()

    def test_lee_marca_en_rojo_y_precarga(self):
        p = self.subir()
        registros = {r['numero']: r for r in sire.de_json(p.registros)}
        self.assertEqual((registros['101']['oc'], registros['101']['clasificacion'], registros['101']['cuenta']),
                         ('501', '7', self.cuenta.codigo))
        # pantalla: lo que falta en rojo y lo que no se reconoce con su motivo
        r = self.client.get(reverse('sire_propuesta', args=[p.pk]))
        self.assertContains(r, '3 sin orden de compra, clasificación o cuenta (en rojo)')
        self.assertContains(r, 'Falta', count=5)  # 102 sin las tres, 103 sin cuenta y 104 sin orden de compra
        self.assertContains(r, 'OC01-00000501')
        self.assertContains(r, 'es de otro proveedor (OTRO PROVEEDOR SAC)')
        self.assertContains(r, 'no acepta movimientos')
        # Excel del cruce: las tres columnas al final y las vacías de lo que falta precargar en rojo
        resp = self.client.get(reverse('sire_propuesta', args=[p.pk]), {'formato': 'excel'})
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(resp.content)).active
        titulos = [c.value for c in hoja[3]]
        self.assertEqual(titulos[-3:][0], 'Orden de compra')
        self.assertTrue(titulos[-2].startswith('Clasificación (1 Mercadería, 2 Materia prima'))
        self.assertEqual(titulos[-1], 'Cuenta contable')
        self.assertEqual(titulos.count('Orden de compra'), 1)  # la del archivo no se repite entre las de SUNAT
        rojas = [(c.row, c.column) for fila in hoja.iter_rows(min_row=4) for c in fila
                 if c.fill and c.fill.fgColor and c.fill.fgColor.rgb == '00FFC7CE']
        por_factura = {}
        for f, _ in rojas:
            numero = hoja.cell(row=f, column=titulos.index('Nro CP o Doc. Nro Inicial (Rango)') + 1).value
            por_factura[numero] = por_factura.get(numero, 0) + 1
        self.assertEqual(por_factura, {'102': 3, '103': 1, '104': 1})
        # precarga: crea 101 (con orden, centro, clasificación y cuenta) y 102 (sin datos propios); 103 y 104 no
        claves = [f'01|F001|{n}|{self.prov.numero_doc}' for n in ('101', '102', '103', '104')]
        r = self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'crear', 'falta': claves}, follow=True)
        self.assertContains(r, '2 comprobante(s) creados')
        self.assertContains(r, 'F001-103: no se creó')
        c101 = Compra.objects.get(serie='F001', numero='101')
        self.assertEqual((c101.orden_compra, c101.centro_costo, c101.clasificacion, c101.cuenta_contable),
                         (self.oc, self.cc, 'SERVICIO', self.cuenta))
        c102 = Compra.objects.get(serie='F001', numero='102')
        self.assertEqual((c102.orden_compra, c102.clasificacion, c102.cuenta_contable), (None, 'MERCADERIA', None))
        self.assertFalse(Compra.objects.filter(numero__in=['103', '104']).exists())

    def test_clasificacion_por_numero_o_nombre(self):
        self.assertEqual([sire._clasificacion(t) for t in ('1', '1 - MERCADERIA', 'Materia prima', 'SERVICIO', '7',
                                                          '8', '9', 'otra cosa')],
                         ['MERCADERIA', 'MERCADERIA', 'MATERIA_PRIMA', 'SERVICIO', 'SERVICIO', 'HONORARIOS', None, None])
        self.assertEqual(sire._orden_compra('00000501', self.prov.numero_doc)[0], self.oc)
        self.assertIn('No existe', sire._orden_compra('999', self.prov.numero_doc)[1])
        # sin las columnas propias el archivo se lee igual (todas opcionales)
        archivo = SimpleUploadedFile('p.xlsx', xlsx([CAB, self.fila('105')[:-3]]))
        encabezados, registros = sire.leer_completo(archivo)
        self.assertEqual((registros[0]['oc'], registros[0]['clasificacion'], registros[0]['cuenta']), ('', '', ''))
        self.assertEqual(date(2026, 9, 10), registros[0]['fecha'])
