"""Cruce del Registro de Compras con la propuesta SIRE (RCE) que SUNAT arma con los comprobantes electrónicos.

El archivo de la propuesta (TXT separado por «|», o el .zip que lo contiene; también CSV o Excel) se lee por los
nombres de sus columnas: Tipo CP/Doc., Serie del CDP, Nro CP o Doc. Nro Inicial, Nro Doc Identidad (del proveedor),
Apellidos Nombres/Razón Social, Fecha de emisión, BI Gravado DG, IGV / IPM DG y Total CP.

Columnas propias (no son de SUNAT, opcionales): antes de precargar las compras se pueden agregar en el Excel
«Orden de compra», «Clasificación» (por número: 1 Mercadería, 2 Materia prima…) y «Cuenta contable». Las compras que se
crean desde la propuesta las toman; las que faltan se marcan en rojo.
"""
import csv
import io
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

D0 = Decimal('0')
TOLERANCIA = Decimal('1.00')


class ErrorSIRE(Exception):
    pass


def _norm(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', texto.strip().lower())


def _dec(valor):
    if valor in (None, ''):
        return D0
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    try:
        return Decimal(str(valor).replace(',', '').strip() or '0')
    except InvalidOperation:
        return D0


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(valor).strip()[:10], fmt).date()
        except ValueError:
            pass
    return None


def _filas(archivo):
    nombre = (getattr(archivo, 'name', '') or '').lower()
    datos = archivo.read()
    if datos[:4] == b'\xd0\xcf\x11\xe0':  # Excel 97-2003 (.xls)
        import xlrd
        try:
            hoja = xlrd.open_workbook(file_contents=datos).sheet_by_index(0)
        except Exception as exc:
            raise ErrorSIRE(f'No se pudo leer el Excel (.xls): {exc}') from exc
        return [[None if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else
                 xlrd.xldate.xldate_as_datetime(c.value, hoja.book.datemode) if c.ctype == xlrd.XL_CELL_DATE else c.value
                 for c in hoja.row(i)] for i in range(hoja.nrows)]
    if nombre.endswith('.zip') or datos[:2] == b'PK' and not nombre.endswith('.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(datos)) as z:
                txt = [n for n in z.namelist() if n.lower().endswith(('.txt', '.csv'))]
                if not txt:
                    raise ErrorSIRE('El .zip no contiene el TXT de la propuesta.')
                datos, nombre = z.read(txt[0]), txt[0].lower()
        except zipfile.BadZipFile as exc:
            raise ErrorSIRE('El archivo .zip está dañado.') from exc
    if nombre.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
        return [list(f) for f in wb.active.iter_rows(values_only=True)]
    for codificacion in ('utf-8-sig', 'latin-1'):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    primera = texto.splitlines()[0] if texto else ''
    sep = max('|;,\t', key=primera.count)
    return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=sep)]


def _texto(valor):
    """Texto de una celda: 1.0 → '1', 6011.0 → '6011'."""
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


# columnas propias (no son de SUNAT) y cómo se reconoce su encabezado: prefijos y nombres exactos
PROPIAS = {'oc': (('orden de compra', 'nro oc', 'n oc', 'nro. oc', 'numero de oc'), ('oc', 'o/c', 'o.c.', 'o.c')),
           'clasificacion': (('clasificacion',), ()),
           'cuenta': (('cuenta contable', 'cta contable', 'cta. contable'), ('cuenta', 'cta'))}


def _es_columna(encabezado, claves):
    prefijos, exactos = claves
    return encabezado.startswith(prefijos) or encabezado in exactos


def _es_propia(encabezado):
    e = _norm(encabezado)
    return any(_es_columna(e, claves) for claves in PROPIAS.values())


def clasificaciones():
    """[(número, código, nombre)] de la clasificación de las compras: el número es el que se escribe en el Excel."""
    from .models import Compra
    return [(n, codigo, nombre) for n, (codigo, nombre) in enumerate(Compra.CLASIFICACION, 1)]


def titulos_propios():
    leyenda = ', '.join(f'{n} {nombre.split(" (")[0]}' for n, _, nombre in clasificaciones())
    return ['Orden de compra', f'Clasificación ({leyenda})', 'Cuenta contable']


def _clasificacion(texto):
    """Código de la clasificación desde «1», «1 - Mercadería», «MERCADERIA» o «Mercadería»; None si no se reconoce."""
    lista = clasificaciones()
    m = re.match(r'\s*(\d+)', texto)
    if m:
        return next((codigo for n, codigo, _ in lista if n == int(m.group(1))), None)
    t = _norm(texto).replace(' ', '_')
    return next((codigo for _, codigo, nombre in lista if t in (codigo.lower(), _norm(nombre).replace(' ', '_'))
                 or _norm(nombre).startswith(_norm(texto))), None)


def _orden_compra(texto, ruc):
    """(orden, error) desde «OC01-00000007», «00000007» o «7»: debe ser del mismo proveedor y no estar anulada."""
    from .models import OrdenCompra
    qs = OrdenCompra.objects.select_related('tercero', 'centro_costo')
    candidatas = list(qs.filter(numero__iexact=texto))
    numero = re.search(r'(\d+)\s*$', texto)
    if not candidatas and numero:
        candidatas = list(qs.filter(numero__endswith=numero.group(1).zfill(8)))
    if not candidatas:
        return None, f'No existe la orden de compra {texto}.'
    propia = [o for o in candidatas if o.tercero.numero_doc == ruc]
    if not propia:
        return None, f'La orden {candidatas[0].numero} es de otro proveedor ({candidatas[0].tercero.nombre}).'
    oc = propia[0]
    if oc.estado == 'ANULADO':
        return None, f'La orden {oc.numero} está anulada.'
    return oc, ''


def _cuenta(texto):
    from contabilidad.models import CuentaContable
    cuenta = CuentaContable.objects.filter(codigo=texto.replace(' ', '')).first()
    if cuenta is None:
        return None, f'No existe la cuenta contable {texto}.'
    if not cuenta.imputable:
        return None, f'La cuenta {cuenta.codigo} no acepta movimientos: use una divisionaria.'
    return cuenta, ''


def propios(filas):
    """Anota en cada comprobante por precargar r['propios'] = {'oc', 'clasificacion', 'cuenta' (resueltos o None),
    'faltan' [campos vacíos], 'errores' [valores que no se reconocen]}."""
    for r in filas:
        p = {'oc': None, 'clasificacion': '', 'cuenta': None, 'faltan': [], 'errores': []}
        oc, clas, cta = (r.get(c) or '' for c in ('oc', 'clasificacion', 'cuenta'))
        if not oc:
            p['faltan'].append('orden de compra')
        else:
            p['oc'], error = _orden_compra(oc, r.get('ruc') or '')
            p['errores'] += [error] if error else []
        if not clas:
            p['faltan'].append('clasificación')
        else:
            p['clasificacion'] = _clasificacion(clas) or ''
            if not p['clasificacion']:
                p['errores'].append(f'Clasificación «{clas}» no reconocida (use 1 a {len(clasificaciones())}).')
        if not cta:
            p['faltan'].append('cuenta contable')
        else:
            p['cuenta'], error = _cuenta(cta)
            p['errores'] += [error] if error else []
        r['propios'] = p
    return filas


def _crudo(valor):
    """Valor de una celda de la propuesta tal como vino, guardable en JSON (sesión)."""
    if isinstance(valor, datetime):
        return valor.strftime('%d/%m/%Y %H:%M:%S') if valor.time() != datetime.min.time() else f'{valor:%d/%m/%Y}'
    if isinstance(valor, date):
        return f'{valor:%d/%m/%Y}'
    if isinstance(valor, Decimal):
        return str(valor)
    return valor


def leer(archivo):
    """[{'ruc', 'nombre', 'tipo', 'serie', 'numero', 'fecha', 'base', 'igv', 'total'}] de la propuesta."""
    return [{k: v for k, v in r.items() if k != 'sunat'} for r in leer_completo(archivo)[1]]


def leer_completo(archivo):
    """(encabezados, registros) de la propuesta con todas sus columnas: cada registro lleva además en 'sunat' los
    valores de la fila tal como los entrega SUNAT, en el orden de los encabezados."""
    filas = _filas(archivo)
    for i, fila in enumerate(filas[:20]):
        enc = [_norm(c) for c in fila]
        if any(e.startswith('serie') for e in enc) and any('total' in e for e in enc):
            break
    else:
        raise ErrorSIRE('No se reconoce el archivo: se espera la propuesta RCE de SIRE (columnas Tipo CP, Serie, '
                        'Nro CP, Nro Doc Identidad y Total CP).')

    def col(*claves, ultima=False):
        hallados = [j for j, e in enumerate(enc) if any(e.startswith(c) for c in claves)]
        if not hallados:
            return None
        return hallados[-1] if ultima else hallados[0]

    c = {'tipo': col('tipo cp', 'tipo de comprobante', 'tipo comprobante'), 'serie': col('serie'),
         'numero': col('nro cp', 'numero', 'nro comprobante', 'nro. cp'),
         'ruc': col('nro doc identidad', 'nro. doc identidad', 'numero de documento de identidad', 'ruc proveedor'),
         'nombre': col('apellidos nombres', 'apellidos y nombres', 'razon social', ultima=True),
         'fecha': col('fecha de emision', 'fecha emision'), 'base': col('bi grav', 'base imponible'),
         'igv': col('igv / ipm', 'igv'), 'total': col('total cp', 'importe total', 'total'),
         'tipo_doc': col('tipo doc identidad', 'tipo de documento de identidad'),
         'vencimiento': col('fecha vcto', 'fecha de vencimiento'), 'moneda': col('moneda', 'cod moneda'),
         'tc': col('tipo cambio', 'tipo de cambio'), 'exonerado': col('mto exonerado', 'valor adq. ng', 'adq no grav'),
         'inafecto': col('mto inafecto')}
    for campo, claves in PROPIAS.items():  # columnas propias agregadas en el Excel (opcionales)
        c[campo] = next((j for j, e in enumerate(enc) if _es_columna(e, claves)), None)
    if None in (c['serie'], c['numero'], c['ruc'], c['total']):
        raise ErrorSIRE('Faltan columnas en la propuesta (serie, número, documento del proveedor o total).')
    valor = lambda fila, k: fila[c[k]] if c[k] is not None and c[k] < len(fila) else ''
    registros = []
    for fila in filas[i + 1:]:
        if not any(str(v or '').strip() for v in fila):
            continue
        serie = str(valor(fila, 'serie') or '').strip().upper()
        numero = str(valor(fila, 'numero') or '').strip()
        if numero.endswith('.0'):
            numero = numero[:-2]
        if not serie or not numero:
            continue
        tipo = str(valor(fila, 'tipo') or '01').strip()
        registros.append({'ruc': str(valor(fila, 'ruc') or '').strip(), 'nombre': str(valor(fila, 'nombre') or ''),
                          'tipo': tipo.zfill(2) if tipo.isdigit() else tipo, 'serie': serie,
                          'numero': numero.lstrip('0') or '0', 'fecha': _fecha(valor(fila, 'fecha')),
                          'base': _dec(valor(fila, 'base')), 'igv': _dec(valor(fila, 'igv')),
                          'total': _dec(valor(fila, 'total')),
                          'tipo_doc': str(valor(fila, 'tipo_doc') or '').strip(),
                          'vencimiento': _fecha(valor(fila, 'vencimiento')),
                          'moneda': str(valor(fila, 'moneda') or 'PEN').strip().upper() or 'PEN',
                          'tc': _dec(valor(fila, 'tc')),
                          'no_gravado': _dec(valor(fila, 'exonerado')) + _dec(valor(fila, 'inafecto')),
                          **{campo: _texto(valor(fila, campo)) for campo in PROPIAS},
                          'sunat': [_crudo(v) for v in fila]})
    encabezados = [str(e if e is not None else '').strip() for e in filas[i]]
    return encabezados, registros


_IMPORTE = re.compile(r'^-?\d+\.\d+$')


def _celda(valor):
    """Importes del TXT ('1180.00', '3.750') como número en el Excel; códigos, RUC, series y números como texto."""
    if isinstance(valor, str) and _IMPORTE.match(valor.strip()):
        return Decimal(valor.strip())
    return valor


# columnas técnicas de la propuesta que no van al Excel: CAR, año y rango final del CDP, clasificación de bienes,
# operadores de proyecto, IMB, CAR original, tipo de nota, estado, incalificable y los campos libres CLU
_SIN_EXPORTAR = {'car sunat', 'ano', 'nro final (rango)', 'clasif de bss y sss', 'id proyecto operadores', 'porcpart',
                 'imb', 'car orig/ ind e o i', 'tipo de nota', 'est. comp.', 'est. comp', 'incal'}


def _se_exporta(encabezado):
    e = _norm(encabezado)
    return e not in _SIN_EXPORTAR and not re.fullmatch(r'clu\d*', e) and not _es_propia(encabezado)


def _clave_fila(r):
    return (r['tipo'], r['serie'], r['numero'], r['ruc'])


def _propias_de(g, compra):
    """Orden de compra, número de clasificación y cuenta: los del archivo o, si ya está registrada, los de la compra."""
    if any(g.get(c) for c in PROPIAS):
        return [g.get('oc') or None, g.get('clasificacion') or None, g.get('cuenta') or None]
    if compra is None or not hasattr(compra, 'orden_compra'):
        return [None, None, None]
    numero = next((n for n, codigo, _ in clasificaciones() if codigo == compra.clasificacion), None)
    return [compra.orden_compra.numero if compra.orden_compra_id else None, numero,
            compra.cuenta_contable.codigo if compra.cuenta_contable_id else None]


def filas_cruce_completo(encabezados, registros, r, extra=None, rendiciones=False, propias=False):
    """(encabezados, filas) del cruce con las columnas de la propuesta SUNAT (sin las técnicas de _SIN_EXPORTAR): el
    resultado del cruce, las columnas de SUNAT en su orden y al final el comprobante del ERP, su total y la diferencia. `extra` = (títulos, función(x))
    para columnas propias de cada pantalla (x es el renglón del cruce o None); `rendiciones` marca las compras que
    nadie rindió ni registró; `propias` agrega al final Orden de compra, Clasificación y Cuenta contable (se llenan
    antes de precargar las compras)."""
    titulos_extra, valor_extra = extra or ([], lambda x: [])
    if propias:
        titulos_extra = list(titulos_extra) + titulos_propios()
        valor_base = valor_extra

        def valor_extra(x, g=None, compra=None):
            return list(valor_base(x)) + _propias_de(g or {}, compra)
    if not registros:  # sin el archivo original: las columnas que el ERP leyó de la propuesta
        encabezados = ['Nro Doc Identidad', 'Apellidos Nombres/ Razón Social', 'Tipo CP/Doc.', 'Serie del CDP',
                       'Nro CP', 'Fecha de emisión', 'BI Gravado', 'IGV', 'Total CP', 'Moneda']
        registros = [{**x, 'sunat': [x['ruc'], x['nombre'], x['tipo'], x['serie'], x['numero'], x['fecha'],
                                     x['base'], x['igv'], x['total'], x.get('moneda')]}
                     for x in r['coinciden'] + r['diferencias'] + r['solo_sunat']]
    resultado = {}
    for x in r['coinciden']:
        resultado[_clave_fila(x)] = ('Coincide', x, 0)
    for x in r['diferencias']:
        resultado[_clave_fila(x)] = ('Diferencia de importe', x, x['dif'])
    for x in r['solo_sunat']:
        motivo = (' (registrado en otro periodo)' if x['otro_periodo'] else
                  ' (nadie la rindió)' if rendiciones and x['nadie_rindio'] else '')
        resultado[_clave_fila(x)] = ('Solo en SUNAT' + motivo, x, None)
    ancho = max([len(encabezados)] + [len(g.get('sunat') or []) for g in registros])
    sunat = list(encabezados) + [f'Columna {n + 1}' for n in range(len(encabezados), ancho)]
    van = [j for j, e in enumerate(sunat) if _se_exporta(e)]
    sunat = [sunat[j] for j in van]
    erp =['ERP: Doc. tercero', 'ERP: Nombre', 'ERP: Tipo', 'ERP: Serie', 'ERP: Número', 'ERP: Total',
           'Diferencia (ERP - SUNAT)']
    filas = []
    for g in registros:
        texto, x, dif = resultado.get(_clave_fila(g), ('', None, None))
        d = x.get('compra') if x else None
        valores = [_celda(v) for v in g.get('sunat') or []]
        valores += [None] * (ancho - len(valores))
        filas.append([texto] + [valores[j] for j in van] + _erp(d) + [dif] +
                     list(valor_extra(x, g, d) if propias else valor_extra(x)))
    for d in r['solo_erp']:
        filas.append(['Solo en el ERP'] + [None] * len(van) + _erp(d) + [None] +
                     list(valor_extra(None, None, d) if propias else valor_extra(None)))
    return ['Resultado del cruce'] + sunat + erp + list(titulos_extra), filas


def celdas_en_rojo(encabezados, filas):
    """[(fila, columna)] de las columnas propias vacías en los comprobantes por precargar (solo en SUNAT)."""
    columnas = [j for j, e in enumerate(encabezados) if _es_propia(e)]
    return [(i, j) for i, f in enumerate(filas) if str(f[0] or '').startswith('Solo en SUNAT')
            and 'otro periodo' not in str(f[0]) for j in columnas if f[j] in (None, '')]


def _erp(d):
    if d is None:
        return [None] * 6
    return [d.tercero.numero_doc, d.tercero.nombre, d.tipo_comprobante, d.serie, d.numero, d.total]


FECHAS, IMPORTES = ('fecha', 'vencimiento'), ('base', 'igv', 'total', 'tc', 'no_gravado')


def a_json(registros):
    """Registros leídos → datos guardables en JSON (sesión o PropuestaSIRE.registros)."""
    return [{**r, **{k: r[k].isoformat() if r.get(k) else None for k in FECHAS if k in r},
             **{k: str(r[k]) for k in IMPORTES if k in r}} for r in registros]


def de_json(datos):
    return [{**r, **{k: date.fromisoformat(r[k]) if r.get(k) else None for k in FECHAS if k in r},
             **{k: Decimal(r[k]) for k in IMPORTES if k in r}} for r in datos]


def _clave(ruc, tipo, serie, numero):
    return (ruc, tipo, serie.upper(), str(numero).lstrip('0') or '0')


def _rendidas(registros):
    """{clave: gasto rendido} de las facturas de la propuesta que llegaron por rendición de gastos."""
    from finanzas.models import LineaGasto
    rucs = {r['ruc'] for r in registros}
    lineas = (LineaGasto.objects.filter(ruc__in=rucs).exclude(estado='RECHAZADA')
              .exclude(rendicion__estado='RECHAZADA').select_related('rendicion__entrega__responsable'))
    return {_clave(l.ruc, l.tipo_comprobante, l.serie, l.numero): l for l in lineas}


def comparar(registros, periodo):
    """Cruza la propuesta con las compras registradas del periodo (vigentes). Marca las que vinieron por rendición
    de gastos y las que nadie rindió ni registró."""
    from .models import Compra
    compras = Compra.objects.filter(periodo=periodo).exclude(estado='ANULADO').select_related('tercero')
    erp = {_clave(c.tercero.numero_doc, c.tipo_comprobante, c.serie, c.numero): c for c in compras}
    sunat = {_clave(r['ruc'], r['tipo'], r['serie'], r['numero']): r for r in registros}
    rendidas = _rendidas(registros)
    coinciden, diferencias, solo_sunat, solo_erp = [], [], [], []
    for clave, r in sunat.items():
        c = erp.get(clave)
        rendida = rendidas.get(clave)
        if c is None:
            # registrada en otro periodo (SUNAT permite anotarla dentro de los 12 meses siguientes)
            otra = Compra.objects.filter(tercero__numero_doc=r['ruc'], tipo_comprobante=r['tipo'],
                                         serie=r['serie'], numero__in=[r['numero'], r['numero'].zfill(8)]) \
                .exclude(estado='ANULADO').first()
            solo_sunat.append({**r, 'otro_periodo': otra, 'rendicion': rendida,
                               'nadie_rindio': otra is None and rendida is None})
        elif abs(c.total - r['total']) > TOLERANCIA:
            diferencias.append({**r, 'compra': c, 'dif': c.total - r['total'], 'rendicion': rendida})
        else:
            coinciden.append({**r, 'compra': c, 'rendicion': rendida})
    for clave, c in erp.items():
        if clave not in sunat:
            solo_erp.append(c)
    return {'coinciden': coinciden, 'diferencias': diferencias, 'solo_sunat': solo_sunat, 'solo_erp': solo_erp,
            'por_rendicion': sum(1 for x in coinciden + diferencias if x['rendicion']),
            'nadie_rindio': [x for x in solo_sunat if x['nadie_rindio']]}


def _clave_venta(tipo, serie, numero):
    return (tipo, serie.upper(), str(numero).lstrip('0') or '0')


def comparar_ventas(registros, periodo):
    """Cruza la propuesta RVIE con los comprobantes de venta del periodo (la empresa es la emisora: la clave es tipo,
    serie y número)."""
    from ventas.models import Venta
    ventas = (Venta.objects.filter(periodo=periodo).exclude(estado='ANULADO').exclude(es_historico=True)
              .select_related('tercero'))
    erp = {_clave_venta(v.tipo_comprobante, v.serie, v.numero): v for v in ventas}
    sunat = {_clave_venta(r['tipo'], r['serie'], r['numero']): r for r in registros}
    coinciden, diferencias, solo_sunat = [], [], []
    for clave, r in sunat.items():
        v = erp.get(clave)
        if v is None:
            otra = Venta.objects.filter(tipo_comprobante=r['tipo'], serie=r['serie'],
                                        numero__in=[r['numero'], r['numero'].zfill(8)]).first()
            solo_sunat.append({**r, 'otro_periodo': otra, 'rendicion': None, 'nadie_rindio': otra is None})
        elif abs(v.total - r['total']) > TOLERANCIA:
            diferencias.append({**r, 'compra': v, 'dif': v.total - r['total'], 'rendicion': None})
        else:
            coinciden.append({**r, 'compra': v, 'rendicion': None})
    solo_erp = [v for clave, v in erp.items() if clave not in sunat]
    return {'coinciden': coinciden, 'diferencias': diferencias, 'solo_sunat': solo_sunat, 'solo_erp': solo_erp,
            'por_rendicion': 0, 'nadie_rindio': [x for x in solo_sunat if x['nadie_rindio']]}


def cruzar(libro, registros, periodo):
    return comparar_ventas(registros, periodo) if libro == 'RVIE' else comparar(registros, periodo)


def _tercero(r, tipo_tercero):
    from core.models import Tercero
    doc = (r.get('ruc') or '').strip()
    if not doc or doc in ('0', '-') or set(doc) == {'0'}:  # boleta a clientes varios
        tercero, _ = Tercero.objects.get_or_create(numero_doc='00000000', defaults={
            'nombre': 'CLIENTES VARIOS', 'tipo': tipo_tercero, 'tipo_doc': '0'})
        return tercero
    tipo_doc = r.get('tipo_doc') or ('6' if len(doc) == 11 else '1')
    tercero, creado = Tercero.objects.get_or_create(numero_doc=doc, defaults={
        'nombre': (r.get('nombre') or doc)[:200], 'tipo': tipo_tercero,
        'tipo_doc': tipo_doc if tipo_doc in ('6', '1', '4', '7', '0') else '0'})
    if not creado and tercero.tipo not in (tipo_tercero, 'AMBOS'):
        tercero.tipo = 'AMBOS'
        tercero.save(update_fields=['tipo'])
    return tercero


def crear_faltantes(libro, registros, periodo, claves, usuario=None):
    """Crea en el ERP los comprobantes de la propuesta marcados (los que están solo en SUNAT). Las ventas quedan
    como ya informadas a SUNAT (emitidas fuera del sistema): no se reenvían al OSE ni mueven stock. Las compras toman
    las columnas propias del Excel (orden de compra, clasificación y cuenta; opcionales): un valor que no se reconoce
    no deja crear esa compra. Devuelve (creados, errores)."""
    from django.db import IntegrityError, transaction

    from ventas.models import Venta

    from .models import Compra
    resultado = cruzar(libro, registros, periodo)
    faltan = {f'{r["tipo"]}|{r["serie"]}|{r["numero"]}|{r["ruc"]}': r for r in resultado['solo_sunat']
              if not r['otro_periodo']}
    creados, errores = 0, []
    for clave in claves:
        r = faltan.get(clave)
        if r is None:
            continue
        if libro != 'RVIE':
            p = propios([r])[0]['propios']
            if p['errores']:
                errores.append(f'{r["serie"]}-{r["numero"]}: no se creó. ' + ' '.join(p['errores']))
                continue
        try:
            with transaction.atomic():
                tercero = _tercero(r, 'CLIENTE' if libro == 'RVIE' else 'PROVEEDOR')
                base, igv, no_grav = r['base'], r['igv'], r.get('no_gravado') or D0
                total = r['total'] or base + igv + no_grav
                datos = dict(tipo_comprobante=r['tipo'], serie=r['serie'], numero=r['numero'],
                             fecha_emision=r['fecha'], fecha_vencimiento=r.get('vencimiento') or r['fecha'],
                             tercero=tercero, moneda=r.get('moneda') or 'PEN',
                             tipo_cambio=r.get('tc') or Decimal('1'), tipo_operacion='GRAVADA' if igv else 'EXONERADA',
                             base_imponible=base, no_gravado=no_grav, exonerado=no_grav, igv=igv, total=total,
                             periodo=periodo, glosa=f'Creado desde la propuesta SIRE {periodo}')
                if libro == 'RVIE':
                    Venta.objects.create(**datos, descontar_stock=False, estado_sunat='ACEPTADO',
                                         sunat_descripcion='Informado en SUNAT (propuesta SIRE): emitido fuera del '
                                                           'sistema')
                else:
                    oc = p['oc']
                    Compra.objects.create(**datos, orden_compra=oc, clasificacion=p['clasificacion'] or 'MERCADERIA',
                                          cuenta_contable=p['cuenta'], centro_costo=oc.centro_costo if oc else None)
            creados += 1
        except IntegrityError:
            errores.append(f'{r["serie"]}-{r["numero"]}: ya existe en el sistema.')
        except Exception as exc:  # dato de la propuesta que el modelo no acepta: se informa y se sigue
            errores.append(f'{r["serie"]}-{r["numero"]}: {exc}')
    return creados, errores
