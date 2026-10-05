"""Estados de cuenta bancarios: lectura de los formatos de BCP, BBVA e Interbank (Excel o CSV) y conciliación
automática contra los movimientos de tesorería.

Los bancos exportan el estado de cuenta con unas filas de cabecera (empresa, cuenta, periodo) antes de la tabla, y
cada uno nombra distinto las columnas. Por eso se busca la fila de encabezados entre las primeras filas y se reconoce
cada columna por sus nombres habituales:

- BCP: Fecha · Fecha valuta · Descripción operación · Monto · Saldo · Operación - Número
- BBVA: F. Operación · F. Valor · Concepto · Importe · Nº. Doc.
- Interbank: Fecha de operación · Nro. de operación · Descripción · Cargo · Abono · Saldo contable
- Plantilla del ERP: fecha · descripcion · monto · operacion
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction

D0 = Decimal('0')
TOLERANCIA_DIAS = 5
PALABRAS_GASTO = ('ITF', 'COMISION', 'MANTENIMIENTO', 'PORTES', 'MANT.', 'CARGO POR', 'COM.', 'IMPUESTO')

COLUMNAS = {
    'fecha': ['fecha de operacion', 'f. operacion', 'fecha operacion', 'fecha'],
    'descripcion': ['descripcion operacion', 'descripcion', 'concepto', 'detalle', 'glosa', 'movimiento'],
    'monto': ['monto', 'importe', 'importe (s/)', 'importe (us$)'],
    'cargo': ['cargo', 'cargos', 'debe', 'retiros'],
    'abono': ['abono', 'abonos', 'haber', 'depositos'],
    'saldo': ['saldo contable', 'saldo disponible', 'saldo'],
    'operacion': ['operacion - numero', 'nro. de operacion', 'n. de operacion', 'no. operacion', 'n° operacion',
                  'nro. operacion', 'numero de operacion', 'operacion', 'nº. doc.', 'n. doc.', 'no. doc.',
                  'nro. doc.', 'referencia'],
}
FIRMAS = [('BCP', ('fecha valuta', 'descripcion operacion')), ('BBVA', ('f. operacion', 'f. valor')),
          ('INTERBANK', ('fecha de operacion', 'cargo', 'abono'))]


class ErrorExtracto(Exception):
    pass


def _norm(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', texto.strip().lower())


def _filas_archivo(archivo):
    nombre = (getattr(archivo, 'name', '') or '').lower()
    contenido = archivo.read()
    if nombre.endswith(('.csv', '.txt')):
        texto = None
        for codificacion in ('utf-8-sig', 'latin-1'):
            try:
                texto = contenido.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        separador = ';' if texto.count(';') > texto.count(',') else (',' if ',' in texto else '\t')
        return [list(f) for f in csv.reader(io.StringIO(texto), delimiter=separador)]
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as exc:
        raise ErrorExtracto(f'No se pudo leer el archivo (use .xlsx o .csv): {exc}')
    return [list(f) for f in wb.active.iter_rows(values_only=True)]


def _columnas(encabezados):
    """{campo: índice} reconociendo cada encabezado por sus nombres habituales (el más específico primero)."""
    normal = [_norm(e) for e in encabezados]
    mapa = {}
    for campo, alias in COLUMNAS.items():
        for a in map(_norm, alias):
            if a in normal and normal.index(a) not in mapa.values():
                mapa[campo] = normal.index(a)
                break
    return mapa


def _numero(valor):
    if valor in (None, ''):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace('S/', '').replace('US$', '').replace('$', '').replace(' ', '')
    negativo = texto.startswith('(') and texto.endswith(')') or texto.endswith('-')
    texto = texto.strip('()-') if negativo else texto
    if re.search(r',\d{1,2}$', texto) and '.' in texto:  # 1.234,56
        texto = texto.replace('.', '').replace(',', '.')
    elif re.search(r',\d{1,2}$', texto):  # 1234,56
        texto = texto.replace(',', '.')
    else:
        texto = texto.replace(',', '')
    try:
        n = Decimal(texto)
    except InvalidOperation:
        return None
    return -n if negativo else n


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or '').strip()[:10]
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y', '%d.%m.%Y'):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            pass
    return None


def leer(archivo):
    """(formato, [{'fecha', 'descripcion', 'operacion', 'monto', 'saldo'}]) del estado de cuenta."""
    filas = _filas_archivo(archivo)
    for i, fila in enumerate(filas[:40]):
        mapa = _columnas(fila)
        if 'fecha' in mapa and ('monto' in mapa or ('cargo' in mapa and 'abono' in mapa)):
            break
    else:
        raise ErrorExtracto('No se encontró la tabla de movimientos: se esperan columnas de fecha, descripción y '
                            'monto (o cargo y abono). Use el archivo Excel/CSV que descarga de la banca por '
                            'internet o la plantilla.')
    normal = {_norm(e) for e in fila}
    formato = next((banco for banco, firma in FIRMAS if all(_norm(f) in normal for f in firma)), 'PLANTILLA')
    lineas = []
    for fila in filas[i + 1:]:
        celda = lambda campo: fila[mapa[campo]] if campo in mapa and mapa[campo] < len(fila) else None
        fecha = _fecha(celda('fecha'))
        if fecha is None:
            continue  # filas de totales, saldos o pie de página
        if 'monto' in mapa:
            monto = _numero(celda('monto'))
        else:
            monto = (_numero(celda('abono')) or D0) - abs(_numero(celda('cargo')) or D0)
        if not monto:
            continue
        operacion = celda('operacion')
        if isinstance(operacion, float) and operacion.is_integer():
            operacion = int(operacion)
        lineas.append({'fecha': fecha, 'descripcion': str(celda('descripcion') or '').strip()[:250],
                       'operacion': str(operacion or '').strip()[:40], 'monto': monto.quantize(Decimal('0.01')),
                       'saldo': _numero(celda('saldo'))})
    if not lineas:
        raise ErrorExtracto('El archivo no tiene movimientos con fecha y monto.')
    return formato, lineas


@transaction.atomic
def cargar(cuenta, archivo, usuario=None):
    """Registra el extracto (sin repetir líneas ya cargadas antes) y concilia automáticamente."""
    from .models import Extracto, LineaExtracto
    formato, lineas = leer(archivo)
    existentes = set(LineaExtracto.objects.filter(extracto__cuenta=cuenta).values_list(
        'fecha', 'monto', 'operacion', 'descripcion'))
    nuevas = [l for l in lineas if (l['fecha'], l['monto'], l['operacion'], l['descripcion']) not in existentes]
    con_saldo = [l for l in lineas if l['saldo'] is not None]
    extracto = Extracto.objects.create(
        cuenta=cuenta, formato=formato, nombre_archivo=getattr(archivo, 'name', '')[:200],
        desde=min(l['fecha'] for l in lineas), hasta=max(l['fecha'] for l in lineas), cargado_por=usuario,
        saldo_final=max(con_saldo, key=lambda l: l['fecha'])['saldo'] if con_saldo else None)
    LineaExtracto.objects.bulk_create([LineaExtracto(extracto=extracto, **l) for l in nuevas])
    return extracto, len(lineas) - len(nuevas), conciliar(extracto)


def _digitos(texto):
    return re.sub(r'\D', '', texto or '').lstrip('0')


def conciliar(extracto):
    """Empareja cada línea pendiente con un movimiento no conciliado de la cuenta, del mismo tipo e importe:
    1) mismo número de operación; 2) la fecha más cercana dentro de ±5 días. Cada movimiento se usa una vez."""
    from .models import Movimiento
    pendientes = list(extracto.lineas.filter(estado='PENDIENTE'))
    if not pendientes:
        return 0
    desde = min(l.fecha for l in pendientes) - timedelta(days=TOLERANCIA_DIAS)
    hasta = max(l.fecha for l in pendientes) + timedelta(days=TOLERANCIA_DIAS)
    libres = list(Movimiento.objects.filter(cuenta=extracto.cuenta, conciliado=False, fecha__range=[desde, hasta],
                                            lineas_extracto__isnull=True))
    n = 0
    for regla in ('operacion', 'fecha'):
        for linea in pendientes:
            if linea.estado != 'PENDIENTE':
                continue
            candidatos = [m for m in libres if m.tipo == linea.tipo and m.monto == abs(linea.monto)
                          and abs((m.fecha - linea.fecha).days) <= TOLERANCIA_DIAS]
            if regla == 'operacion':
                op = _digitos(linea.operacion)
                candidatos = [m for m in candidatos if op and len(op) >= 3 and _digitos(m.numero_operacion) and (
                    _digitos(m.numero_operacion).endswith(op) or op.endswith(_digitos(m.numero_operacion)))]
                texto = 'Mismo importe y N° de operación'
            else:
                texto = 'Mismo importe y fecha cercana'
            if not candidatos:
                continue
            mov = min(candidatos, key=lambda m: (abs((m.fecha - linea.fecha).days), m.pk))
            _enlazar(linea, mov, texto)
            libres.remove(mov)
            n += 1
    return n


def _enlazar(linea, mov, regla, estado='CONCILIADA'):
    mov.conciliado, mov.fecha_conciliacion = True, linea.fecha
    mov.save(update_fields=['conciliado', 'fecha_conciliacion'])
    linea.movimiento, linea.estado, linea.regla = mov, estado, regla
    linea.save(update_fields=['movimiento', 'estado', 'regla'])


def enlazar_manual(linea, mov):
    if mov.cuenta_id != linea.extracto.cuenta_id or mov.tipo != linea.tipo:
        raise ErrorExtracto('El movimiento debe ser de la misma cuenta y del mismo tipo (ingreso / egreso).')
    if mov.conciliado:
        raise ErrorExtracto(f'El movimiento {mov.voucher} ya está conciliado.')
    _enlazar(linea, mov, 'Manual' if mov.monto == abs(linea.monto) else
             f'Manual (diferencia {abs(linea.monto) - mov.monto:+.2f})')


def es_gasto_bancario(linea):
    return linea.monto < 0 and any(p in linea.descripcion.upper() for p in PALABRAS_GASTO)


def crear_movimiento(linea, usuario=None, concepto=None):
    """Registra en tesorería una línea del banco que no estaba (comisiones, ITF, abonos no identificados)."""
    from .models import Movimiento
    if linea.estado != 'PENDIENTE':
        raise ErrorExtracto('La línea ya fue conciliada o ignorada.')
    mov = Movimiento.objects.create(
        cuenta=linea.extracto.cuenta, fecha=linea.fecha, tipo=linea.tipo,
        concepto=concepto or ('GASTO_BANCARIO' if es_gasto_bancario(linea) else 'OTRO'),
        medio_pago='TRANSFERENCIA', numero_operacion=linea.operacion, monto=abs(linea.monto),
        glosa=linea.descripcion or 'Según estado de cuenta', creado_por=usuario)
    _enlazar(linea, mov, 'Creado desde el extracto', estado='CREADA')
    return mov


def quitar(linea):
    """Deshace la conciliación de la línea (el movimiento creado desde ella se conserva, sin conciliar)."""
    if linea.movimiento_id:
        linea.movimiento.conciliado, linea.movimiento.fecha_conciliacion = False, None
        linea.movimiento.save(update_fields=['conciliado', 'fecha_conciliacion'])
    linea.movimiento, linea.estado, linea.regla = None, 'PENDIENTE', ''
    linea.save(update_fields=['movimiento', 'estado', 'regla'])
