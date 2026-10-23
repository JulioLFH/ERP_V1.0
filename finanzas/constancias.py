"""Constancias de depósito de detracciones desde el reporte de SUNAT (Operaciones en Línea › Detracciones ›
Consulta de constancias), exportado a Excel (.xlsx), CSV o TXT. SUNAT da una constancia por cada comprobante.

Las columnas se reconocen por su nombre: N° de constancia, fecha de pago, importe o monto depositado, tipo, serie y
número del comprobante, y los documentos del adquirente (cliente) y del proveedor. También sirve una planilla propia
con al menos «serie», «número» y «constancia».

- Clientes (ventas): lo que los clientes depositaron en la cuenta del BN de la empresa → cobranzas de detracción.
- Compras: las constancias de un lote de depósito masivo.
"""
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import Movimiento
from .operaciones import ErrorTesoreria, _periodo_abierto


def _norm(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', texto)).strip()


def _tiene(h, *palabras):
    return any(p in h for p in palabras)


# (campo, prueba sobre el encabezado normalizado); se toma la primera columna que cumple
COLUMNAS = [
    ('constancia', lambda h: 'constancia' in h and not _tiene(h, 'fecha')),
    ('fecha', lambda h: 'fecha' in h and _tiene(h, 'pago', 'deposito')),
    ('fecha', lambda h: 'fecha' in h and not _tiene(h, 'emision', 'comprobante')),
    ('importe', lambda h: _tiene(h, 'importe', 'monto') and not _tiene(h, 'total del lote')),
    ('tipo', lambda h: 'tipo' in h and _tiene(h, 'comprobante', 'cpe', 'cdp')),
    ('serie', lambda h: 'serie' in h),
    ('numero', lambda h: _tiene(h, 'numero', 'nro', 'num') and _tiene(h, 'comprobante', 'cpe', 'cdp')
     and not _tiene(h, 'serie')),
    ('numero', lambda h: h in ('numero', 'nro', 'n', 'num', 'correlativo')),
    ('comprobante', lambda h: h == 'comprobante'),
    ('ruc_adquirente', lambda h: _tiene(h, 'adquir', 'cliente', 'usuario') and not _tiene(h, 'tipo', 'nombre',
                                                                                         'razon')),
    ('ruc_proveedor', lambda h: _tiene(h, 'proveedor') and not _tiene(h, 'tipo', 'nombre', 'razon')),
]


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    return texto[:-2] if re.fullmatch(r'\d+\.0', texto) else texto


def _fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(_texto(valor)[:10], fmt).date()
        except ValueError:
            continue
    return None


def _importe(valor):
    try:
        return Decimal(_texto(valor).replace(',', '')) if _texto(valor) else None
    except InvalidOperation:
        return None


def _filas(archivo):
    nombre = (getattr(archivo, 'name', '') or '').lower()
    if nombre.endswith('.xls'):
        raise ErrorTesoreria('El archivo .xls (Excel antiguo) no se puede leer: ábralo en Excel y guárdelo como .xlsx '
                             'o .csv.')
    datos = archivo.read()
    if nombre.endswith('.xlsx') or datos[:2] == b'PK':
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(datos), data_only=True, read_only=True).active
        return [list(f) for f in hoja.iter_rows(values_only=True)]
    for codificacion in ('utf-8-sig', 'latin-1'):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    lineas = [l for l in texto.splitlines() if l.strip()]
    if not lineas:
        return []
    sep = max('|;,\t', key=lambda s: sum(l.count(s) for l in lineas[:20]))
    return [list(f) for f in csv.reader(io.StringIO('\n'.join(lineas)), delimiter=sep)]


def leer(archivo):
    """[{'constancia', 'fecha', 'importe', 'tipo', 'serie', 'numero', 'ruc_adquirente', 'ruc_proveedor'}]."""
    filas = _filas(archivo)
    for i, fila in enumerate(filas[:20]):
        encabezados = [_norm(c) for c in fila]
        if any('constancia' in h for h in encabezados):
            break
    else:
        raise ErrorTesoreria('No se encontró la columna del N° de constancia en el archivo.')
    columnas = {}
    for campo, prueba in COLUMNAS:
        if campo in columnas:
            continue
        j = next((j for j, h in enumerate(encabezados) if prueba(h) and j not in columnas.values()), None)
        if j is not None:
            columnas[campo] = j
    if 'serie' not in columnas and 'comprobante' not in columnas:
        raise ErrorTesoreria('El archivo debe traer la serie y el número del comprobante de cada constancia.')
    valor = lambda fila, campo: fila[columnas[campo]] if campo in columnas and columnas[campo] < len(fila) else None
    registros = []
    for fila in filas[i + 1:]:
        constancia = _texto(valor(fila, 'constancia'))
        if not constancia or not re.search(r'\d', constancia):
            continue
        serie, numero = _texto(valor(fila, 'serie')).upper(), _texto(valor(fila, 'numero'))
        if not serie and '-' in _texto(valor(fila, 'comprobante')):
            serie, numero = _texto(valor(fila, 'comprobante')).upper().split('-', 1)
        tipo = _texto(valor(fila, 'tipo'))
        registros.append({
            'constancia': constancia, 'fecha': _fecha(valor(fila, 'fecha')), 'importe': _importe(valor(fila, 'importe')),
            'tipo': tipo.zfill(2) if tipo.isdigit() else '', 'serie': serie.strip(),
            'numero': numero.strip().lstrip('0') or '0',
            'ruc_adquirente': re.sub(r'\D', '', _texto(valor(fila, 'ruc_adquirente'))),
            'ruc_proveedor': re.sub(r'\D', '', _texto(valor(fila, 'ruc_proveedor')))})
    return registros


# ---------------------------------------------------------------- clientes: cobranzas de detracción
def emparejar_clientes(registros):
    """Cada constancia con su factura de venta y lo que se hará: REGISTRAR, YA (ya registrada), NO_ENCONTRADA o
    SIN_DETRACCION (la venta no tiene detracción por cobrar)."""
    from ventas.models import Venta
    from .multiple import detraccion_pendiente
    salida = []
    for r in registros:
        qs = Venta.objects.filter(serie=r['serie'], numero__in=[r['numero'], r['numero'].zfill(8)],
                                  estado='REGISTRADO').select_related('tercero')
        if r['tipo']:
            qs = qs.filter(tipo_comprobante=r['tipo'])
        if r['ruc_adquirente']:
            qs = qs.filter(tercero__numero_doc=r['ruc_adquirente'])
        venta = qs.first()
        item = {**r, 'venta': venta, 'pendiente': None}
        if venta is None:
            item['estado'] = 'NO_ENCONTRADA'
        elif venta.detraccion_constancia == r['constancia'] or venta.movimientos.filter(
                concepto='DETRACCION', numero_operacion=r['constancia']).exists():
            item['estado'] = 'YA'
        else:
            item['pendiente'] = detraccion_pendiente(venta) * venta.tc_efectivo
            item['estado'] = 'REGISTRAR' if item['pendiente'] > 0 else 'SIN_DETRACCION'
        if item['importe'] is None and item['pendiente']:
            item['importe'] = item['pendiente']
        salida.append(item)
    return salida


@transaction.atomic
def registrar_clientes(items, cuenta, usuario):
    """Cobranza de detracción (concepto Detracción, en la cuenta del BN) por cada constancia a registrar."""
    if not cuenta.es_detracciones:
        raise ErrorTesoreria('Elija su cuenta de detracciones del Banco de la Nación.')
    creados = []
    for it in items:
        if it['estado'] != 'REGISTRAR':
            continue
        venta, fecha = it['venta'], it['fecha']
        if fecha is None or not it['importe']:
            raise ErrorTesoreria(f'{venta}: la constancia {it["constancia"]} no tiene fecha o importe.')
        _periodo_abierto(fecha)
        creados.append(Movimiento.objects.create(
            cuenta=cuenta, fecha=fecha, tipo='INGRESO', concepto='DETRACCION', medio_pago='DEPOSITO',
            numero_operacion=it['constancia'][:40], tercero=venta.tercero, venta=venta, monto=it['importe'],
            creado_por=usuario, glosa=f'Detracción de {venta} (constancia SUNAT {it["constancia"]})'[:250]))
    return creados


# ---------------------------------------------------------------- compras: constancias de un lote
def emparejar_lote(registros, lote):
    """({id de línea: constancia}, [no encontradas]) para un lote de depósito masivo."""
    def clave(serie, numero):
        return serie.strip().upper(), (numero.strip().lstrip('0') or '0')
    por_doc = {}
    for l in lote.lineas.select_related('compra__tercero'):
        por_doc.setdefault(clave(l.compra.serie, l.compra.numero), []).append(l)
    constancias, faltan = {}, []
    for r in registros:
        candidatas = [l for l in por_doc.get(clave(r['serie'], r['numero']), [])
                      if not r['ruc_proveedor'] or l.compra.tercero.numero_doc == r['ruc_proveedor']]
        if candidatas:
            constancias[candidatas[0].pk] = r['constancia']
        else:
            faltan.append(f'{r["serie"]}-{r["numero"]}')
    return constancias, faltan
