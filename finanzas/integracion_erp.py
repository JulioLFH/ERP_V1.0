"""Integración con un ERP externo (cliente que mantiene SAP u otro): exporta los gastos rendidos y contabilizados
con cuenta, centro de costo, indicador de impuesto y enlace al sustento; recibe la respuesta del ERP (aceptado o
error por documento), reintenta los errores y concilia Ceiba frente al ERP externo.

Archivo de exportación (CSV con «;»): una fila por gasto y centro de costo.
Archivo de respuesta (CSV o Excel): columnas id (el id_envio de la exportación), estado (OK / ERROR), documento
(número en el ERP externo) y mensaje."""
import csv
import io
from collections import defaultdict
from decimal import Decimal

from django.db import transaction

from .models import EnvioERP, EnvioERPDetalle, LineaGasto

D0 = Decimal('0')
COLUMNAS = ['id_envio', 'rendicion', 'fecha_contable', 'fecha_documento', 'tipo_comprobante', 'documento', 'ruc',
            'proveedor', 'cuenta', 'cuenta_ceiba', 'centro_costo', 'indicador_impuesto', 'moneda', 'base', 'igv',
            'total', 'importe_pen', 'porcentaje', 'responsable', 'glosa', 'sustento']


class ErrorIntegracion(Exception):
    pass


def _pen(linea):
    from .rendiciones import _pen as en_soles
    return en_soles(linea)


def por_exportar():
    """Gastos contabilizados que el ERP externo aún no aceptó y no están en un envío sin respuesta."""
    pendientes = EnvioERPDetalle.objects.filter(estado__in=['ENVIADO', 'ACEPTADO']).values('linea_id')
    return (LineaGasto.objects.filter(rendicion__estado__in=['CONTABILIZADA', 'REEMBOLSADA', 'DEVUELTA'],
                                      estado='APROBADA').exclude(pk__in=pendientes)
            .select_related('rendicion__entrega__responsable', 'categoria', 'indicador', 'centro_costo'))


@transaction.atomic
def generar(usuario, solo_errores=False):
    """Nuevo lote con lo pendiente (o, en un reintento, con lo que el ERP devolvió con error)."""
    if solo_errores:
        ids = set(EnvioERPDetalle.objects.filter(estado='ERROR').values_list('linea_id', flat=True)) - set(
            EnvioERPDetalle.objects.filter(estado__in=['ENVIADO', 'ACEPTADO']).values_list('linea_id', flat=True))
        lineas = LineaGasto.objects.filter(pk__in=ids).select_related('rendicion__entrega')
    else:
        lineas = por_exportar()
    lineas = list(lineas)
    if not lineas:
        raise ErrorIntegracion('No hay gastos por exportar.' if not solo_errores else 'No hay errores por reintentar.')
    envio = EnvioERP.objects.create(usuario=usuario if usuario and usuario.is_authenticated else None,
                                    es_reintento=solo_errores)
    for l in lineas:
        intento = EnvioERPDetalle.objects.filter(linea=l).count() + 1
        EnvioERPDetalle.objects.create(envio=envio, linea=l, importe_pen=_pen(l), intento=intento)
    return envio


def filas(envio, url_sustento=lambda s: ''):
    """Filas del archivo del lote: una por gasto y por centro de costo (si está distribuido)."""
    from core.sustentos import de
    salida = []
    for d in envio.detalles.select_related('linea__rendicion__entrega__responsable', 'linea__categoria',
                                           'linea__indicador', 'linea__centro_costo', 'linea__proyecto',
                                           'linea__compra'):
        l = d.linea
        r = l.rendicion
        cuenta = l.categoria.cuenta_en(l.fecha)
        externa = getattr(getattr(cuenta, 'mapeo_erp', None), 'codigo_externo', '') or cuenta.codigo
        sustentos = list(de(l))
        enlace = url_sustento(sustentos[0]) if sustentos else ''
        partes = [(x.centro_costo, x.porcentaje) for x in l.distribucion.select_related('centro_costo')] or \
                 [(l.centro_costo or (l.proyecto.centro_costo if l.proyecto_id else None), Decimal('100'))]
        for centro, pct in partes:
            factor = pct / 100
            salida.append({
                'id_envio': d.pk, 'rendicion': r.numero,
                'fecha_contable': r.fecha_contable.isoformat() if r.fecha_contable else '',
                'fecha_documento': l.fecha.isoformat(), 'tipo_comprobante': l.tipo_comprobante,
                'documento': l.documento, 'ruc': l.ruc, 'proveedor': l.razon_social, 'cuenta': externa,
                'cuenta_ceiba': cuenta.codigo, 'centro_costo': centro.codigo if centro else '',
                'indicador_impuesto': (l.indicador.codigo_externo or l.indicador.codigo) if l.indicador_id else '',
                'moneda': l.moneda, 'base': (l.base * factor).quantize(Decimal('0.01')),
                'igv': (l.igv * factor).quantize(Decimal('0.01')), 'total': (l.total * factor).quantize(Decimal('0.01')),
                'importe_pen': (d.importe_pen * factor).quantize(Decimal('0.01')), 'porcentaje': pct,
                'responsable': r.entrega.responsable.numero_doc, 'glosa': l.glosa, 'sustento': enlace})
    return salida


def csv_de(envio, url_sustento=lambda s: ''):
    salida = io.StringIO()
    salida.write('﻿')
    w = csv.DictWriter(salida, fieldnames=COLUMNAS, delimiter=';')
    w.writeheader()
    for f in filas(envio, url_sustento):
        w.writerow(f)
    return salida.getvalue()


def _leer(archivo):
    nombre = (getattr(archivo, 'name', '') or '').lower()
    datos = archivo.read()
    if nombre.endswith('.xlsx'):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
        tabla = [list(f) for f in wb.active.iter_rows(values_only=True)]
    else:
        texto = datos.decode('utf-8-sig', errors='replace')
        sep = ';' if texto.splitlines()[0].count(';') >= texto.splitlines()[0].count(',') else ','
        tabla = list(csv.reader(io.StringIO(texto), delimiter=sep))
    if not tabla:
        raise ErrorIntegracion('El archivo de respuesta está vacío.')
    enc = [str(c or '').strip().lower() for c in tabla[0]]
    col = {k: next((i for i, e in enumerate(enc) if e in alias), None) for k, alias in (
        ('id', ('id', 'id_envio')), ('estado', ('estado', 'status')),
        ('documento', ('documento', 'documento_erp', 'belnr', 'referencia')), ('mensaje', ('mensaje', 'message')))}
    if col['id'] is None or col['estado'] is None:
        raise ErrorIntegracion('La respuesta debe tener las columnas id y estado (OK o ERROR).')
    valor = lambda f, k: str(f[col[k]] if col[k] is not None and col[k] < len(f) and f[col[k]] is not None else '')
    return [(valor(f, 'id').strip(), valor(f, 'estado').strip().upper(), valor(f, 'documento').strip(),
             valor(f, 'mensaje').strip()) for f in tabla[1:] if any(c not in (None, '') for c in f)]


@transaction.atomic
def registrar_respuesta(envio, archivo):
    """Marca cada documento del lote como aceptado (con su número en el ERP externo) o con error (con el
    mensaje). Devuelve (aceptados, errores, no encontrados)."""
    detalles = {str(d.pk): d for d in envio.detalles.all()}
    ok = mal = desconocidos = 0
    for ident, estado, documento, mensaje in _leer(archivo):
        d = detalles.get(ident.split('.')[0])
        if d is None:
            desconocidos += 1
            continue
        if estado in ('OK', 'ACEPTADO', 'S', 'SUCCESS'):
            d.estado, d.referencia_externa, d.mensaje = 'ACEPTADO', documento[:40], ''
            ok += 1
        else:
            d.estado, d.mensaje = 'ERROR', (mensaje or estado or 'Rechazado por el ERP externo')[:300]
            mal += 1
        d.save()
    restantes = envio.detalles.exclude(estado='ACEPTADO').exists()
    envio.estado = 'PARCIAL' if restantes else 'CONCILIADO'
    envio.save(update_fields=['estado'])
    return ok, mal, desconocidos


def conciliacion():
    """Por cuenta: lo contabilizado en Ceiba (gastos de rendiciones) frente a lo aceptado por el ERP externo."""
    ceiba, erp = defaultdict(lambda: D0), defaultdict(lambda: D0)
    lineas = LineaGasto.objects.filter(rendicion__estado__in=['CONTABILIZADA', 'REEMBOLSADA', 'DEVUELTA'],
                                       estado='APROBADA').select_related('categoria', 'rendicion__entrega')
    for l in lineas:
        ceiba[l.categoria.cuenta_en(l.fecha).codigo] += _pen(l)
    aceptados = EnvioERPDetalle.objects.filter(estado='ACEPTADO').select_related('linea__categoria')
    for d in aceptados:
        erp[d.linea.categoria.cuenta_en(d.linea.fecha).codigo] += d.importe_pen
    filas_ = [{'cuenta': c, 'ceiba': ceiba[c], 'erp': erp[c], 'diferencia': ceiba[c] - erp[c]}
              for c in sorted(set(ceiba) | set(erp))]
    return {'filas': filas_, 'ceiba': sum((f['ceiba'] for f in filas_), D0), 'erp': sum((f['erp'] for f in filas_), D0),
            'errores': EnvioERPDetalle.objects.filter(estado='ERROR').exclude(
                linea__in=EnvioERPDetalle.objects.filter(estado='ACEPTADO').values('linea')).select_related(
                'linea__rendicion', 'envio')}
