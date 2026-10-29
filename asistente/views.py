"""Pantallas del asistente de planta «Lucho»: el chat al pie de la máquina, los casos, las fichas de parámetros (con
su aprobación e importación), la base de conocimiento y la configuración de procesos y defectos."""
from datetime import timedelta
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from core.permisos import puede

from . import motor, plantillas
from . import servicios as srv
from .forms import ConfigForm, DefectoForm, FichaForm, ManualForm, ParametroForm, ProcesoForm
from .models import (AjusteDefecto, CasoPlanta, ConfigAsistente, Consulta, Defecto, FichaParametros, Manual,
                     ParametroProceso, Proceso)


def _entero(valor):
    return int(valor) if str(valor or '').isdigit() else None


def _contexto(request, caso=None):
    """Lo que Lucho sabe sin teclear: el del caso abierto o el de la orden elegida (y el equipo si se cambia)."""
    from core.models import Producto
    from produccion.models import Equipo, OrdenProduccion
    if caso is not None:
        ficha = caso.ficha or srv.vigente(caso.producto, caso.equipo, caso.proceso, caso.molde)
        return {'orden': caso.orden, 'producto': caso.producto, 'equipo': caso.equipo, 'molde': caso.molde,
                'proceso': caso.proceso or (ficha.proceso if ficha else None), 'material': caso.material,
                'lote': caso.lote, 'ficha': ficha, 'puesto': caso.equipo.centro if caso.equipo_id else None,
                'equipos': [caso.equipo] if caso.equipo_id else [], 'moldes': []}
    datos = request.POST if request.method == 'POST' else request.GET
    orden = OrdenProduccion.objects.filter(pk=_entero(datos.get('orden')),
                                           estado__in=OrdenProduccion.ABIERTAS).first()
    ctx = srv.contexto_de_orden(orden) if orden else {'equipos': [], 'moldes': []}
    equipo = Equipo.objects.filter(pk=_entero(datos.get('equipo')), activo=True).first()
    if equipo is not None:
        ctx['equipo'] = equipo
    if orden is None:
        ctx['producto'] = Producto.objects.filter(pk=_entero(datos.get('producto'))).first()
        ctx['proceso'] = Proceso.objects.filter(pk=_entero(datos.get('proceso')), activo=True).first() or (
            Proceso.objects.filter(activo=True, puestos=equipo.centro_id).first() if equipo and equipo.centro_id
            else None)
    if ctx.get('equipo') is not None and ctx.get('producto') is not None:
        ctx['ficha'] = srv.vigente(ctx['producto'], ctx['equipo'], ctx.get('proceso'), ctx.get('molde'))
        if ctx['ficha'] and not ctx.get('proceso'):
            ctx['proceso'] = ctx['ficha'].proceso
    ctx['puesto'] = ctx['equipo'].centro if ctx.get('equipo') is not None else None
    return ctx


@login_required
def chat(request):
    """Preguntar a Lucho: foto primero (texto opcional), con la orden abierta; la conversación es el caso."""
    from core.sustentos import ErrorSustento, adjuntar, validar_imagen
    caso = CasoPlanta.objects.filter(pk=_entero(request.GET.get('caso') or request.POST.get('caso'))).first()
    ctx = _contexto(request, caso)
    if request.method == 'POST':
        texto = (request.POST.get('texto') or '').strip()[:2000]
        foto = request.FILES.get('foto')
        defecto = Defecto.objects.filter(pk=_entero(request.POST.get('defecto')), activo=True).first()
        if foto:
            try:
                validar_imagen(foto)
            except ErrorSustento as exc:
                messages.error(request, f'Foto: {exc}')
                return redirect(request.get_full_path())
        if not (texto or foto or defecto):
            messages.error(request, 'Tome la foto del defecto, elija el defecto o escriba una línea.')
            return redirect(request.get_full_path())
        # LUC-05: lo ajustado son los parámetros que el operario dijo haber cambiado (no todo lo recomendado)
        ajustados = [n for c in (caso.consultas.all() if caso else []) for n in (c.ajustes or {})]
        r = motor.responder(texto, ctx, defecto, ajustados=ajustados, caso=caso, con_foto=bool(foto))
        if r['fuera_de_dominio']:  # LUC-10: queda con el caso (y la orden) en la bitácora
            motor.registrar_consulta(request.user, ctx, texto, r, caso=caso, con_foto=bool(foto))
            messages.warning(request, r['texto'])
            return redirect(f'{reverse("asistente:chat")}?caso={caso.pk}' if caso else request.get_full_path())
        if caso is None:
            caso = srv.abrir_caso(request.user, ctx, defecto=r['defecto'], problema=texto, foto=foto)
        else:
            if foto:
                adjuntar(caso, foto, request.user, 'Foto del defecto')
            if caso.defecto_id is None and r['defecto'] is not None:
                caso.defecto = r['defecto']
                caso.save(update_fields=['defecto'])
        motor.registrar_consulta(request.user, ctx, texto, r, caso=caso, con_foto=bool(foto))
        return redirect(f'{reverse("asistente:chat")}?caso={caso.pk}#fin')
    defectos = Defecto.objects.filter(activo=True).select_related('proceso')
    if ctx.get('proceso') is not None:
        defectos = defectos.filter(proceso=ctx['proceso'])
    from produccion.models import Equipo
    return render(request, 'asistente/chat.html', {
        'caso': caso, 'ctx': ctx, 'consultas': list(caso.consultas.order_by('creado', 'id')) if caso else [],
        'defectos': defectos, 'ordenes': srv.ordenes_abiertas()[:60],
        'equipos': Equipo.objects.filter(activo=True, tipo__in=srv.MAQUINAS).order_by('codigo'),
        'procesos': Proceso.objects.filter(activo=True),
        'puede_registrar': puede(request.user, 'asistente.registrar_caso')})


# ---------------------------------------------------------------- casos
@login_required
def casos(request):
    qs = CasoPlanta.objects.select_related('producto', 'equipo', 'molde', 'defecto', 'orden', 'autor')
    estado = request.GET.get('estado', '')
    if estado == 'SIN_RESOLVER':
        qs = qs.exclude(estado='RESUELTO')
    elif estado in dict(CasoPlanta.ESTADOS):
        qs = qs.filter(estado=estado)
    for campo in ('equipo', 'molde', 'producto', 'defecto'):
        if _entero(request.GET.get(campo)):
            qs = qs.filter(**{f'{campo}_id': _entero(request.GET[campo])})
    q = (request.GET.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(numero__icontains=q) | Q(problema__icontains=q) | Q(causa__icontains=q) |
                       Q(accion__icontains=q) | Q(lote__icontains=q) | Q(producto__nombre__icontains=q))
    from produccion.models import Equipo
    return render(request, 'asistente/casos.html', {
        'casos': qs[:300], 'estado': estado, 'q': q,
        'equipos': Equipo.objects.filter(activo=True, tipo__in=srv.MAQUINAS).order_by('codigo'),
        'moldes': Equipo.objects.filter(activo=True, tipo='MOLDE').order_by('codigo'),
        'defectos': Defecto.objects.filter(activo=True).select_related('proceso'),
        'f_equipo': _entero(request.GET.get('equipo')), 'f_molde': _entero(request.GET.get('molde')),
        'f_defecto': _entero(request.GET.get('defecto'))})


@login_required
def caso(request, pk):
    c = get_object_or_404(CasoPlanta.objects.select_related('producto', 'equipo', 'molde', 'proceso', 'ficha',
                                                            'defecto', 'orden', 'autor'), pk=pk)
    if request.method == 'POST':
        p = request.POST
        filas = [{'nombre': p.get(f'param_nombre_{i}'), 'valor': p.get(f'param_valor_{i}')} for i in range(30)]
        parada = None
        if p.get('parada_motivo') and p.get('parada_minutos'):
            from produccion.models import MotivoParada
            parada = {'motivo': MotivoParada.objects.filter(pk=_entero(p['parada_motivo']), activo=True).first(),
                      'minutos': p['parada_minutos']}
        try:
            c, propuesta = srv.resolver(
                c, request.user, p.get('resuelto') == '1', p.get('causa'), p.get('accion'), p.get('resultado'),
                p.get('piezas'), _entero(p.get('minutos')) or 0,
                defecto=Defecto.objects.filter(pk=_entero(p.get('defecto'))).first(), parametros=filas,
                fotos=request.FILES.getlist('fotos'), parada=parada, abrir_nc=p.get('abrir_nc') == '1',
                avisar_mantenimiento=p.get('avisar') == '1', prioridad=p.get('prioridad') or 'NORMAL')
            partes = [f'Caso {c.numero} {c.get_estado_display().lower()}.']
            if c.merma_registrada:
                partes.append(f'Merma de {c.piezas_perdidas:,.2f} a la orden {c.orden.numero} '
                              f'(S/ {c.costo:,.2f}).')
            if c.no_conformidad_id:
                partes.append(f'No conformidad {c.no_conformidad.codigo}.')
            if c.aviso_id:
                partes.append(f'Aviso de mantenimiento {c.aviso.numero}.')
            if propuesta is not None:
                partes.append(f'Los parámetros salieron de la ficha: se propuso la revisión {propuesta} para '
                              f'aprobación.')
            messages.success(request, ' '.join(partes))
        except srv.ErrorAsistente as exc:
            messages.error(request, f'No se cerró el caso: {exc}')
        return redirect('asistente:caso', pk)
    from produccion.models import MotivoParada
    setpoints = list(c.ficha.setpoints.all()) if c.ficha_id else []
    return render(request, 'asistente/caso.html', {
        'c': c, 'consultas': c.consultas.order_by('creado', 'id'), 'setpoints': setpoints,
        'filas_param': setpoints + [None] * max(3, 8 - len(setpoints)),
        'defectos': Defecto.objects.filter(activo=True, **({'proceso': c.proceso} if c.proceso_id else {})),
        'motivos': MotivoParada.para_puesto(c.equipo.centro if c.equipo_id else None),
        'prioridades': [('NORMAL', 'Normal'), ('ALTA', 'Alta'), ('EMERGENCIA', 'Emergencia')],
        'propuestas': c.fichas_propuestas.all()})


# ---------------------------------------------------------------- fichas de parámetros
@login_required
def fichas(request):
    qs = FichaParametros.objects.select_related('producto', 'equipo', 'molde', 'proceso')
    estado = request.GET.get('estado', '')
    if estado in dict(FichaParametros.ESTADOS):
        qs = qs.filter(estado=estado)
    else:
        qs = qs.exclude(estado='RETIRADA')
    if _entero(request.GET.get('proceso')):
        qs = qs.filter(proceso_id=_entero(request.GET['proceso']))
    q = (request.GET.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(numero__icontains=q) | Q(producto__nombre__icontains=q) | Q(producto__codigo__icontains=q) |
                       Q(equipo__nombre__icontains=q) | Q(equipo__codigo__icontains=q) | Q(molde__codigo__icontains=q))
    return render(request, 'asistente/fichas.html', {
        'fichas': qs[:300], 'estado': estado, 'q': q, 'estados': FichaParametros.ESTADOS,
        'procesos': Proceso.objects.all(), 'f_proceso': _entero(request.GET.get('proceso'))})


def _filas_setpoints(request, n=30):
    p = request.POST
    return [{k: p.get(f'sp_{k}_{i}') for k in ('nombre', 'unidad', 'minimo', 'objetivo', 'maximo', 'nota')}
            for i in range(n)]


def _ficha_form(request, ficha):
    form = FichaForm(request.POST or None, instance=ficha)
    filas = None
    if request.method == 'POST':
        filas = _filas_setpoints(request)
        if form.is_valid():
            try:
                f = srv.guardar_ficha(form.save(commit=False), filas, request.user)
                if request.POST.get('proponer'):
                    srv.proponer(f, request.user)
                    messages.success(request, f'Ficha {f} enviada a aprobación.')
                else:
                    messages.success(request, f'Ficha {f} guardada como borrador.')
                return redirect('asistente:ficha', f.pk)
            except srv.ErrorAsistente as exc:
                form.add_error(None, str(exc))
    if filas is None:
        if ficha.pk:
            filas = [{'nombre': s.nombre, 'unidad': s.unidad, 'minimo': s.minimo, 'objetivo': s.objetivo,
                      'maximo': s.maximo, 'nota': s.nota} for s in ficha.setpoints.all()]
        else:
            proceso = Proceso.objects.filter(pk=_entero(request.GET.get('proceso'))).first()
            filas = [{'nombre': p.nombre, 'unidad': p.unidad} for p in proceso.parametros.all()] if proceso else []
    filas = [f for f in filas if any(f.values())]
    filas += [{}] * max(2, 12 - len(filas))
    return render(request, 'asistente/ficha_form.html', {
        'form': form, 'ficha': ficha, 'filas': filas, 'procesos': Proceso.objects.filter(activo=True),
        'siguiente': len(filas)})


@login_required
def ficha_nueva(request):
    return _ficha_form(request, FichaParametros())


@login_required
def ficha_editar(request, pk):
    ficha = get_object_or_404(FichaParametros, pk=pk)
    if ficha.estado not in ('BORRADOR', 'RECHAZADA'):
        messages.error(request, 'Solo se edita una ficha en borrador o rechazada; cree una nueva revisión.')
        return redirect('asistente:ficha', pk)
    return _ficha_form(request, ficha)


@login_required
def ficha(request, pk):
    f = get_object_or_404(FichaParametros.objects.select_related('producto', 'equipo', 'molde', 'proceso',
                                                                 'propuesta_por', 'aprobada_por'), pk=pk)
    if request.method == 'POST':
        accion = request.POST.get('accion')
        try:
            if accion == 'proponer':
                srv.proponer(f, request.user)
                messages.success(request, f'Ficha {f} enviada a aprobación.')
            elif accion == 'aprobar':
                srv.aprobar(f, request.user)
                messages.success(request, f'Ficha {f} aprobada: es el estándar vigente.')
            elif accion == 'rechazar':
                srv.rechazar(f, request.user, request.POST.get('motivo'))
                messages.success(request, f'Ficha {f} rechazada.')
            elif accion == 'revision':
                nueva = srv.nueva_revision(f, request.user)
                messages.success(request, f'Revisión {nueva} creada como borrador: ajuste los valores y envíela.')
                return redirect('asistente:ficha_editar', nueva.pk)
        except srv.ErrorAsistente as exc:
            messages.error(request, str(exc))
        return redirect('asistente:ficha', pk)
    revisiones = FichaParametros.objects.filter(producto=f.producto, equipo=f.equipo, molde=f.molde,
                                                proceso=f.proceso).order_by('-revision')
    return render(request, 'asistente/ficha.html', {
        'f': f, 'setpoints': f.setpoints.all(), 'revisiones': revisiones,
        'casos': f.casos.select_related('defecto')[:20],
        'puede_aprobar_esta': f.estado == 'POR_APROBAR' and puede(request.user, 'asistente.aprobar_ficha') and (
            f.propuesta_por_id != request.user.pk)})


@login_required
def importar(request):
    if request.GET.get('plantilla'):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Fichas'
        ws.append(srv.COLUMNAS)
        ws.append(['PT-001', 'INY-01', 'MOL-01', 'INY', 'Temperatura de barril', '°C', 210, 220, 230,
                   'Manual del fabricante'])
        ws.append(['PT-001', 'INY-01', 'MOL-01', 'INY', 'Presión de mantenimiento', 'bar', 550, 575, 600, ''])
        for col, ancho in zip('ABCDEFGHIJ', (18, 16, 22, 16, 28, 9, 10, 10, 10, 30)):
            ws.column_dimensions[col].width = ancho
        ayuda = wb.create_sheet('Códigos')
        ayuda.append(['Procesos (código)', 'Nombre'])
        for p in Proceso.objects.filter(activo=True):
            ayuda.append([p.codigo, p.nombre])
        salida = BytesIO()
        wb.save(salida)
        r = HttpResponse(salida.getvalue(),
                         content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r['Content-Disposition'] = 'attachment; filename="plantilla_fichas_parametros.xlsx"'
        return r
    if request.method == 'POST' and request.FILES.get('archivo'):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(request.FILES['archivo'], read_only=True, data_only=True)
            filas = list(wb.worksheets[0].iter_rows(values_only=True))
            cabecera = [str(c or '').strip() for c in filas[0]] if filas else []
            datos = [(n, dict(zip(cabecera, fila))) for n, fila in enumerate(filas[1:], 2)
                     if any(v not in (None, '') for v in fila)]
            faltan = [c for c in srv.COLUMNAS if c not in cabecera]
            if faltan:
                raise srv.ErrorAsistente('Use la plantilla: faltan las columnas ' + ', '.join(faltan))
            creadas = srv.importar(datos, request.user)
            messages.success(request, f'{len(creadas)} ficha(s) importadas y enviadas a aprobación: '
                                      + ', '.join(str(f) for f in creadas[:10]))
            return redirect(f'{reverse("asistente:fichas")}?estado=POR_APROBAR')
        except srv.ErrorAsistente as exc:
            messages.error(request, f'No se importó nada: {exc}')
        except Exception:  # noqa: BLE001 (archivo dañado o no es Excel)
            messages.error(request, 'No se pudo leer el archivo: use la plantilla en formato .xlsx.')
    return render(request, 'asistente/importar.html', {'columnas': srv.COLUMNAS})


# ---------------------------------------------------------------- conocimiento
@login_required
def manuales(request):
    qs = Manual.objects.select_related('equipo', 'puesto', 'material', 'fabricante', 'proceso')
    if not request.GET.get('inactivos'):
        qs = qs.filter(activo=True)
    if request.GET.get('tipo') in dict(Manual.TIPOS):
        qs = qs.filter(tipo=request.GET['tipo'])
    q = (request.GET.get('q') or '').strip()
    if q:
        qs = qs.filter(Q(titulo__icontains=q) | Q(contenido__icontains=q) | Q(equipo__nombre__icontains=q) |
                       Q(material__nombre__icontains=q) | Q(fabricante__nombre__icontains=q))
    return render(request, 'asistente/manuales.html', {'manuales': qs[:300], 'q': q, 'tipos': Manual.TIPOS,
                                                       'tipo': request.GET.get('tipo', '')})


@login_required
def manual(request, pk):
    m = get_object_or_404(Manual.objects.select_related('equipo', 'puesto', 'material', 'fabricante', 'proceso',
                                                        'documento', 'validado_por'), pk=pk)
    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion == 'validar':  # LUC-07: otra persona con el permiso de aprobar (como las fichas)
            from core.segregacion import propio
            error = propio(request.user, m)
            if error:
                messages.error(request, error)
            elif not (m.secciones or m.contenido.strip()):
                messages.error(request, 'No hay qué validar: adjunte el archivo del manual o escriba los puntos clave.')
            else:
                m.validado, m.validado_por, m.validado_en = True, request.user, timezone.now()
                m.save(update_fields=['validado', 'validado_por', 'validado_en'])
                messages.success(request, f'«{m}» validado: Lucho ya lo cita.')
        elif accion == 'quitar_validacion':
            m.validado, m.validado_por, m.validado_en = False, None, None
            m.save(update_fields=['validado', 'validado_por', 'validado_en'])
            messages.success(request, f'«{m}»: Lucho deja de citarlo hasta que se valide de nuevo.')
        elif accion == 'leer':
            from .lectura import leer_manual
            messages.info(request, leer_manual(m) or 'El manual no tiene un PDF ni un TXT adjunto.')
        return redirect('asistente:manual', m.pk)
    return render(request, 'asistente/manual.html', {'m': m})


def _manual_form(request, m):
    form = ManualForm(request.POST or None, request.FILES or None, instance=m)
    contenido_antes = m.contenido
    if request.method == 'POST' and form.is_valid():
        nuevo = form.instance.pk is None
        obj = form.save(commit=False)
        if nuevo:
            obj.creado_por = request.user
        elif obj.contenido != contenido_antes and obj.validado:
            # LUC-07: lo validado era otro texto: vuelve a validarse antes de que Lucho lo cite
            obj.validado, obj.validado_por, obj.validado_en = False, None, None
            messages.warning(request, 'Cambió el contenido: el manual vuelve a «sin validar» hasta que otra persona '
                                      'lo valide.')
        obj.save()
        archivo = form.cleaned_data.get('archivo')
        if archivo:  # LUC-01: se adjunta y Lucho lo lee
            from core.sustentos import adjuntar
            messages.info(request, obj.al_adjuntar(adjuntar(obj, archivo, request.user, 'Archivo del manual')))
        messages.success(request, f'«{obj}» guardado.' + (' Adjunte el archivo del manual abajo.'
                                                          if nuevo and not archivo else ''))
        return redirect('asistente:manual', obj.pk)
    return render(request, 'asistente/formulario.html', {
        'form': form, 'titulo': 'Editar manual' if m.pk else 'Nuevo manual o conocimiento',
        'cancelar': reverse('asistente:manual', args=[m.pk]) if m.pk else reverse('asistente:manuales'),
        'anchos': ['contenido'],
        'ayuda': 'Amárrelo al equipo, al puesto, al material o al proceso: así el operario lo encuentra en su puesto '
                 'y Lucho lo cita. Solo se cita el conocimiento marcado como validado.'})


@login_required
def manual_nuevo(request):
    return _manual_form(request, Manual())


@login_required
def manual_editar(request, pk):
    return _manual_form(request, get_object_or_404(Manual, pk=pk))


# ---------------------------------------------------------------- configuración
@login_required
def procesos(request):
    form = ProcesoForm(request.POST or None) if request.POST.get('accion') != 'plantillas' else ProcesoForm()
    if request.method == 'POST':
        if request.POST.get('accion') == 'plantillas':
            n = plantillas.cargar()
            messages.success(request, f'{n} proceso(s) de ejemplo cargados con sus parámetros y defectos.' if n else
                             'Los procesos de ejemplo ya estaban cargados.')
            return redirect('asistente:procesos')
        if form.is_valid():
            p = form.save()
            messages.success(request, f'Proceso «{p}» creado: agregue sus parámetros y defectos.')
            return redirect('asistente:proceso', p.pk)
    return render(request, 'asistente/procesos.html', {
        'procesos': Proceso.objects.annotate(n_param=Count('parametros', distinct=True),
                                             n_def=Count('defectos', distinct=True),
                                             n_fichas=Count('fichas', distinct=True)),
        'form': form})


@login_required
def proceso(request, pk):
    p = get_object_or_404(Proceso, pk=pk)
    form = ProcesoForm(instance=p)
    form_param, form_def = ParametroForm(prefix='par'), DefectoForm(prefix='def')
    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion == 'proceso':
            form = ProcesoForm(request.POST, instance=p)
            if form.is_valid():
                form.save()
                messages.success(request, 'Proceso guardado.')
                return redirect('asistente:proceso', pk)
        elif accion == 'parametro':
            form_param = ParametroForm(request.POST, prefix='par', instance=ParametroProceso(proceso=p))
            if form_param.is_valid():
                if p.parametros.filter(nombre__iexact=form_param.cleaned_data['nombre']).exists():
                    form_param.add_error('nombre', 'Ese parámetro ya está en el proceso.')
                else:
                    form_param.save()
                    return redirect('asistente:proceso', pk)
        elif accion == 'defecto':
            form_def = DefectoForm(request.POST, prefix='def', instance=Defecto(proceso=p))
            if form_def.is_valid():
                if p.defectos.filter(codigo=form_def.cleaned_data['codigo']).exists():
                    form_def.add_error('codigo', 'Ese código ya está en el proceso.')
                else:
                    form_def.save()
                    return redirect('asistente:proceso', pk)
        elif accion == 'ajuste':
            d = p.defectos.filter(pk=_entero(request.POST.get('defecto'))).first()
            par = p.parametros.filter(pk=_entero(request.POST.get('parametro'))).first()
            if d and par and request.POST.get('sentido') in dict(AjusteDefecto.SENTIDOS):
                AjusteDefecto.objects.update_or_create(defecto=d, parametro=par, defaults={
                    'sentido': request.POST['sentido'], 'prioridad': _entero(request.POST.get('prioridad')) or 1,
                    'nota': (request.POST.get('nota') or '').strip()[:200]})
            else:
                messages.error(request, 'Elija el defecto, el parámetro y el sentido del ajuste.')
            return redirect(f'{reverse("asistente:proceso", args=[pk])}#defecto-{d.pk if d else ""}')
        elif accion == 'quitar_ajuste':
            AjusteDefecto.objects.filter(pk=_entero(request.POST.get('ajuste')), defecto__proceso=p).delete()
            return redirect('asistente:proceso', pk)
        elif accion == 'defecto_activo':
            d = get_object_or_404(Defecto, pk=_entero(request.POST.get('defecto')), proceso=p)
            d.activo = not d.activo
            d.save(update_fields=['activo'])
            return redirect('asistente:proceso', pk)
    return render(request, 'asistente/proceso.html', {
        'p': p, 'form': form, 'form_param': form_param, 'form_def': form_def,
        'parametros': p.parametros.all(), 'sentidos': AjusteDefecto.SENTIDOS,
        'defectos': p.defectos.prefetch_related('ajustes__parametro')})


@login_required
def configuracion(request):
    form = ConfigForm(request.POST or None, instance=ConfigAsistente.actual())
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Configuración del asistente guardada.')
        return redirect('asistente:configuracion')
    return render(request, 'asistente/formulario.html', {
        'form': form, 'titulo': 'Configuración del asistente', 'cancelar': reverse('asistente:chat'),
        'ayuda': 'Desde cuántas consultas fuera de tema en 7 días se marca el uso inadecuado (queda en la bitácora y '
                 'el supervisor lo ve en la campana y en «Uso inadecuado»).'})


@login_required
def uso(request):
    if request.method == 'POST' and request.POST.get('accion') == 'descartar':
        # LUC-02: falso positivo: sale del conteo y del aviso, pero queda en la bitácora con su motivo
        c = get_object_or_404(Consulta, pk=_entero(request.POST.get('id')))
        motivo = (request.POST.get('motivo') or '').strip()
        if len(motivo) < 5:
            messages.error(request, 'Indique por qué no aplica (mínimo 5 caracteres).')
        else:
            c.descartada, c.descartada_por, c.motivo_descarte = True, request.user, motivo[:200]
            c.supervisor_avisado = False
            c.save(update_fields=['descartada', 'descartada_por', 'motivo_descarte', 'supervisor_avisado'])
            from core.auditoria import registrar
            registrar('MODIFICAR', c, {'descartada': [False, True], 'motivo': [None, motivo[:200]]},
                      'Uso inadecuado descartado (falso positivo)')
            messages.success(request, 'Registro descartado como falso positivo: ya no cuenta para el aviso.')
        return redirect('asistente:uso')
    desde = timezone.now() - timedelta(days=90)
    qs = (Consulta.objects.filter(creado__gte=desde).exclude(motivo_uso='')
          .select_related('usuario', 'caso', 'orden', 'descartada_por'))
    resumen = (qs.filter(descartada=False).values('usuario__username', 'usuario__first_name', 'usuario__last_name')
               .annotate(n=Count('id'), marcadas=Count('id', filter=Q(supervisor_avisado=True))).order_by('-n'))
    return render(request, 'asistente/uso.html', {'consultas': qs[:200], 'resumen': resumen})
