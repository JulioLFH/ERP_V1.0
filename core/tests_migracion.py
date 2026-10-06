"""Migración desde los Excel del sistema anterior y carga masiva de trabajadores (v1.18)."""
import os
import tempfile
from datetime import date
from decimal import Decimal
from io import StringIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from contabilidad.models import CentroCosto, CuentaContable
from core.forms import digito_ruc
from planillas.models import Trabajador
from ventas.models import Venta

from .carga_masiva import VALIDADORES, cargar
from .models import Almacen, Producto, Serie, StockLote, Tercero

D = Decimal


def guardar(carpeta, nombre, filas):
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    wb.save(os.path.join(carpeta, nombre))


class MigracionTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        with patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450'))):
            call_command('seed_demo', verbosity=0)
        cls.admin = User.objects.create_superuser('mig', 'm@i.pe', 'x')

    def setUp(self):
        p = patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450')))
        p.start()
        self.addCleanup(p.stop)

    def _carpeta(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        c = tmp.name
        ruc = f'2060000001{digito_ruc("2060000001")}'
        self.ruc = ruc
        guardar(c, 'Data_Productos_API.xlsx', [
            ['codigo_producto', 'nombre_producto', 'codigo_de_barras', 'peso_kg', 'categoria_producto',
             'marca_producto', 'unidad_de_producto', 'unidad_de_compra', 'tipo_producto', 'trazabilidad_seguimiento',
             'puede_vender', 'puede_comprar', 'descripcion_compra', 'descripcion_venta'],
            ['X9001', 'YOGURT FRESA 1L', '775001', 1.05, 'PT YOGURTS', 'MIO', 'Unidades', 'Unidades',
             'Producto Almacenable', 'por lotes', 'SI', 'NO', None, None],
            ['X9002', 'AZUCAR', None, 0, 'MATERIA PRIMA', '-', 'kg', 'kg', 'Producto Almacenable', 'sin seguimiento',
             'NO', 'SI', None, None],
            ['X9003', 'FLETE', None, 0, 'SERVICIOS', None, 'Unidades', 'Unidades', 'Servicio', 'sin seguimiento',
             'NO', 'SI', None, None],
            [None, 'SIN CODIGO', None, 0, 'All', None, 'Unidades', 'Unidades', 'Consumible', 'sin seguimiento', 'NO',
             'SI', None, None]])
        guardar(c, 'Data_Contactos_API.xlsx', [
            ['nombre', 'nombre_completo', 'tipo_contacto', 'es_empresa', 'es_cliente', 'es_empleado', 'es_proveedor',
             'codigo_postal', 'direccion_completa', 'calle', 'email', 'telefono', 'movil', 'tipo_documento',
             'numero_documento', 'equipo_de_venta', 'lista_de_precios', 'termino_de_pago'],
            ['ANDINA', 'ANDINA SAC', 'Contacto', 'SI', 'SI', 'NO', 'SI', '150103', 'AV. LIMA 1', None, 'a@b.pe',
             '01-555', None, 'RUC', ruc, 'Canal Moderno', 'Mayorista Lima (PEN)', '30 días'],
            ['ANDINA', 'ANDINA SAC', 'Direccion de Entrega', 'NO', 'SI', 'NO', 'NO', None, 'ALMACEN', None, None,
             None, None, 'RUC', ruc, None, None, None],
            ['PEREZ', 'PEREZ LOPEZ JUAN CARLOS', 'Contacto', 'NO', 'NO', 'SI', 'NO', None, None, None, None, None,
             None, 'DNI', '41112222', None, None, None]])
        guardar(c, 'Data_Saldo(Valorado)_Scraping.xlsx', [
            ['IDOdoo', 'Categoría de Producto N1', 'Marca', 'Codigo Producto', 'Producto', 'Unidad', 'Precio De Venta',
             'Saldo Cantidad', 'Saldo Soles', 'C/U', 'Almacen'],
            [1, 'PT YOGURTS', 'MIO', 'X9001', 'YOGURT FRESA 1L', 'Unidades', 6.5, 100, 350, 3.5, 'LPTER/Existencias'],
            [2, 'MATERIA PRIMA', '-', 'X9002', 'AZUCAR', 'kg', 0, 50.0000000001, 150, 3, 'LPROD/Existencias'],
            [2, 'MATERIA PRIMA', '-', 'X9002', 'AZUCAR', 'kg', 0, -4.5e-13, 0, 0, 'LPTER/Existencias'],
            [2, 'MATERIA PRIMA', '-', 'X9002', 'AZUCAR', 'kg', 0, -3, -9, 3, 'LPROD/TRANSITO']])
        guardar(c, 'Data_Saldo(Valorado)xlote_Scraping.xlsx', [
            ['N.Producto', 'Lote', 'Cod. Producto', 'Unidad', 'Categoria 1', 'Categoria 2', 'Categoria 3', 'N.Almacén',
             'Stock', 'Reservado', 'Disponible'],
            ['YOGURT', 'LT 901, FV 06112026', 'X9001', 'Unidades', None, None, None, 'LPTER/Existencias', 60, 0, 60]])
        guardar(c, 'Data_Kardex_API_2025.xlsx', [['sede', 'almacen'], ['LPTER/Existencias', 'Lácteos PT'],
                                                 ['LPROD/Existencias', 'Lácteos Producción']])
        guardar(c, 'Data_Contable_2026.xlsx', [
            ['CtaCont', 'DescCtaCont', 'Centro_costo'], ['70211100', 'Ventas productos', 'SIN_CC'],
            ['62110001', 'Sueldos', '[924002] PLANTA LACTEOS'], ['63110001', 'Fletes', '[951003] CANAL MODERNO']])
        guardar(c, 'Data_Cuenta_Planillas.xlsx', [['CTA', 'CECO', 'CATEGORIA', 'Subcategoria'],
                                                  ['92322200', '922009', 'MANTENIMIENTO', 'Gasto de personal']])
        guardar(c, 'Data_FacturacionVentas_API_2025_2026.xlsx', [
            ['nro_documento_completo', 'tipo_de_documento', 'fecha_de_factura', 'fecha_vencimiento', 'estado',
             'estado_pago', 'contacto_factura', 'tipo_documento', 'numero_documento', 'moneda', 'tipo_de_cambio',
             'monto_total_factura'],
            ['FX01-00000120', 'Factura', '2026-09-01 00:00:00', '2026-10-01 00:00:00', 'Publicado', 'Sin Pagar',
             'ANDINA SAC', 'RUC', ruc, 'PEN', 3.5, 1180],
            ['FX01-00000120', 'Factura', '2026-09-01 00:00:00', '2026-10-01 00:00:00', 'Publicado', 'Sin Pagar',
             'ANDINA SAC', 'RUC', ruc, 'PEN', 3.5, 1180],  # otra línea del mismo comprobante
            ['FX01-00000121', 'Factura', '2026-09-02 00:00:00', None, 'Publicado', 'Pagado', 'ANDINA SAC', 'RUC',
             ruc, 'PEN', 3.5, 500],
            ['BX01-00000009', 'Boleta', '2026-09-03 00:00:00', None, 'Publicado', 'Sin Pagar', 'CLIENTE VARIOS',
             'DNI', '40001111', 'PEN', 3.5, 59],
            ['FC09-00000003', 'Nota de crédito', '2026-09-04 00:00:00', None, 'Publicado', 'Sin Pagar', 'ANDINA SAC',
             'RUC', ruc, 'PEN', 3.5, 20],
            ['FX01-00000122', 'Factura', '2026-09-05 00:00:00', None, 'Publicado', 'Pagado Parcialmente',
             'ANDINA SAC', 'RUC', ruc, 'PEN', 3.5, 300]])
        return c

    def test_migracion_completa(self):
        carpeta = self._carpeta()
        reporte = os.path.join(carpeta, 'obs.xlsx')
        salida = StringIO()
        call_command('migrar_excels', carpeta, '--corte', '2026-09-27', '--reporte', reporte, stdout=salida)
        # maestros
        yog, azucar, flete = (Producto.objects.get(codigo=c) for c in ('X9001', 'X9002', 'X9003'))
        self.assertEqual((yog.clase, yog.control, yog.precio_venta, yog.puede_comprarse),
                         ('PRODUCTO_TERMINADO', 'LOTE', D('6.5'), False))
        self.assertEqual((azucar.clase, azucar.unidad, flete.clase), ('MATERIA_PRIMA', 'KGM', 'SERVICIO'))
        t = Tercero.objects.get(numero_doc=self.ruc)
        self.assertEqual((t.tipo, t.dias_credito, t.ubigeo, t.lista_precios.nombre),
                         ('AMBOS', 30, '150103', 'Mayorista Lima (PEN)'))
        self.assertEqual(CuentaContable.objects.get(codigo='70211100').naturaleza, 'ACREEDORA')
        self.assertEqual(CentroCosto.objects.get(codigo='924002').tipo, 'PRODUCCION')
        self.assertEqual(CentroCosto.objects.get(codigo='951003').tipo, 'VENTAS')
        self.assertEqual(Almacen.objects.get(codigo='LPTER').nombre, 'Lácteos PT')
        # stock: lote con coma normalizado, el resto al lote INICIAL; residuo -4.5e-13 ignorado; negativo reportado
        self.assertEqual((yog.__class__.objects.get(pk=yog.pk).stock, Producto.objects.get(pk=azucar.pk).stock),
                         (D('100'), D('50')))
        lotes = dict(StockLote.objects.filter(lote__producto=yog).values_list('lote__codigo', 'cantidad'))
        self.assertEqual(lotes, {'LT 901 FV 06112026': D('60'), 'INICIAL': D('40')})
        self.assertEqual(StockLote.objects.get(lote__codigo='LT 901 FV 06112026').lote.vencimiento,
                         date(2026, 11, 6))
        # por cobrar: solo los "Sin Pagar" (factura y boleta); cliente de la boleta creado
        docs = Venta.objects.filter(es_saldo_inicial=True).order_by('serie')
        self.assertEqual([(v.serie, v.numero, v.total) for v in docs],
                         [('BX01', '9', D('59.00')), ('FX01', '120', D('1180.00'))])
        self.assertEqual(Serie.objects.get(tipo='01', serie='FX01').correlativo, 122)
        self.assertEqual(Serie.objects.get(tipo='07', serie='FC09').correlativo, 3)
        # trabajador por completar
        tr = Trabajador.objects.get(numero_doc='41112222')
        self.assertEqual((tr.apellido_paterno, tr.apellido_materno, tr.nombres, tr.datos_completos),
                         ('PEREZ', 'LOPEZ', 'JUAN CARLOS', False))
        texto = salida.getvalue()
        self.assertIn('Por cobrar: 2 comprobantes', texto)
        self.assertTrue(os.path.exists(reporte))
        from openpyxl import load_workbook
        hojas = load_workbook(reporte).sheetnames
        self.assertIn('Stock no cargado', hojas)
        self.assertIn('Por cobrar no cargado', hojas)

    def test_simulacion_no_graba(self):
        carpeta = self._carpeta()
        call_command('migrar_excels', carpeta, '--reporte', os.path.join(carpeta, 'o.xlsx'), '--simular',
                     stdout=StringIO())
        self.assertFalse(Producto.objects.filter(codigo='X9001').exists())

    def test_carga_masiva_trabajadores(self):
        CentroCosto.objects.create(codigo='924001', nombre='Planta', tipo='PRODUCCION')
        fila = {'numero_doc': 4567891, 'apellido_paterno': 'QUISPE', 'nombres': 'ANA', 'fecha_ingreso': '01/03/2024',
                'sueldo': 1500, 'sistema_pensiones': 'AFP', 'afp': 'integra', 'asignacion_familiar': 'SI',
                'centro_costo': '924001', 'regimen': 'Pequeña', 'quinta_remuneracion_previa': '12000',
                'quinta_retencion_previa': 80}
        r = VALIDADORES['trabajadores'](fila, False)
        self.assertEqual(r['codigo'], '04567891')  # el DNI que Excel guardó como número recupera el cero
        cargar('trabajadores', [r], self.admin)
        t = Trabajador.objects.get(numero_doc='04567891')
        self.assertEqual((t.afp.codigo, t.regimen, t.asignacion_familiar, t.centro_costo.codigo, t.datos_completos),
                         ('INTEGRA', 'PEQUENA', True, '924001', True))
        self.assertEqual((t.quinta_anio, t.quinta_remuneracion_previa, t.quinta_retencion_previa),
                         (date.today().year, D('12000'), D('80')))
        from .carga_masiva import ErrorFila
        with self.assertRaises(ErrorFila):
            VALIDADORES['trabajadores']({**fila, 'afp': 'NOEXISTE'}, True)
        with self.assertRaises(ErrorFila):
            VALIDADORES['trabajadores'](fila, False)  # ya existe
