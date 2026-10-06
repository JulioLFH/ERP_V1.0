"""Historial de ventas del sistema anterior (Odoo), para consulta y reportes.

    python manage.py importar_historial_ventas "<carpeta>" [--reporte historial.xlsx] [--simular]

Lee Data_FacturacionVentas_API_2025_2026.xlsx (una fila por línea) y crea cada comprobante con su detalle, marcado
como histórico: aparece en los comprobantes y en los reportes de ventas, pero no va al registro de ventas/PLE ni a
SUNAT (ya se declaró con el sistema anterior), no mueve stock, no genera asientos y no se edita.
Los comprobantes sin pagar ya cargados como saldo inicial reciben su detalle real (siguen en cuentas por cobrar).
Los cobrados quedan con saldo cero. Es idempotente: lo ya importado se omite.
"""
import os
import time
from collections import OrderedDict, defaultdict
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from .migrar_excels import _dec, _fecha, _txt

D0 = Decimal('0')
TIPOS = {'Factura': '01', 'Boleta': '03', 'Nota de crédito': '07', 'Nota de Débito': '08'}
TIPOS_DOC = {'RUC': '6', 'DNI': '1', 'Cédula Extranjera': '4'}
COLUMNAS = ['nro_documento_completo', 'tipo_de_documento', 'fecha_de_factura', 'fecha_vencimiento', 'estado',
            'estado_pago', 'contacto_factura', 'tipo_documento', 'numero_documento', 'vendedor', 'tipo_de_cambio',
            'moneda', 'monto_total_factura', 'factura_relacionada', 'codigo_producto', 'producto', 'cantidad',
            'impuestos', 'precio_unitario', 'descuento', 'monto_sin_igv_linea', 'nro_orden_pedido',
            'termino_de_pago']
TOLERANCIA = Decimal('0.10')


class _Simulacion(Exception):
    pass


class Command(BaseCommand):
    help = 'Importa el historial de facturación del sistema anterior para consulta y reportes.'

    def add_arguments(self, parser):
        parser.add_argument('carpeta')
        parser.add_argument('--reporte', default='historial_ventas.xlsx')
        parser.add_argument('--simular', action='store_true')
        parser.add_argument('--rehacer', action='store_true',
                            help='Borra el historial ya importado (cancelado) y lo vuelve a importar')

    def handle(self, carpeta, reporte, simular, rehacer, **_):
        self.rehacer = rehacer
        from openpyxl import load_workbook
        ruta = os.path.join(carpeta, 'Data_FacturacionVentas_API_2025_2026.xlsx')
        if not os.path.exists(ruta):
            raise CommandError(f'No existe {ruta}')
        inicio = time.time()
        wb = load_workbook(ruta, read_only=True, data_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        enc = [_txt(c) for c in next(it)]
        faltan = [c for c in COLUMNAS if c not in enc]
        if faltan:
            raise CommandError(f'Faltan las columnas {faltan}')
        ix = {c: enc.index(c) for c in COLUMNAS}
        docs = OrderedDict()
        for f in it:
            r = {c: f[i] if i < len(f) else None for c, i in ix.items()}
            tipo = TIPOS.get(_txt(r['tipo_de_documento']))
            numero = _txt(r['nro_documento_completo'])
            if not tipo or '-' not in numero:
                continue
            docs.setdefault((tipo, numero), []).append(r)
        wb.close()
        self.stdout.write(f'[{time.time() - inicio:5.0f}s] {len(docs)} comprobantes en el archivo')
        self.obs = []
        resumen = defaultdict(int)
        try:
            with transaction.atomic():
                self._importar(docs, resumen, inicio)
                if simular:
                    raise _Simulacion()
        except _Simulacion:
            self.stdout.write('Simulación: no se grabó nada.')
        self._reporte(reporte, resumen)
        for k, v in resumen.items():
            self.stdout.write(f'  {k}: {v}')
        self.stdout.write(self.style.SUCCESS(f'Observaciones: {os.path.abspath(reporte)}'))

    def _importar(self, docs, resumen, inicio):
        from core.models import Producto, Tercero
        from ventas.models import Venta, VentaItem
        if self.rehacer:
            borrados = Venta.objects.filter(es_historico=True, es_saldo_inicial=False).delete()[1].get(
                'ventas.Venta', 0)
            resumen['Borrados para rehacer'] = borrados
        productos = {p.codigo: p.pk for p in Producto.objects.only('pk', 'codigo')}
        terceros = dict(Tercero.objects.values_list('numero_doc', 'pk'))
        existentes = {(t, s, n): (pk, si, hist) for t, s, n, pk, si, hist in Venta.objects.values_list(
            'tipo_comprobante', 'serie', 'numero', 'pk', 'es_saldo_inicial', 'es_historico')}
        for n, ((tipo, completo), lineas) in enumerate(docs.items(), 1):
            if n % 2000 == 0:
                self.stdout.write(f'[{time.time() - inicio:5.0f}s] {n} procesados')
            serie, _, correlativo = completo.partition('-')
            serie, numero = serie.upper(), correlativo.lstrip('0') or '0'
            r = lineas[0]
            previo = existentes.get((tipo, serie, numero))
            if previo and previo[2] and not (self.rehacer and previo[1]):
                resumen['Ya importados (omitidos)'] += 1
                continue
            if previo and not previo[1]:
                self.obs.append((tipo, completo, 'Ya existe un comprobante con ese número registrado en Ceiba: '
                                                 'no se importó'))
                resumen['Omitidos por número ya usado en Ceiba'] += 1
                continue
            doc_num = _txt(r['numero_documento'])
            if doc_num not in terceros:
                tipo_doc = TIPOS_DOC.get(_txt(r['tipo_documento']), '0')
                t = Tercero.objects.create(tipo='CLIENTE', tipo_doc=tipo_doc, numero_doc=doc_num[:15] or '-',
                                           nombre=_txt(r['contacto_factura'])[:200] or doc_num or 'CLIENTE VARIOS')
                terceros[t.numero_doc] = t.pk
                resumen['Clientes creados'] += 1
            items = self._items(lineas, productos)
            todo_gratuito = all(i['gratuita'] for i in items)
            for i in items:
                if i['gratuita'] and not todo_gratuito:  # bonificación dentro de una venta: no se cobra
                    i['descripcion'] = f'{i["descripcion"]} (transferencia gratuita, valor ref. {i["precio_ref"]})'[:250]
                    i['precio_unitario'] = D0
            # el sistema anterior suma las bonificaciones al total (y luego las revierte); aquí no se cobran
            gratuito_odoo = sum((i['neto_odoo'] for i in items if i['gratuita']), D0) * Decimal('1.18')
            emision = _fecha(r['fecha_de_factura'])
            vence = _fecha(r['fecha_vencimiento']) or emision
            moneda = 'USD' if _txt(r['moneda']) == 'USD' else 'PEN'
            total_odoo = _dec(r['monto_total_factura'])
            glosa = 'Historial del sistema anterior'
            if _txt(r['factura_relacionada']):
                glosa += f' · modifica {_txt(r["factura_relacionada"])}'
            if _txt(r['nro_orden_pedido']):
                glosa += f' · pedido {_txt(r["nro_orden_pedido"])}'
            if previo:  # saldo inicial sin pagar: se le pone el detalle real si cuadra con su total
                v = Venta.objects.get(pk=previo[0])
                v.es_historico = True
                v.vendedor = _txt(r['vendedor'])[:80]
                v.save(update_fields=['es_historico', 'vendedor'])
                if self._cuadra(items, v):
                    v.items.all().delete()
                    self._crear_items(v, items, VentaItem)
                    v.tipo_operacion = 'GRAVADA'
                    v.calcular_totales()
                    diferencia = total_odoo - v.total
                    if abs(diferencia) > TOLERANCIA:
                        raise CommandError(f'{completo}: el detalle no cuadra con el saldo inicial')
                    # el saldo por cobrar debe quedar exacto: el céntimo de redondeo va al IGV
                    v.igv += diferencia
                    v.total += diferencia
                    v.save()
                    resumen['Saldos iniciales completados con su detalle'] += 1
                else:
                    self.obs.append((tipo, completo, 'Saldo inicial: el detalle del archivo no cuadra con el total; '
                                                     'se dejó la línea única'))
                    resumen['Saldos iniciales sin detalle (no cuadra)'] += 1
                continue
            v = Venta(tipo_comprobante=tipo, serie=serie, numero=numero, tercero_id=terceros[doc_num],
                      fecha_emision=emision, fecha_vencimiento=max(vence, emision), moneda=moneda,
                      tipo_cambio=_dec(r['tipo_de_cambio']) if moneda == 'USD' else Decimal('1'),
                      tipo_operacion='GRATUITA' if todo_gratuito else 'GRAVADA', forma_pago='CREDITO' if vence > emision else 'CONTADO',
                      vendedor=_txt(r['vendedor'])[:80], descontar_stock=False, es_historico=True,
                      estado='ANULADO' if _txt(r['estado']) == 'Cancelado' else 'REGISTRADO',
                      estado_sunat='ACEPTADO', sunat_descripcion='Emitido con el sistema anterior', glosa=glosa)
            v.save()
            self._crear_items(v, items, VentaItem)
            v.calcular_totales()
            v.save()
            resumen['Comprobantes importados'] += 1
            if v.estado == 'REGISTRADO' and _txt(r['estado_pago']) == 'Pagado Parcialmente':
                self.obs.append((tipo, completo, 'Pagado parcialmente en el sistema anterior: el saldo pendiente no '
                                                 'viene en el archivo; cárguelo como saldo inicial por cobrar'))
                resumen['Pagados parcialmente (saldo por cargar)'] += 1
            if abs(v.total - (total_odoo - gratuito_odoo)) > TOLERANCIA:
                self.obs.append((tipo, completo, f'Total en Ceiba {v.total} vs sistema anterior {total_odoo}'))
                resumen['Con diferencia de total (ver observaciones)'] += 1

    def _items(self, lineas, productos):
        items = []
        for l in lineas:
            cantidad = _dec(l['cantidad']) or Decimal('1')
            precio = _dec(l['precio_unitario'])
            impuesto = _txt(l['impuestos'])
            descripcion = (_txt(l['producto']) or 'Sin descripción')[:250]
            if impuesto == 'IGV-INC-V':
                precio = (precio / Decimal('1.18')).quantize(Decimal('0.0001'))
            gratuita = impuesto in ('IGV-TRG', 'TRF-GRAT')
            items.append({'gratuita': gratuita, 'precio_ref': precio.quantize(Decimal('0.0001')),'producto_id': productos.get(_txt(l['codigo_producto']).upper()),
                          'descripcion': descripcion, 'cantidad': cantidad,
                          'precio_unitario': precio.quantize(Decimal('0.0001')),
                          'descuento_pct': _dec(l['descuento']), 'neto_odoo': _dec(l['monto_sin_igv_linea'])})
        return items

    def _cuadra(self, items, venta):
        base = sum((i['neto_odoo'] for i in items if not i['gratuita']), D0)
        return abs(base * Decimal('1.18') - venta.total) <= TOLERANCIA

    @staticmethod
    def _crear_items(venta, items, VentaItem):
        objetos = []
        for i in items:
            item = VentaItem(documento=venta, producto_id=i['producto_id'], descripcion=i['descripcion'],
                             cantidad=i['cantidad'], precio_unitario=i['precio_unitario'],
                             descuento_pct=i['descuento_pct'])
            item.subtotal = (item.cantidad * item.precio_neto).quantize(Decimal('0.01'))
            objetos.append(item)
        VentaItem.objects.bulk_create(objetos)

    def _reporte(self, ruta, resumen):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'Resumen'
        for k, v in resumen.items():
            ws.append([k, v])
        h = wb.create_sheet('Observaciones')
        h.append(['Tipo', 'Comprobante', 'Observación'])
        for o in self.obs:
            h.append(list(o))
        wb.save(ruta)
