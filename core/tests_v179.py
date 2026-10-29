"""Checklist v1.78.0 (24/10 13h50): inventario físico (INV-17 a INV-24), cruce de tres vías y recepciones (ALM-06 a
ALM-13), listas con columnas y Excel (LIS-01), aprobaciones (APR-02 a APR-06), movimientos de personal (PLA-01) y
usuario nuevo con mínimo privilegio (SEG-05)."""
from datetime import timedelta
from decimal import Decimal as D
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from compras.models import Compra, CompraItem, OrdenCompra, OrdenCompraItem
from inventario import fisico, servicios
from inventario.models import ConteoCiclico, Operacion, TipoOperacion

from . import aprobaciones as motor
from .models import Almacen, PerfilUsuario, PreferenciaLista, Producto, ReglaAprobacion, Tercero

HOY = timezone.localdate()


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_superuser('admin', 'a@a.com', 'x')
        cls.almacen = Almacen.objects.create(codigo='PRI', nombre='Almacén principal', es_principal=True)
        cls.prov = Tercero.objects.create(tipo='PROVEEDOR', tipo_doc='6', numero_doc='20100070970',
                                          nombre='PROVEEDOR DEMO')
        cls.harina = Producto.objects.create(codigo='HAR', nombre='Harina', clase='MERCADERIA', unidad='KGM')

    def setUp(self):
        self.client.force_login(self.admin)
        p = patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450')))
        p.start()
        self.addCleanup(p.stop)

    def oc(self, cantidad='100', precio='5'):
        oc = OrdenCompra.objects.create(numero=f'OC01-{OrdenCompra.objects.count() + 20:08d}', tercero=self.prov,
                                        estado='APROBADO')
        OrdenCompraItem.objects.create(documento=oc, producto=self.harina, descripcion='Harina',
                                       cantidad=D(cantidad), precio_unitario=D(precio))
        oc.calcular_totales()
        oc.save()
        return oc

    def recibir(self, oc, cantidad, costo=None, referencia='', fecha=None):
        op = Operacion.objects.create(tipo=TipoOperacion.objects.get(codigo='REC_COMPRA'), fecha=fecha or HOY,
                                      orden_compra=oc, almacen_destino=self.almacen, referencia=referencia)
        op.items.create(producto=self.harina, cantidad=D(cantidad), costo_unitario=D(costo) if costo else None)
        return op


class InventarioFisicoTest(Base):
    def setUp(self):
        super().setUp()
        self.harina.mover_stock(D('8'), 'Saldo', costo=D('10'), almacen=self.almacen)

    def test_lectura_estricta_de_cantidades(self):
        """INV-17 a INV-19: coma decimal, sin notación científica, negativos, letras ni números imposibles."""
        self.assertEqual(fisico.leer_cantidad('10,5'), D('10.5'))
        self.assertEqual(fisico.leer_cantidad(' 7 '), D('7'))
        self.assertIsNone(fisico.leer_cantidad(''))
        for malo in ('1e3', '-5', 'abc', '99999999999999', '1.500', '1,500', '10,5,1'):
            with self.assertRaises(fisico.ErrorFisico, msg=malo):
                fisico.leer_cantidad(malo)
        doc = fisico.crear(self.almacen, 'GENERAL', self.admin)
        item = doc.items.get()
        r = self.client.post(reverse('inventario:fisico', args=[doc.pk]), {'accion': 'contar', f'c_{item.pk}': 'abc'},
                             follow=True)
        self.assertContains(r, 'no se guardó nada')
        item.refresh_from_db()
        self.assertIsNone(item.contado)
        self.client.post(reverse('inventario:fisico', args=[doc.pk]), {'accion': 'contar', f'c_{item.pk}': '7,5'})
        item.refresh_from_db()
        self.assertEqual(item.contado, D('7.50'))

    def test_conteos_superpuestos_ciego_y_reversion(self):
        doc = fisico.crear(self.almacen, 'GENERAL', self.admin)
        with self.assertRaisesMessage(fisico.ErrorFisico, doc.numero):  # INV-20
            fisico.crear(self.almacen, 'GENERAL', self.admin)
        r = self.client.get(reverse('inventario:fisico', args=[doc.pk]))
        self.assertContains(r, 'Conteo a ciegas')  # INV-21: mientras se cuenta nadie ve la foto
        self.assertNotContains(r, 'Sistema (foto)')
        item = doc.items.get()
        fisico.registrar(doc, {item.pk: D('8')}, self.admin)
        doc.refresh_from_db()
        self.assertEqual(doc.estado, 'APROBADO')  # INV-22: sin diferencias, listo para contabilizar
        fisico.contabilizar(doc, self.admin)
        doc2 = fisico.crear(self.almacen, 'GENERAL', self.admin)
        fisico.registrar(doc2, {doc2.items.get().pk: D('7')}, self.admin)
        doc2.refresh_from_db()
        self.assertEqual(doc2.estado, 'CONTADO')
        fisico.registrar(doc2, {doc2.items.get().pk: D('7')}, self.admin, reconteo=True)  # fuera de la tolerancia
        doc2.items.update(motivo='MERMA')
        fisico.aprobar(doc2, self.admin)
        fisico.contabilizar(doc2, self.admin)
        doc2.refresh_from_db()
        self.assertEqual(self.harina.stock_en(self.almacen), D('7'))
        ajuste = doc2.ajuste_salida
        self.assertIn(doc2.numero, ' '.join(servicios.errores_anulacion(ajuste)))  # INV-23: no se anula suelto
        fisico.revertir(doc2, self.admin, 'Recuento equivocado')
        doc2.refresh_from_db()
        self.assertEqual((doc2.estado, self.harina.stock_en(self.almacen)), ('APROBADO', D('8')))
        import io

        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(fisico.acta(doc2).content)).active  # INV-24: totales y firmas
        textos = ' '.join(str(v) for fila in hoja.iter_rows(values_only=True) for v in fila if v)
        self.assertIn('Faltantes S/', textos)
        self.assertIn('Firma', textos)
        ConteoCiclico.objects.filter(pk=doc2.pk).update(estado='CONTADO', fecha=HOY - timedelta(days=10))
        self.assertIn(doc2.pk, [d.pk for d in fisico.abiertos_antiguos()])


class CruceRecepcionTest(Base):
    def test_cruce_de_tres_vias(self):
        """ALM-06: no se factura más de lo recibido; precio fuera de tolerancia queda retenido para pago."""
        from compras.cruce import errores_cantidad, retener_si_corresponde
        oc = self.oc()
        servicios.confirmar(self.recibir(oc, '50'), self.admin)
        self.assertTrue(errores_cantidad(oc, [(self.harina, D('100'))], ingresa=False))
        self.assertEqual(errores_cantidad(oc, [(self.harina, D('50'))], ingresa=False), [])
        self.assertEqual(errores_cantidad(oc, [(self.harina, D('100'))], ingresa=True), [])  # la factura ingresa
        c = Compra.objects.create(tipo_comprobante='01', serie='F001', numero='2002', tercero=self.prov,
                                  fecha_emision=HOY, orden_compra=oc, ingresar_almacen=False)
        CompraItem.objects.create(documento=c, producto=self.harina, descripcion='Harina', cantidad=D('50'),
                                  precio_unitario=D('6.50'))
        c.calcular_totales()
        c.save()
        self.assertTrue(errores_cantidad(oc, [(self.harina, D('5'))], ingresa=False))  # ya se facturó lo recibido
        self.assertTrue(retener_si_corresponde(c))
        c.refresh_from_db()
        self.assertTrue(c.pago_bloqueado)

    def test_recepcion_costo_de_la_oc_fecha_guia_y_anulacion(self):
        oc = self.oc()
        op = self.recibir(oc, '30', costo='9', referencia='Guía T001-905')
        servicios.confirmar(op, self.admin)
        op.items.get().refresh_from_db()
        self.assertEqual(op.items.get().costo_unitario, D('5.0000'))  # ALM-07: el costo es el de la orden
        futura = self.recibir(oc, '1', fecha=HOY + timedelta(days=30))
        self.assertTrue(any('futura' in e for e in servicios.errores_confirmacion(futura)))  # ALM-08
        futura.delete()
        # ALM-09: anular una recepción facturada devuelve la orden a pendiente y retiene la factura
        OrdenCompra.objects.filter(pk=oc.pk).update(estado='ATENDIDO')
        c = Compra.objects.create(tipo_comprobante='01', serie='F001', numero='2001', tercero=self.prov,
                                  fecha_emision=HOY, orden_compra=oc, ingresar_almacen=False)
        CompraItem.objects.create(documento=c, producto=self.harina, descripcion='Harina', cantidad=D('30'),
                                  precio_unitario=D('5'))
        c.calcular_totales()
        c.save()
        servicios.anular(op, self.admin, 'Recepción duplicada')
        oc.refresh_from_db()
        c.refresh_from_db()
        self.assertEqual(oc.estado, 'APROBADO')
        self.assertTrue(c.pago_bloqueado)
        self.assertIn('Con diferencia', c.motivo_bloqueo_pago)

    def test_borrador_no_queda_si_no_se_confirma_y_guia_repetida(self):
        """ALM-10 y ALM-11: la guía repetida avisa; si la confirmación falla no queda un borrador suelto."""
        oc = self.oc()
        servicios.confirmar(self.recibir(oc, '10', referencia='Guía T001-905'), self.admin)
        datos = {'fecha': HOY.isoformat(), 'orden_compra': oc.pk, 'almacen_destino': self.almacen.pk,
                 'referencia': 'Guía T001-905', 'glosa': '', 'items-TOTAL_FORMS': '1', 'items-INITIAL_FORMS': '0',
                 'items-MIN_NUM_FORMS': '1', 'items-MAX_NUM_FORMS': '1000', 'items-0-producto': self.harina.pk,
                 'items-0-cantidad': '5', 'accion': 'confirmar'}
        r = self.client.post(reverse('inventario:nueva') + '?tipo=REC_COMPRA', datos, follow=True)
        self.assertContains(r, 'ya se usó')
        antes = Operacion.objects.count()
        r = self.client.post(reverse('inventario:nueva') + '?tipo=REC_COMPRA', {**datos, 'items-0-cantidad': '500',
                                                                               'referencia': 'G-2'})
        self.assertContains(r, 'no se guardó')
        self.assertEqual(Operacion.objects.count(), antes)
        r = self.client.post(reverse('inventario:nueva') + '?tipo=REC_COMPRA',
                             {**datos, 'fecha': (HOY + timedelta(days=3)).isoformat(), 'referencia': 'G-3'})
        self.assertContains(r, 'no puede ser futura')

    def test_lista_busca_por_orden_columnas_y_excel(self):
        """ALM-13 y LIS-01."""
        oc = self.oc()
        servicios.confirmar(self.recibir(oc, '10', referencia='Guía 77'), self.admin)
        url = reverse('inventario:lista')
        r = self.client.get(url + f'?q={oc.numero}')
        self.assertContains(r, oc.numero)
        r = self.client.get(url + '?q=20100070970')  # RUC del proveedor
        self.assertContains(r, 'PROVEEDOR DEMO')
        self.client.post(reverse('lista_preferencias'), {'lista': 'inventario.operaciones', 'accion': 'columnas',
                                                         'col': ['numero', 'guia', 'registro'], 'vuelta': url})
        self.assertEqual(PreferenciaLista.objects.get(usuario=self.admin).columnas, ['numero', 'guia', 'registro'])
        r = self.client.get(url)
        self.assertContains(r, 'Guía 77')
        self.assertIn('spreadsheetml', self.client.get(url + '?formato=excel')['Content-Type'])
        self.assertEqual(self.client.get(url + '?agrupar=tercero').status_code, 200)
        self.client.post(reverse('lista_preferencias'), {'lista': 'compras.ordenes', 'accion': 'guardar_filtro',
                                                         'nombre': 'Aprobadas', 'query': 'estado=APROBADO',
                                                         'vuelta': reverse('compras:oc_lista')})
        r = self.client.get(reverse('compras:oc_lista'))
        self.assertContains(r, 'Aprobadas')
        self.assertIn('spreadsheetml', self.client.get(reverse('compras:oc_lista') + '?formato=excel')['Content-Type'])


class AprobacionesTest(Base):
    def test_nivel_sin_aprobador_se_salta_o_usa_el_reemplazo(self):
        """APR-02: el gerente general (sin jefe) no deja la orden trabada."""
        gerente = User.objects.create_user('gerente', password='x')
        operaciones = User.objects.create_user('operaciones', password='x')
        ReglaAprobacion.objects.create(documento='OC', nivel=1, tipo='JEFE')
        ReglaAprobacion.objects.create(documento='OC', nivel=2, tipo='USUARIO', aprobador=operaciones)
        avisos = []
        niveles = motor.calcular_niveles('OC', gerente, D('50000'), avisos=avisos)
        self.assertEqual([n['usuario'] for n in niveles], [operaciones.pk])
        self.assertIn('no tiene jefe', avisos[0])
        ReglaAprobacion.objects.filter(nivel=1).update(reemplazo=self.admin)
        niveles = motor.calcular_niveles('OC', gerente, D('50000'))
        self.assertEqual([n['usuario'] for n in niveles], [self.admin.pk, operaciones.pk])

    def test_ruta_muestra_el_rechazo(self):
        """APR-06."""
        from compras import solped
        from contabilidad.models import CentroCosto
        uno, dos = User.objects.create_user('uno', password='x'), User.objects.create_user('dos', password='x')
        ReglaAprobacion.objects.create(documento='SOLPED', nivel=1, tipo='USUARIO', aprobador=uno)
        ReglaAprobacion.objects.create(documento='SOLPED', nivel=2, tipo='USUARIO', aprobador=dos)
        cc = CentroCosto.objects.create(codigo='ADM', nombre='Administración')
        sp = solped.crear({'centro_costo': cc, 'tipo': 'MATERIAL', 'motivo': 'Prueba'},
                          [{'producto': self.harina, 'cantidad': D('5')}], self.admin)
        solped.enviar(sp, self.admin)
        solped.liberar(sp, uno, 'ok')
        solped.rechazar(sp, dos, 'Sin presupuesto este mes')
        sp.refresh_from_db()
        ruta = motor.ruta('SOLPED', sp, sp.niveles, sp.nivel_aprobado, rechazado=True)
        self.assertEqual([f['estado'] for f in ruta], ['APROBADO', 'RECHAZADO'])
        self.assertEqual(ruta[1]['comentario'], 'Sin presupuesto este mes')
        r = self.client.get(reverse('solicitudes:detalle', args=[sp.pk]))
        self.assertContains(r, 'Rechazado por')

    def test_pedido_se_envia_a_aprobacion_y_sale_en_la_bandeja(self):
        """APR-03 y APR-04."""
        from ventas.models import Cotizacion, CotizacionItem
        vendedor, jefe = User.objects.create_user('vendedor', password='x'), User.objects.create_user('jefe', password='x')
        ReglaAprobacion.objects.create(documento='PEDIDO', nivel=1, tipo='USUARIO', aprobador=jefe)
        cliente = Tercero.objects.create(tipo='CLIENTE', tipo_doc='6', numero_doc='20600000011', nombre='CLIENTE SAC')
        cot = Cotizacion.objects.create(numero='PD01-00000005', tipo='PED', tercero=cliente)
        CotizacionItem.objects.create(documento=cot, descripcion='Servicio', cantidad=1, precio_unitario=D('50000'))
        cot.calcular_totales()
        cot.save()
        from ventas.aprobacion_pedido import enviar, estado_aprobacion
        with patch('ventas.aprobacion_pedido._solicitante', return_value=vendedor):
            ruta, debe_enviar, espera = estado_aprobacion(cot, vendedor)
            self.assertTrue(debe_enviar)
            self.assertEqual([f['nombre'] for f in ruta], ['jefe'])
            self.assertIn('nivel 1', enviar(cot, vendedor))
            espera = estado_aprobacion(cot, vendedor)[2]
            self.assertTrue(espera.endswith('nivel 1: jefe'), espera)
            self.assertIn('usted registró este pedido', espera)  # AUD-VEN-08: al creador se le explica
            from core.views_aprobaciones import _pendientes
            self.assertIn(cot, [f[2] for f in _pendientes(jefe)])

    def test_movimiento_sin_reglas_lo_aprueba_el_permiso_visible(self):
        """PLA-01."""
        from planillas.organizacion import quien_aprueba
        from types import SimpleNamespace
        from django.contrib.auth.models import Group
        from .modulos import GRUPOS
        planillero = User.objects.create_user('planillero', password='x')
        planillero.groups.add(Group.objects.get_or_create(name=GRUPOS['empleados'])[0])  # v1.83: módulo Empleados
        PerfilUsuario.objects.create(usuario=planillero, acciones=['empleados.aprobar_movimientos'])
        m = SimpleNamespace(niveles=[{'nivel': 1, 'usuario': None, 'nombre': 'Aprobador de movimientos de personal'}],
                            nivel_aprobado=0, solicitante_id=self.admin.pk)
        nombres = quien_aprueba(m)
        self.assertIn('planillero', nombres)
        self.assertNotIn('admin', nombres)
        self.assertEqual(motor.PERMISOS['MOVIMIENTO'], 'empleados.aprobar_movimientos')

    def test_usuario_nuevo_sin_acciones(self):
        """SEG-05: mínimo privilegio y aviso de acciones de módulos no asignados."""
        from .usuarios import UsuarioForm, acciones_sin_modulo
        self.assertEqual(UsuarioForm(editor=self.admin).initial.get('acciones'), [])
        u = User.objects.create_user('ventas1', password='x')
        PerfilUsuario.objects.create(usuario=u, acciones=['compras.aprobar_oc'])
        self.assertTrue(acciones_sin_modulo(u))
