"""v1.98.0: el plan de cuentas se exporta en el formato de la plantilla de carga masiva y se vuelve a importar."""
import io

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from contabilidad.models import CuentaContable
from contabilidad import pcge

from . import carga_masiva as cm


class ExportarPlanCuentasTest(TestCase):
    def setUp(self):
        self.client.force_login(User.objects.create_superuser('t', 't@t.com', 'x'))
        if not CuentaContable.objects.filter(codigo='941').exists():
            from contabilidad.models import CuentaDefecto
            pcge.cargar(CuentaContable, CuentaDefecto)

    def exportar(self, **params):
        r = self.client.get(reverse('contabilidad:plan'), {'formato': 'excel', **params})
        self.assertEqual(r.status_code, 200)
        self.assertIn('plan_de_cuentas_', r['Content-Disposition'])
        return load_workbook(io.BytesIO(r.content))

    def test_mismo_formato_que_la_plantilla(self):
        wb = self.exportar()
        ws = wb['Datos']
        plantilla = cm.plantilla('cuentas')['Datos']
        self.assertEqual([c.value for c in ws[1]], [c.value for c in plantilla[1]])
        self.assertEqual(ws.max_row - 1, CuentaContable.objects.count())
        self.assertIn('Instrucciones', wb.sheetnames)
        # el filtro de la pantalla se respeta
        ws = self.exportar(clase='6')['Datos']
        self.assertTrue(all(str(f[0]).startswith('6') for f in ws.iter_rows(min_row=2, values_only=True)))
        self.assertContains(self.client.get(reverse('contabilidad:plan')), '?formato=excel')

    def test_se_completa_y_se_vuelve_a_importar(self):
        c = CuentaContable.objects.filter(imputable=True, codigo__startswith='63').exclude(destino_debe=None).first()
        antes = {k: getattr(c, k) for k in ('naturaleza', 'columna_balance', 'moneda', 'tipo_cambio', 'analisis',
                                            'centro_costo_regla', 'destino_debe_id', 'destino_haber_id')}
        wb = self.exportar()
        ws = wb['Datos']
        for fila in ws.iter_rows(min_row=2):
            if fila[0].value == c.codigo:
                fila[1].value = 'Nombre corregido en Excel'
        ws.append(['634398', 'Cuenta nueva desde el Excel', '', 'Sí', '941', '7911'])
        buf = io.BytesIO()
        wb.save(buf)
        url = reverse('carga_masiva')
        r = self.client.post(url, {'tipo': 'cuentas', 'accion': 'validar', 'actualizar': '1',
                                   'archivo': SimpleUploadedFile('plan.xlsx', buf.getvalue())})
        self.assertNotContains(r, 'con error')
        self.client.post(url, {'tipo': 'cuentas', 'accion': 'confirmar', 'actualizar': '1'})
        c.refresh_from_db()
        self.assertEqual(c.nombre, 'Nombre corregido en Excel')
        self.assertEqual({k: getattr(c, k) for k in antes}, antes)  # lo demás queda igual
        nueva = CuentaContable.objects.get(codigo='634398')
        self.assertEqual((nueva.destino_debe.codigo, nueva.destino_haber.codigo), ('941', '7911'))
