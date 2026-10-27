from datetime import date, timedelta
from decimal import Decimal

from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone


def elegidos(request, nombre):
    """PLAT-13: valores marcados de un filtro de lista (varias casillas = varios valores), sin vacíos."""
    return [v for v in request.GET.getlist(nombre) if v not in ('', None)]


def filtrar(qs, request, nombre, campo=None):
    """Filtra por los valores elegidos (OR dentro del filtro); sin elegidos no filtra."""
    valores = elegidos(request, nombre)
    return qs.filter(**{f'{campo or nombre}__in': valores}) if valores else qs


def id_o_none(valor):
    """Identificador numérico de un formulario o None (vacío, «yo», texto): filter(pk=None) no encuentra nada en
    lugar del error 500 «Field 'id' expected a number» (EXC-03)."""
    valor = str(valor or '').strip()
    return int(valor) if valor.isdigit() else None


def excel_response(nombre, titulo, encabezados, filas, rojas=()):
    """Excel con título, encabezados y filas. `rojas`: [(fila, columna)] (desde 0, dentro de `filas`) que se pintan
    de rojo (datos que faltan)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = ''.join(c for c in titulo if c not in '[]:*?/\\')[:30] or 'Hoja1'
    ws.append([titulo])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(encabezados)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E79')
    for fila in filas:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in fila])
    rojo = PatternFill('solid', fgColor='FFC7CE')
    for i, j in rojas:
        ws.cell(row=4 + i, column=j + 1).fill = rojo
    for col in ws.columns:
        ancho = max(len(str(c.value or '')) for c in col[2:]) if len(col) > 2 else 10
        ws.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 10), 45)
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{nombre}.xlsx"'
    wb.save(resp)
    return resp


def txt_response(nombre, lineas):
    contenido = '\r\n'.join(lineas) + ('\r\n' if lineas else '')
    resp = HttpResponse(contenido, content_type='text/plain; charset=utf-8')
    resp['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return resp


def leer_excel(archivo):
    """Devuelve lista de dicts usando la primera fila como encabezados (en minúsculas)."""
    from openpyxl import load_workbook

    wb = load_workbook(archivo, data_only=True, read_only=True)
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    enc = [str(c or '').strip().lower() for c in filas[0]]
    return [dict(zip(enc, f)) for f in filas[1:] if any(v not in (None, '') for v in f)]


def a_fecha(valor):
    if isinstance(valor, date):
        return valor if not hasattr(valor, 'date') else valor.date()
    from datetime import datetime
    texto = str(valor).strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Fecha inválida: {valor}')


def periodo_actual(request, qs=None):
    """Periodo pedido; si no hay, el actual; y si el actual está vacío, el último con registros."""
    if request.GET.get('periodo'):
        return request.GET['periodo']
    actual = timezone.localdate().strftime('%Y%m')
    if qs is not None and not qs.filter(periodo=actual).exists():
        ultimo = qs.exclude(periodo='').order_by('-periodo').values_list('periodo', flat=True).first()
        if ultimo:
            return ultimo
    return actual


def _fin_de_mes(d):
    siguiente = (d.replace(day=28) + timedelta(days=4)).replace(day=1)
    return siguiente - timedelta(days=1)


def rango_por_defecto(request, qs=None, campo='fecha'):
    """(desde, hasta) ISO: lo pedido por GET, si no el mes en curso completo; si ese mes no tiene
    registros, el mes del último registro."""
    if request.GET.get('desde') and request.GET.get('hasta'):
        return request.GET['desde'], request.GET['hasta']
    hoy = timezone.localdate()
    desde, hasta = hoy.replace(day=1), _fin_de_mes(hoy)
    if qs is not None and not qs.filter(**{f'{campo}__range': [desde, hasta]}).exists():
        ultima = qs.order_by(f'-{campo}').values_list(campo, flat=True).first()
        if ultima:
            desde, hasta = ultima.replace(day=1), _fin_de_mes(ultima)
    return desde.isoformat(), hasta.isoformat()


def fmt_fecha(f):
    return f.strftime('%d/%m/%Y') if f else ''


def leer_numero(texto):
    """Decimal escrito a mano: 1,109.90 / 1109.90 / 287,50 (coma decimal) / 1.109,90. Con punto y coma, el último
    es el decimal; con solo comas, es decimal si hay una sola y no le siguen exactamente 3 dígitos (1,110 = mil
    ciento diez). Lanza decimal.InvalidOperation si no es un número."""
    t = str(texto or '').strip().replace(' ', '').replace(' ', '')
    if not t:
        return Decimal('0')
    if ',' in t and '.' in t:
        decimal_sep = ',' if t.rfind(',') > t.rfind('.') else '.'
        miles = '.' if decimal_sep == ',' else ','
        t = t.replace(miles, '').replace(decimal_sep, '.')
    elif ',' in t:
        partes = t.split(',')
        t = f'{partes[0]}.{partes[1]}' if len(partes) == 2 and len(partes[1]) != 3 else t.replace(',', '')
    return Decimal(t)


def lineas_formset(formset):
    """[(producto, cantidad)] de las filas válidas y no eliminadas de un formset de ítems."""
    lineas = []
    for f in formset.forms:
        cd = getattr(f, 'cleaned_data', None) or {}
        if cd and not cd.get('DELETE') and cd.get('cantidad'):
            lineas.append((cd.get('producto'), cd['cantidad']))
    return lineas


def saldo_minimo_desde(producto, almacen, fecha):
    """(saldo, fecha) más bajo del producto en el almacén desde `fecha` en adelante: el saldo al cierre de ese día y
    después de cada movimiento posterior. Es lo máximo que puede salir con esa fecha sin dejar el kardex en negativo
    en ningún momento (p. ej. un consumo del 05/10 no puede usar un ingreso del 10/10). Se calcula desde el stock
    actual descontando los movimientos posteriores, así no depende de los saldos migrados."""
    from .models import Kardex
    actual = producto.stock_en(almacen)
    posteriores = list(Kardex.objects.filter(producto=producto, almacen=almacen, fecha__gt=fecha)
                       .order_by('fecha', 'id').values_list('fecha', 'tipo', 'cantidad'))
    if not posteriores:
        return actual, fecha
    neto = sum((c if t == 'ENTRADA' else -c for _, t, c in posteriores), Decimal('0'))
    saldo = actual - neto  # al cierre del día `fecha`
    minimo = (saldo, fecha)
    for dia, tipo, cantidad in posteriores:
        saldo += cantidad if tipo == 'ENTRADA' else -cantidad
        if saldo < minimo[0]:
            minimo = (saldo, dia)
    return minimo


def faltantes_stock(lineas, almacen, devolver=None, fecha=None):
    """Mensajes por cada producto cuya salida supera el stock del almacén.

    devolver: {producto_id: cantidad} que el mismo documento ya había descontado (al editar).
    fecha: fecha del movimiento; si es anterior a hoy, la salida no puede dejar el kardex en negativo en esa fecha
    ni en ningún movimiento posterior.
    """
    from collections import defaultdict

    from django.utils import timezone

    from .models import Almacen, Empresa
    if Empresa.actual().permitir_stock_negativo:
        return []
    almacen = almacen or Almacen.principal()
    pedido, productos = defaultdict(Decimal), {}
    for producto, cantidad in lineas:
        if producto and producto.es_inventariable:
            pedido[producto.pk] += Decimal(cantidad)
            productos[producto.pk] = producto
    mensajes = []
    for pk, cantidad in pedido.items():
        extra = (devolver or {}).get(pk, Decimal('0'))
        if fecha and fecha < timezone.localdate():
            minimo, dia = saldo_minimo_desde(productos[pk], almacen, fecha)
            disponible = minimo + extra
            if cantidad > disponible:
                cuando = f'al {fecha:%d/%m/%Y}' if dia == fecha else f'el {dia:%d/%m/%Y}'
                mensajes.append(
                    f'Stock insuficiente de "{productos[pk].nombre}" en {almacen} {cuando}: con fecha '
                    f'{fecha:%d/%m/%Y} pueden salir {max(disponible, Decimal("0")):,.2f} y se necesitan '
                    f'{cantidad:,.2f}. Registre primero el ingreso con una fecha anterior o use la fecha de hoy.')
            continue
        disponible = productos[pk].stock_en(almacen) + extra
        if cantidad > disponible:
            mensajes.append(f'Stock insuficiente de "{productos[pk].nombre}" en {almacen}: hay {disponible:,.2f} y '
                            f'se necesitan {cantidad:,.2f}.')
    return mensajes


def anio_param(request, nombre='anio'):
    """Año de un filtro de la URL: ACT-01, «?anio=abc», «-1» o «99999» vuelven al año en curso (antes, error 500)."""
    hoy = timezone.localdate().year
    valor = (request.GET.get(nombre) or '').strip()
    return int(valor) if valor.isdigit() and 2000 <= int(valor) <= hoy + 1 else hoy


DIAS_ENVIO_SUNAT = 3  # días calendario que SUNAT admite para enviar un comprobante electrónico


class _Rechazo(Exception):
    pass


def _reglas_pedido(form, formset):
    """CRM-23 y OBR-38: en cotizaciones, pedidos y órdenes de compra la cantidad es mayor que cero, el precio no es
    negativo y la entrega no es anterior a la fecha del documento."""
    d = form.cleaned_data
    errores = []
    lineas = [f.cleaned_data for f in formset.forms
              if getattr(f, 'cleaned_data', None) and not f.cleaned_data.get('DELETE')]
    for n, l in enumerate(lineas, 1):
        if l.get('cantidad') is not None and l['cantidad'] <= 0:
            errores.append(f'Ítem {n}: la cantidad debe ser mayor a cero.')
        if l.get('precio_unitario') is not None and l['precio_unitario'] < 0:
            errores.append(f'Ítem {n}: el precio no puede ser negativo.')
    fecha = d.get('fecha') or getattr(form.instance, 'fecha', None)
    entrega = d.get('fecha_entrega')
    if fecha and entrega and entrega < fecha:
        errores.append(f'La fecha de entrega ({entrega:%d/%m/%Y}) no puede ser anterior a la del documento '
                       f'({fecha:%d/%m/%Y}).')
    return errores


def reglas_comprobante(form, formset):
    """Reglas de compras y ventas en el servidor (pantallas y API, no solo el navegador): cantidad y precio mayores a
    cero (valor 0 solo en operaciones gratuitas; las rebajas y devoluciones van con nota de crédito) (COM-18,
    VEN-15); vencimiento no anterior a la emisión (COM-19); emisión de una venta no futura ni fuera del plazo de envío
    a SUNAT (VEN-16); no se edita un documento con su asiento en un periodo cerrado (CON-15)."""
    from django.utils import timezone

    from .models import ComprobanteBase
    doc = form.instance
    if doc._meta.label_lower in ('ventas.cotizacion', 'compras.ordencompra'):
        return _reglas_pedido(form, formset)
    if not isinstance(doc, ComprobanteBase):
        return []
    d = form.cleaned_data
    errores = []
    gratuita = (d.get('tipo_operacion') or getattr(doc, 'tipo_operacion', '')) == 'GRATUITA'
    historico = getattr(doc, 'es_historico', False) or getattr(doc, 'es_saldo_inicial', False)
    lineas = [f.cleaned_data for f in formset.forms
              if getattr(f, 'cleaned_data', None) and not f.cleaned_data.get('DELETE')]
    for n, l in enumerate(lineas, 1):
        if l.get('cantidad') is not None and l['cantidad'] <= 0:
            errores.append(f'Ítem {n}: la cantidad debe ser mayor a cero (las devoluciones y rebajas se registran con '
                           f'una nota de crédito).')
        precio = l.get('precio_unitario')
        if precio is not None and (precio < 0 or (precio == 0 and not gratuita)):
            errores.append(f'Ítem {n}: el precio debe ser mayor a cero (valor 0 solo en una operación gratuita).')
    emision, vence = d.get('fecha_emision', doc.fecha_emision), d.get('fecha_vencimiento')
    if emision and vence and vence < emision:
        errores.append(f'El vencimiento ({vence:%d/%m/%Y}) no puede ser anterior a la emisión ({emision:%d/%m/%Y}).')
    errores += _reglas_sunat(doc, d, emision, historico)
    if doc._meta.label_lower == 'ventas.venta' and emision and not historico:
        hoy = timezone.localdate()
        if emision > hoy:
            errores.append(f'La fecha de emisión ({emision:%d/%m/%Y}) no puede ser posterior a hoy.')
        tipo = d.get('tipo_comprobante') or doc.tipo_comprobante
        # VEN-20: toda serie electrónica (empieza con letra), con o sin facturación electrónica configurada en Ceiba;
        # la serie física numérica (0001) puede registrarse después
        electronica = (d.get('serie') or doc.serie or 'F')[:1].isalpha()
        if (not doc.pk and tipo in ('01', '03', '07', '08') and electronica
                and emision < hoy - timedelta(days=DIAS_ENVIO_SUNAT)):
            errores.append(f'SUNAT acepta comprobantes electrónicos hasta {DIAS_ENVIO_SUNAT} días después de su '
                           f'emisión: use una fecha desde el {hoy - timedelta(days=DIAS_ENVIO_SUNAT):%d/%m/%Y}.')
    if doc.pk:
        from contabilidad.models import Asiento, PeriodoContable
        campo = 'venta' if doc._meta.label_lower == 'ventas.venta' else 'compra'
        cerrados = set(PeriodoContable.objects.filter(cerrado=True).values_list('periodo', flat=True))
        en_cerrado = Asiento.objects.filter(**{campo: doc}, periodo__in=cerrados).first()
        if en_cerrado:
            errores.append(f'Su asiento {en_cerrado.numero} está en el periodo cerrado {en_cerrado.periodo[4:]}/'
                           f'{en_cerrado.periodo[:4]}: corríjalo con una nota de crédito o débito, no editándolo.')
    return errores


import re as _re

FECHA_MINIMA = date(2000, 1, 1)
LARGO_GLOSA = 500
SERIE_VALIDA = _re.compile(r'^([A-Z][A-Z0-9]{3}|\d{1,4})$')
TIPOS_CON_SERIE = ('01', '03', '07', '08', '02', '04')  # factura, boleta, notas, RH y liquidación de compra
PERCEPCION_TASAS = (Decimal('0'), Decimal('0.5'), Decimal('1'), Decimal('2'))


def _reglas_sunat(doc, d, emision, historico):
    """COM-20 y VEN-20: serie y número del comprobante, periodo de registro, fechas imposibles, tipo de cambio,
    tasas de detracción, retención y percepción, recibo por honorarios a una persona natural, boleta de más de S/ 700
    con el documento del cliente, nota del mismo cliente y moneda que el comprobante (VEN-19) y glosa con límite."""
    from contabilidad.automatico import periodo_valido
    errores = []
    venta = doc._meta.label_lower == 'ventas.venta'
    tipo = d.get('tipo_comprobante') or doc.tipo_comprobante
    serie = (d.get('serie') or '').strip().upper()
    numero = str(d.get('numero') or '').strip()
    if tipo in TIPOS_CON_SERIE and not historico:
        if serie and not SERIE_VALIDA.match(serie):
            errores.append(f'Serie «{serie}» no válida: 4 caracteres que empiezan con letra (F001, B001, E001, FC01) '
                           f'o numérica del comprobante físico (0001).')
        if numero and not (numero.isdigit() and len(numero) <= 8):
            errores.append(f'Número «{numero}» no válido: correlativo numérico de hasta 8 dígitos.')
    if emision and emision < FECHA_MINIMA:
        errores.append(f'La fecha de emisión {emision:%d/%m/%Y} no es válida (desde el {FECHA_MINIMA:%d/%m/%Y}).')
    periodo = d.get('periodo')
    if periodo:
        if not periodo_valido(periodo):
            errores.append(f'Periodo «{periodo}» no válido: AAAAMM con mes 01 a 12.')
        elif emision and not venta:
            meses = (int(periodo[:4]) - emision.year) * 12 + int(periodo[4:]) - emision.month
            if meses < 0:
                errores.append('El periodo de registro no puede ser anterior al mes de emisión.')
            elif meses > 12:
                errores.append('Fuera del plazo de anotación: el periodo de registro es hasta 12 meses después de la '
                               'emisión (Ley del IGV, art. 19).')
            from django.utils import timezone
            if periodo > timezone.localdate().strftime('%Y%m'):
                errores.append('El periodo de registro no puede ser futuro.')
    tc = d.get('tipo_cambio')
    if tc is not None and (tc < 0 or (tc == 0 and (d.get('moneda') or getattr(doc, 'moneda', 'PEN')) == 'PEN')):
        errores.append('El tipo de cambio debe ser mayor a cero (vacío o 0 en dólares = el del día).')
    if (d.get('detraccion_pct') or 0) < 0 or (d.get('detraccion_pct') or 0) > 15:
        errores.append('Detracción entre 0 % y 15 % (tasas del anexo del SPOT).')
    honorarios = tipo == '02'
    retencion = d.get('retencion_pct') or Decimal('0')
    permitidas = (Decimal('0'), Decimal('8')) if honorarios and not venta else (Decimal('0'), Decimal('3'))
    if retencion not in permitidas:
        errores.append('Retención: 0 % u 8 % (renta de 4ta categoría) en el recibo por honorarios.' if honorarios
                       and not venta else 'Retención del IGV: 0 % o 3 %.')
    if (d.get('percepcion_pct') or Decimal('0')) not in PERCEPCION_TASAS:
        errores.append('Percepción del IGV: 0 %, 0.5 %, 1 % o 2 %.')
    if (d.get('icbper') or 0) < 0:
        errores.append('El ICBPER (impuesto a las bolsas) no puede ser negativo.')
    tercero = d.get('tercero') or getattr(doc, 'tercero', None)
    if honorarios and not venta and tercero is not None and not (
            tercero.tipo_doc == '6' and tercero.numero_doc[:2] in ('10', '15', '17')):
        errores.append(f'El recibo por honorarios lo emite una persona natural con RUC 10: {tercero.nombre} '
                       f'({tercero.numero_doc}) no lo es.')
    ref = d.get('doc_referencia')
    if tipo in ('07', '08') and ref is not None and tercero is not None:
        if ref.tercero_id != tercero.pk:
            errores.append(f'La nota debe ser del mismo cliente o proveedor que {ref} ({ref.tercero.nombre}).')
        moneda = d.get('moneda') or getattr(doc, 'moneda', '')
        if moneda and ref.moneda != moneda:
            errores.append(f'La nota debe estar en la misma moneda que {ref} ({ref.moneda}).')
    if len(d.get('glosa') or '') > LARGO_GLOSA:
        errores.append(f'La glosa tiene más de {LARGO_GLOSA} caracteres.')
    return errores


def _boleta_sin_documento(doc):
    """VEN-20: SUNAT exige identificar al cliente (DNI) en boletas de más de S/ 700."""
    if doc._meta.label_lower != 'ventas.venta' or doc.tipo_comprobante != '03' or doc.es_historico:
        return ''
    t = doc.tercero
    sin_doc = t.tipo_doc == '0' or not (t.numero_doc or '').strip('0').strip('-').strip()
    if sin_doc and (doc.total_pen or doc.total) > 700:
        return (f'Una boleta de más de S/ 700 (S/ {doc.total:,.2f}) identifica al cliente con su DNI u otro '
                f'documento: {t.nombre} no lo tiene.')
    return ''


def _nota_excede(doc):
    """VEN-17: el total de las notas de crédito de un comprobante no supera su total; la anulación (motivo 01 o 02)
    es por el saldo completo."""
    if doc.tipo_comprobante != '07' or not doc.doc_referencia_id:
        return ''
    ref = doc.doc_referencia
    previas = sum((n.total for n in type(doc).objects.filter(doc_referencia=ref, tipo_comprobante='07',
                                                             estado='REGISTRADO').exclude(pk=doc.pk)), Decimal('0'))
    saldo = (ref.total or Decimal('0')) - previas
    if doc.total > saldo:
        return (f'La nota de crédito (S/ {doc.total:,.2f}) supera lo que queda de {ref} (total S/ {ref.total:,.2f}, '
                f'notas anteriores S/ {previas:,.2f}): puede ser hasta S/ {saldo:,.2f}.')
    if getattr(doc, 'motivo_nota', '') in ('01', '02') and doc.total != saldo:
        return (f'Una nota de crédito por anulación (motivo {doc.motivo_nota}) es por el total pendiente de {ref}: '
                f'S/ {saldo:,.2f}.')
    return ''


def procesar_documento(data, form_class, formset_class, instance, al_guardar=None, validar=None):
    """Valida y guarda cabecera + ítems (pantallas y API). Devuelve (doc o None si hay errores, form, formset)."""
    form = form_class(data, instance=instance)
    formset = formset_class(data, instance=form.instance)
    if form.is_valid() and formset.is_valid():
        for mensaje in reglas_comprobante(form, formset) + (list(validar(form, formset)) if validar else []):
            form.add_error(None, mensaje)
    if not (form.is_valid() and formset.is_valid()):
        return None, form, formset
    pk_original = instance.pk
    try:
        return _guardar_documento(form, formset, instance, al_guardar)
    except _Rechazo as exc:
        form.add_error(None, str(exc))
        form.instance.pk = pk_original  # se deshizo: el formulario vuelve a ser el de antes
        modelo_item = formset.model
        existentes = set(modelo_item.objects.filter(pk__in=[f.instance.pk for f in formset.forms if f.instance.pk])
                         .values_list('pk', flat=True))
        for f in formset.forms:
            if f.instance.pk not in existentes:
                f.instance.pk = None
        return None, form, formset


def _guardar_documento(form, formset, instance, al_guardar):
    with transaction.atomic():
        referencia_anterior = None
        if instance.pk and getattr(instance, 'stock_aplicado', False):
            anterior = type(instance).objects.get(pk=instance.pk)
            if _firma_stock(anterior, _lineas_guardadas(anterior)) == _firma_stock(form.instance,
                                                                                   _lineas_form(formset)):
                # nada que mueva el almacén cambió (tipo, serie, glosa…): el kardex no se toca (INV-16)
                referencia_anterior = str(anterior)
            else:
                # revertir con los datos guardados (almacén e ítems anteriores); al_guardar vuelve a aplicar
                anterior.revertir_stock()
                instance.stock_aplicado = False
        if getattr(instance, 'moneda', None) == 'USD' and instance.tipo_cambio in (None, 0, 1):
            from .tipo_cambio import venta_del_dia
            fecha = getattr(instance, 'fecha_emision', None) or getattr(instance, 'fecha', None)
            instance.tipo_cambio = venta_del_dia(fecha)
        nuevo = instance.pk is None
        doc = form.save()
        formset.instance = doc
        formset.save()
        doc.calcular_totales()
        if nuevo and hasattr(doc, 'proponer_detraccion'):
            doc.proponer_detraccion()  # % de detracción del producto si la factura no trae otro
        doc.save()
        excede = (_nota_excede(doc) or _boleta_sin_documento(doc)) if hasattr(doc, 'doc_referencia_id') else ''
        if excede:
            raise _Rechazo(excede)  # deshace lo grabado (transacción)
        if referencia_anterior and referencia_anterior != str(doc):
            # cambió el tipo o el número: el kardex sigue apuntando al documento (anulaciones y lotes lo buscan)
            from .models import Kardex
            Kardex.objects.filter(referencia=referencia_anterior).update(referencia=str(doc))
            Kardex.objects.filter(referencia=f'Reversión {referencia_anterior}').update(
                referencia=f'Reversión {doc}')
        if al_guardar:
            al_guardar(doc)
    return doc, form, formset


def _firma_stock(doc, lineas):
    """Lo que define el movimiento de almacén de un comprobante: si no cambia, editar no toca el kardex."""
    almacen = getattr(doc, 'almacen_id', None)
    if almacen is None and hasattr(doc, 'almacen_id'):  # vacío = el principal (así lo aplica aplicar_stock)
        from .models import Almacen
        almacen = Almacen.principal().pk
    return (almacen, getattr(doc, 'fecha_emision', None),
            Decimal(doc.tipo_cambio or 0).quantize(Decimal('0.0001')), getattr(doc, 'moneda', None),
            getattr(doc, 'tipo_operacion', None), getattr(doc, 'tipo_comprobante', '') in ('07', '08'),
            getattr(doc, 'ingresar_almacen', None), getattr(doc, 'descontar_stock', None),
            getattr(doc, 'doc_referencia_id', None), sorted(lineas))


def _norma(valor, decimales='0.0001'):
    return Decimal(valor or 0).quantize(Decimal(decimales))


def _lineas_guardadas(doc):
    return [(i.producto_id, _norma(i.cantidad), _norma(i.precio_unitario), _norma(getattr(i, 'descuento_pct', 0)),
             getattr(i, 'afectacion', '') or '') for i in doc.items.all() if i.producto_id]


def _lineas_form(formset):
    lineas = []
    for f in formset.forms:
        c = getattr(f, 'cleaned_data', None) or {}
        if not c or c.get('DELETE') or not c.get('producto'):
            continue
        lineas.append((c['producto'].pk, _norma(c.get('cantidad')), _norma(c.get('precio_unitario')),
                       _norma(c.get('descuento_pct')), c.get('afectacion', '') or ''))
    return lineas


def guardar_documento(request, form_class, formset_class, instance, template, contexto, al_guardar=None,
                      initial=None, items_iniciales=None, validar=None):
    """Alta/edición de un documento con detalle de ítems (compras, ventas, OC, cotizaciones, guías).

    validar(form, formset) -> [mensajes]: validaciones de negocio antes de guardar (ej. stock).
    """
    if request.method == 'POST':
        doc, form, formset = procesar_documento(request.POST, form_class, formset_class, instance, al_guardar,
                                                validar)
        if doc is not None:
            return redirect(doc_detalle_url(doc))
    else:
        form = form_class(instance=instance, initial=initial)
        if items_iniciales:
            fs_cls = type('FormsetInicial', (formset_class,), {'extra': max(len(items_iniciales) - 1, 0)})
            formset = fs_cls(instance=form.instance, initial=items_iniciales)
        else:
            formset = formset_class(instance=form.instance)
    # solo los productos que ya están en el documento: los demás se piden uno a uno al elegirlos (antes viajaba
    # todo el catálogo, más de 1 MB, en cada formulario)
    from .opciones import datos_productos
    ids = set()
    for f in formset.forms:
        if 'producto' in f.fields:
            valor = f['producto'].value()
            if str(valor or '').isdigit():
                ids.add(int(valor))
    contexto.update(form=form, formset=formset, productos=datos_productos(request.user, ids))
    return render(request, template, contexto)


def doc_detalle_url(doc):
    from django.urls import reverse
    nombres = {
        'compra': 'compras:detalle', 'ordencompra': 'compras:oc_detalle',
        'venta': 'ventas:detalle', 'cotizacion': 'ventas:cot_detalle',
        'guiaremision': 'logistica:detalle',
    }
    return reverse(nombres[doc._meta.model_name], args=[doc.pk])
