"""Carga masiva desde Excel: maestros (productos, clientes/proveedores) y saldos iniciales de inventario.

Flujo: se descarga la plantilla, se sube el archivo, se valida fila por fila sin grabar nada y, si no hay
errores, se confirma la carga (todo o nada).
"""
import unicodedata
from collections import OrderedDict
from datetime import date
from decimal import Decimal, InvalidOperation

from django.db import transaction

from .models import Almacen, Producto, Tercero

D0 = Decimal('0')


class ErrorFila(Exception):
    pass


def normalizar(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return ' '.join(texto.replace('_', ' ').replace('*', '').upper().split())


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # códigos y documentos que Excel guarda como número
    return str(v).strip()


def _dec(v, campo, requerido=False, minimo=None):
    if v in (None, ''):
        if requerido:
            raise ErrorFila(f'{campo}: obligatorio.')
        return None
    try:
        d = Decimal(str(v).replace(',', '').strip())
    except InvalidOperation:
        raise ErrorFila(f'{campo}: "{v}" no es un número.')
    if minimo is not None and d < minimo:
        raise ErrorFila(f'{campo}: debe ser mayor o igual a {minimo}.')
    return d


def _elegir(v, opciones, campo, requerido=True, defecto=''):
    """opciones: {texto normalizado: valor}."""
    t = normalizar(v)
    if not t:
        if requerido and not defecto:
            raise ErrorFila(f'{campo}: obligatorio.')
        return defecto
    if t not in opciones:
        validos = ', '.join(sorted({k for k in opciones if len(k) > 2})[:10])
        raise ErrorFila(f'{campo}: "{v}" no es válido (use: {validos}).')
    return opciones[t]


def _opciones(choices, extras=None):
    ops = {}
    for valor, etiqueta in choices:
        ops[normalizar(valor)] = valor
        ops[normalizar(etiqueta)] = valor
    ops.update(extras or {})
    return ops


CLASES = _opciones(Producto.CLASES, {'SUMINISTRO': 'SUMINISTRO', 'SUMINISTROS': 'SUMINISTRO', 'ACTIVO': 'ACTIVO',
                                     'ACTIVO FIJO': 'ACTIVO', 'SEMIELABORADO': 'SEMIELABORADO'})
UNIDADES = _opciones(Producto.UNIDADES, {'UND': 'NIU', 'UNIDADES': 'NIU', 'KG': 'KGM', 'LT': 'LTR', 'M': 'MTR',
                                         'CAJA': 'BX'})
TIPOS_TERCERO = _opciones(Tercero.TIPOS)
TIPOS_DOC = _opciones(Tercero.TIPOS_DOC, {'CE': '4', 'PASAPORTE': '7'})


# ---------------------------------------------------------------- definiciones
# columna: (clave, título, obligatorio, ejemplo, ayuda)
DEFINICIONES = OrderedDict([
    ('productos', {
        'titulo': 'Productos y servicios', 'icono': 'bi-box-seam', 'modulos': ['inventario', 'compras', 'ventas'],
        'descripcion': 'Maestro de productos. Si el código se deja vacío se genera según el tipo.',
        'columnas': [
            ('codigo', 'Código', False, '', 'Vacío = automático (MP000001...). Si ya existe se actualiza.'),
            ('nombre', 'Nombre', True, 'Harina de trigo x 50 kg', ''),
            ('tipo', 'Tipo de producto', True, 'Materia prima',
             'Mercadería, Materia prima, Semi elaborado, Producto terminado, Suministros, Activo fijo o Servicio'),
            ('unidad', 'Unidad', False, 'KGM', 'NIU (unidad), KGM, LTR, MTR, BX, PK, GLL, ZZ. Vacío = NIU'),
            ('marca', 'Marca', False, '', ''),
            ('codigo_barras', 'Código de barras', False, '', ''),
            ('precio_venta', 'Precio de venta sin IGV', False, '0', ''),
            ('precio_compra', 'Precio de compra sin IGV', False, '95.50', ''),
            ('stock_minimo', 'Stock mínimo', False, '10', ''),
            ('punto_reorden', 'Punto de reorden', False, '20', ''),
            ('stock_maximo', 'Stock máximo', False, '100', ''),
        ]}),
    ('terceros', {
        'titulo': 'Clientes y proveedores', 'icono': 'bi-people', 'modulos': ['contactos', 'compras', 'ventas'],
        'descripcion': 'Se valida el RUC (dígito verificador), el DNI y el ubigeo. Si el documento ya existe se '
                       'actualiza.',
        'columnas': [
            ('tipo', 'Tipo', True, 'Proveedor', 'Cliente, Proveedor o Ambos'),
            ('tipo_doc', 'Tipo de documento', True, 'RUC', 'RUC, DNI, CE, Pasaporte u Otros'),
            ('numero_doc', 'N° documento', True, '20555555556', ''),
            ('nombre', 'Nombre / razón social', True, 'DISTRIBUIDORA EJEMPLO S.A.C.', ''),
            ('direccion', 'Dirección', False, 'Av. Principal 123', ''),
            ('ubigeo', 'Ubigeo', False, '150131', '6 dígitos INEI'),
            ('email', 'Correo', False, '', ''),
            ('telefono', 'Teléfono', False, '', ''),
            ('dias_credito', 'Días de crédito', False, '30', ''),
        ]}),
    ('saldos_cxc', {
        'titulo': 'Saldos iniciales por cobrar', 'icono': 'bi-person-down', 'modulos': ['ventas', 'finanzas'],
        'descripcion': 'Comprobantes de venta pendientes de cobro emitidos antes de usar el sistema. Se cobran '
                       'normalmente; no van al registro de ventas ni a SUNAT y se contabilizan 12 contra 5911 '
                       '(apertura). El cliente debe existir.',
        'columnas': [
            ('numero_doc', 'RUC/DNI del cliente', True, '20555555556', ''),
            ('tipo_comprobante', 'Tipo de comprobante', True, '01', '01 factura, 03 boleta, 12 ticket, 00 otros'),
            ('serie', 'Serie', True, 'F001', ''),
            ('numero', 'Número', True, '1234', ''),
            ('fecha_emision', 'Fecha de emisión', True, '15/09/2026', 'dd/mm/aaaa'),
            ('fecha_vencimiento', 'Fecha de vencimiento', False, '15/10/2026', 'Vacío = la de emisión'),
            ('moneda', 'Moneda', False, 'PEN', 'PEN o USD. Vacío = PEN'),
            ('saldo', 'Saldo pendiente', True, '1180.00', 'Importe que falta cobrar (con IGV), en la moneda'),
            ('tipo_cambio', 'Tipo de cambio', False, '', 'Solo USD. Vacío = T.C. venta de la fecha de emisión'),
        ]}),
    ('saldos_cxp', {
        'titulo': 'Saldos iniciales por pagar', 'icono': 'bi-person-up', 'modulos': ['compras', 'finanzas'],
        'descripcion': 'Comprobantes de compra pendientes de pago recibidos antes de usar el sistema. Se pagan '
                       'normalmente; no van al registro de compras y se contabilizan 5911 (apertura) contra 42. El '
                       'proveedor debe existir.',
        'columnas': [
            ('numero_doc', 'RUC/DNI del proveedor', True, '20555555556', ''),
            ('tipo_comprobante', 'Tipo de comprobante', True, '01', '01 factura, 02 recibo por honorarios, 03 boleta, '
                                                                     '12 ticket, 00 otros'),
            ('serie', 'Serie', True, 'F001', ''),
            ('numero', 'Número', True, '5678', ''),
            ('fecha_emision', 'Fecha de emisión', True, '10/09/2026', 'dd/mm/aaaa'),
            ('fecha_vencimiento', 'Fecha de vencimiento', False, '10/10/2026', 'Vacío = la de emisión'),
            ('moneda', 'Moneda', False, 'PEN', 'PEN o USD. Vacío = PEN'),
            ('saldo', 'Saldo pendiente', True, '2360.00', 'Importe que falta pagar (con IGV), en la moneda'),
            ('tipo_cambio', 'Tipo de cambio', False, '', 'Solo USD. Vacío = T.C. venta de la fecha de emisión'),
        ]}),
    ('saldos', {
        'titulo': 'Saldos iniciales de inventario', 'icono': 'bi-flag', 'modulos': ['inventario'], 'costos': True,
        'descripcion': 'Crea y confirma una operación "Saldo inicial" por almacén y fecha (contrapartida 5911).',
        'columnas': [
            ('codigo', 'Código del producto', True, 'MP000001', 'El producto debe existir y ser inventariable'),
            ('almacen', 'Código del almacén', False, 'ALM01', 'Vacío = almacén principal'),
            ('cantidad', 'Cantidad', True, '150', ''),
            ('costo_unitario', 'Costo unitario S/ sin IGV', True, '92.30', ''),
            ('fecha', 'Fecha', False, '', 'dd/mm/aaaa. Vacío = hoy'),
        ]}),
    ('trabajadores', {
        'titulo': 'Trabajadores (planillas)', 'icono': 'bi-person-badge', 'modulos': ['planillas'],
        'descripcion': 'Datos laborales para la planilla. Si el documento ya existe se actualiza (marque '
                       '"Actualizar existentes").',
        'columnas': [
            ('numero_doc', 'N° documento', True, '45678912', ''),
            ('tipo_doc', 'Tipo de documento', False, 'DNI', 'DNI, CE o Pasaporte. Vacío = DNI'),
            ('apellido_paterno', 'Apellido paterno', True, 'QUISPE', ''),
            ('apellido_materno', 'Apellido materno', False, 'MAMANI', ''),
            ('nombres', 'Nombres', True, 'ROSA ELENA', ''),
            ('fecha_ingreso', 'Fecha de ingreso', True, '01/03/2024', 'dd/mm/aaaa'),
            ('cargo', 'Cargo', False, 'Operaria de producción', ''),
            ('tipo', 'Tipo', False, 'Obrero', 'Empleado u Obrero. Vacío = Empleado'),
            ('regimen', 'Régimen laboral', False, 'General', 'General, Pequeña o Micro. Vacío = General'),
            ('centro_costo', 'Centro de costo (código)', False, '924001', ''),
            ('sueldo', 'Remuneración básica', True, '1500', ''),
            ('asignacion_familiar', 'Asignación familiar', False, 'SI', 'SI o NO'),
            ('sistema_pensiones', 'Sistema de pensiones', True, 'AFP', 'ONP, AFP o Ninguno'),
            ('afp', 'AFP', False, 'INTEGRA', 'Código de la AFP: HABITAT, INTEGRA, PRIMA o PROFUTURO'),
            ('comision', 'Tipo de comisión AFP', False, 'Flujo', 'Flujo o Mixta. Vacío = Flujo'),
            ('cuspp', 'CUSPP', False, '', ''),
            ('banco', 'Banco', False, 'BCP', ''), ('cuenta_sueldo', 'Cuenta de sueldo', False, '', ''),
            ('quinta_remuneracion_previa', 'Remuneraciones previas del año', False, '0',
             'Quinta: lo percibido en el año antes de usar el sistema (incluye gratificaciones)'),
            ('quinta_retencion_previa', 'Retenciones de quinta previas', False, '0', ''),
        ]}),
    # maestros de costos y manufactura: cárguelos en este orden (cada uno usa los anteriores)
    ('centros_beneficio', {
        'titulo': '1. Centros de beneficio', 'icono': 'bi-briefcase', 'modulos': ['contabilidad', 'manufactura'],
        'descripcion': 'Líneas de negocio. Orden de carga: centros de beneficio, centros de costo, puestos de '
                       'trabajo, recetas, hojas de ruta y versiones de fabricación.',
        'columnas': [
            ('codigo', 'Código', True, 'GAL', ''), ('nombre', 'Nombre', True, 'Galletas', ''),
            ('responsable', 'Responsable', False, 'Gerente de línea', ''),
        ]}),
    ('centros_costo', {
        'titulo': '2. Centros de costo', 'icono': 'bi-bullseye', 'modulos': ['contabilidad', 'manufactura'],
        'descripcion': 'Con tipo, jerarquía (centro del que depende) y centro de beneficio. El centro superior '
                       'debe existir o venir antes en el mismo archivo.',
        'columnas': [
            ('codigo', 'Código', True, 'PL-HOR', ''), ('nombre', 'Nombre', True, 'Hornos', ''),
            ('tipo', 'Tipo', True, 'Producción', 'Producción, Servicio, Administración, Ventas o Finanzas'),
            ('padre', 'Depende de (código)', False, 'PL', ''),
            ('centro_beneficio', 'Centro de beneficio (código)', False, 'GAL', ''),
            ('responsable', 'Responsable', False, '', ''),
        ]}),
    ('puestos', {
        'titulo': '3. Puestos de trabajo', 'icono': 'bi-tools', 'modulos': ['manufactura'],
        'descripcion': 'Máquinas, líneas o puestos manuales con sus tarifas, capacidad y calendario.',
        'columnas': [
            ('codigo', 'Código', True, 'HOR1', ''), ('nombre', 'Nombre', True, 'Horno 1', ''),
            ('tipo', 'Tipo', False, 'Máquina', 'Máquina, Línea de producción o Puesto manual'),
            ('costo_hora_mo', 'Tarifa mano de obra S/ h', True, '20', ''),
            ('costo_hora_cif', 'Tarifa máquina y CIF S/ h', True, '30', ''),
            ('centro_costo', 'Centro de costo (código)', False, 'PL-HOR', ''),
            ('horas_turno', 'Horas por turno', False, '8', ''), ('turnos', 'Turnos por día', False, '1', ''),
            ('dias_laborables', 'Días laborables', False, '123456', '1 = lunes … 7 = domingo'),
            ('eficiencia', 'Eficiencia %', False, '100', ''),
        ]}),
    ('recetas', {
        'titulo': '4. Listas de materiales', 'icono': 'bi-diagram-3', 'modulos': ['manufactura'],
        'descripcion': 'Una fila por insumo; las filas con el mismo producto y receta forman una lista. Productos e '
                       'insumos deben existir.',
        'columnas': [
            ('producto', 'Producto a fabricar (código)', True, 'PT000001', ''),
            ('receta', 'Código de la receta', True, 'G1', ''),
            ('cantidad_base', 'Rinde (cantidad por lote)', True, '100', ''),
            ('insumo', 'Insumo (código)', True, 'MP000001', ''),
            ('cantidad', 'Cantidad del insumo', True, '12.5', 'Para el lote indicado en "Rinde"'),
            ('merma', 'Merma %', False, '2', ''),
            ('operacion', 'Se consume en la op.', False, '10', ''),
        ]}),
    ('hojas_ruta', {
        'titulo': '5. Hojas de ruta', 'icono': 'bi-signpost-split', 'modulos': ['manufactura'],
        'descripcion': 'Una fila por operación; las filas con el mismo código forman una hoja. El puesto debe existir.',
        'columnas': [
            ('hoja', 'Código de la hoja', True, 'R-GAL', ''), ('nombre', 'Nombre de la hoja', True, 'Horneado', ''),
            ('secuencia', 'Operación', True, '10', '10, 20, 30…'), ('puesto', 'Puesto (código)', True, 'HOR1', ''),
            ('descripcion', 'Descripción', True, 'Hornear', ''),
            ('preparacion', 'Preparación (h por orden)', False, '0.5', ''),
            ('por_unidad', 'Ejecución (h por unidad)', False, '0.01', ''),
            ('espera', 'Espera (h)', False, '0', ''),
        ]}),
    ('versiones', {
        'titulo': '6. Versiones de fabricación', 'icono': 'bi-layers', 'modulos': ['manufactura'],
        'descripcion': 'Une receta y hoja de ruta para un rango de lote y una vigencia.',
        'columnas': [
            ('producto', 'Producto (código)', True, 'PT000001', ''), ('version', 'Versión', True, '1', ''),
            ('receta', 'Receta (código)', True, 'G1', ''), ('hoja', 'Hoja de ruta (código)', False, 'R-GAL', ''),
            ('lote_desde', 'Lote desde', False, '0', ''), ('lote_hasta', 'Lote hasta', False, '', 'Vacío = sin tope'),
            ('lote_costeo', 'Lote de costeo', False, '100', ''),
            ('vigente_desde', 'Vigente desde', False, '', 'dd/mm/aaaa. Vacío = hoy'),
            ('dias_fabricacion', 'Plazo de fabricación (días)', False, '1', ''),
        ]}),
])

# códigos que vienen en el mismo archivo (para referencias a filas anteriores: centro superior, etc.)
_CONTEXTO = {'codigos': set()}


# ---------------------------------------------------------------- validación por tipo
def _fila_producto(d, actualizar):
    nombre = _txt(d.get('nombre'))
    if not nombre:
        raise ErrorFila('Nombre: obligatorio.')
    clase = _elegir(d.get('tipo'), CLASES, 'Tipo de producto')
    codigo = _txt(d.get('codigo')).upper()
    existente = Producto.objects.filter(codigo=codigo).first() if codigo else None
    if existente and not actualizar:
        raise ErrorFila(f'El código {codigo} ya existe (marque "Actualizar existentes" para modificarlo).')
    if existente and existente.clase != clase:
        raise ErrorFila(f'El producto {codigo} es {existente.get_clase_display()}: el tipo no se puede cambiar.')
    datos = {'nombre': nombre[:200], 'clase': clase,
             'unidad': _elegir(d.get('unidad'), UNIDADES, 'Unidad', requerido=False, defecto='NIU'),
             'marca': _txt(d.get('marca'))[:80], 'codigo_barras': _txt(d.get('codigo_barras'))[:40]}
    for campo, titulo in (('precio_venta', 'Precio de venta'), ('precio_compra', 'Precio de compra'),
                          ('stock_minimo', 'Stock mínimo'), ('punto_reorden', 'Punto de reorden'),
                          ('stock_maximo', 'Stock máximo')):
        valor = _dec(d.get(campo), titulo, minimo=D0)
        if valor is not None:
            datos[campo] = str(valor)
    return {'accion': 'Actualizar' if existente else 'Nuevo', 'codigo': codigo, 'datos': datos,
            'resumen': f'{codigo or "(automático)"} · {nombre} · {dict(Producto.CLASES)[clase]}'}


def _fila_tercero(d, actualizar):
    from .forms import error_ruc
    from . import ubigeo
    tipo = _elegir(d.get('tipo'), TIPOS_TERCERO, 'Tipo')
    tipo_doc = _elegir(d.get('tipo_doc'), TIPOS_DOC, 'Tipo de documento')
    numero = _txt(d.get('numero_doc'))
    if not numero:
        raise ErrorFila('N° documento: obligatorio.')
    if tipo_doc == '6' and error_ruc(numero):
        raise ErrorFila(error_ruc(numero))
    if tipo_doc == '1' and (len(numero) != 8 or not numero.isdigit()):
        raise ErrorFila('El DNI debe tener 8 dígitos.')
    nombre = _txt(d.get('nombre'))
    if not nombre:
        raise ErrorFila('Nombre: obligatorio.')
    ubi = _txt(d.get('ubigeo')).zfill(6) if _txt(d.get('ubigeo')) else ''
    if ubi and not ubigeo.existe(ubi):
        raise ErrorFila(f'Ubigeo {ubi} no existe.')
    email = _txt(d.get('email'))
    if email and '@' not in email:
        raise ErrorFila(f'Correo inválido: {email}.')
    existente = Tercero.objects.filter(numero_doc=numero).first()
    if existente and not actualizar:
        raise ErrorFila(f'El documento {numero} ya existe (marque "Actualizar existentes").')
    dias = _dec(d.get('dias_credito'), 'Días de crédito', minimo=D0)
    datos = {'tipo': tipo, 'tipo_doc': tipo_doc, 'numero_doc': numero, 'nombre': nombre[:200],
             'direccion': _txt(d.get('direccion'))[:250], 'ubigeo': ubi, 'email': email[:254],
             'telefono': _txt(d.get('telefono'))[:50]}
    if dias is not None:
        datos['dias_credito'] = int(dias)
    if existente and existente.tipo != tipo and 'AMBOS' not in (tipo, existente.tipo):
        datos['tipo'] = 'AMBOS'  # era cliente y ahora también proveedor (o al revés)
    return {'accion': 'Actualizar' if existente else 'Nuevo', 'codigo': numero, 'datos': datos,
            'resumen': f'{numero} · {nombre}'}


def _fila_saldo(d, actualizar):
    from inventario.cierre import error_cierre
    from .utils import a_fecha
    codigo = _txt(d.get('codigo')).upper()
    p = Producto.objects.filter(codigo=codigo).first()
    if not p:
        raise ErrorFila(f'Producto {codigo or "(vacío)"} no existe: cárguelo primero en el maestro de productos.')
    if not p.es_inventariable:
        raise ErrorFila(f'{codigo} es {p.get_clase_display()}: no lleva inventario.')
    alm_cod = _txt(d.get('almacen')).upper()
    alm = Almacen.objects.filter(codigo__iexact=alm_cod).first() if alm_cod else Almacen.principal()
    if not alm:
        raise ErrorFila(f'Almacén {alm_cod} no existe.')
    cantidad = _dec(d.get('cantidad'), 'Cantidad', requerido=True)
    if cantidad <= 0:
        raise ErrorFila('Cantidad: debe ser mayor a cero.')
    costo = _dec(d.get('costo_unitario'), 'Costo unitario', requerido=True, minimo=D0)
    try:
        fecha = a_fecha(d.get('fecha')) if d.get('fecha') not in (None, '') else date.today()
    except ValueError as exc:
        raise ErrorFila(str(exc))
    if fecha > date.today():
        raise ErrorFila('La fecha no puede ser futura.')
    if error_cierre(fecha):
        raise ErrorFila(error_cierre(fecha))
    return {'accion': 'Saldo inicial', 'codigo': codigo,
            'datos': {'producto': p.pk, 'almacen': alm.pk, 'cantidad': str(cantidad), 'costo': str(costo),
                      'fecha': fecha.isoformat()},
            'resumen': f'{codigo} · {p.nombre} · {alm.nombre} · {cantidad:,.2f} × S/ {costo:,.4f}'}


def _fila_documento(d, es_venta):
    from .utils import a_fecha
    numero_doc = _txt(d.get('numero_doc'))
    tipos_tercero = ['CLIENTE', 'AMBOS'] if es_venta else ['PROVEEDOR', 'AMBOS']
    tercero = Tercero.objects.filter(numero_doc=numero_doc).first()
    if not tercero:
        raise ErrorFila(f'{"Cliente" if es_venta else "Proveedor"} {numero_doc or "(vacío)"} no existe: cárguelo '
                        f'primero en clientes y proveedores.')
    if tercero.tipo not in tipos_tercero:
        raise ErrorFila(f'{tercero.nombre} está registrado como {tercero.get_tipo_display().lower()}.')
    tipo = _txt(d.get('tipo_comprobante')).zfill(2)
    permitidos = ('01', '03', '12', '00') if es_venta else ('01', '02', '03', '12', '14', '00')
    if tipo not in permitidos:
        raise ErrorFila(f'Tipo de comprobante "{tipo}" no válido (use {", ".join(permitidos)}).')
    serie, numero = _txt(d.get('serie')).upper(), _txt(d.get('numero'))
    if not serie or not numero:
        raise ErrorFila('Serie y número son obligatorios.')
    try:
        emision = a_fecha(d.get('fecha_emision'))
        vence = a_fecha(d.get('fecha_vencimiento')) if d.get('fecha_vencimiento') not in (None, '') else emision
    except ValueError as exc:
        raise ErrorFila(str(exc))
    if emision > date.today():
        raise ErrorFila('La fecha de emisión no puede ser futura.')
    moneda = (_txt(d.get('moneda')) or 'PEN').upper()
    if moneda not in ('PEN', 'USD'):
        raise ErrorFila('Moneda: PEN o USD.')
    saldo = _dec(d.get('saldo'), 'Saldo pendiente', requerido=True)
    if saldo <= 0:
        raise ErrorFila('Saldo pendiente: debe ser mayor a cero.')
    tc = _dec(d.get('tipo_cambio'), 'Tipo de cambio', minimo=Decimal('0.001')) if moneda == 'USD' else Decimal('1')
    if tc is None:
        from .tipo_cambio import venta_del_dia
        tc = venta_del_dia(emision)
    from compras.models import Compra
    from ventas.models import Venta
    modelo = Venta if es_venta else Compra
    filtro = {'tipo_comprobante': tipo, 'serie': serie, 'numero': numero}
    if not es_venta:
        filtro['tercero'] = tercero
    if modelo.objects.filter(**filtro).exists():
        raise ErrorFila(f'El comprobante {serie}-{numero} ya está registrado.')
    return {'accion': 'Saldo inicial', 'codigo': f'{tercero.numero_doc}-{tipo}-{serie}-{numero}',
            'datos': {'tercero': tercero.pk, 'tipo': tipo, 'serie': serie, 'numero': numero,
                      'emision': emision.isoformat(), 'vence': vence.isoformat(), 'moneda': moneda,
                      'saldo': str(saldo), 'tc': str(tc)},
            'resumen': f'{tercero.nombre} · {tipo} {serie}-{numero} · vence {vence:%d/%m/%Y} · '
                       f'{"US$" if moneda == "USD" else "S/"} {saldo:,.2f}'}


def _producto(codigo, campo):
    p = Producto.objects.filter(codigo=_txt(codigo).upper()).first()
    if not p:
        raise ErrorFila(f'{campo}: el producto {_txt(codigo) or "(vacío)"} no existe.')
    return p


def _fila_beneficio(d, actualizar):
    from contabilidad.models import CentroBeneficio
    codigo, nombre = _txt(d.get('codigo')).upper(), _txt(d.get('nombre'))
    if not codigo or not nombre:
        raise ErrorFila('Código y nombre son obligatorios.')
    existe = CentroBeneficio.objects.filter(codigo=codigo).exists()
    if existe and not actualizar:
        raise ErrorFila(f'{codigo} ya existe (marque "Actualizar existentes").')
    return {'accion': 'Actualizar' if existe else 'Nuevo', 'codigo': codigo,
            'datos': {'nombre': nombre[:100], 'responsable': _txt(d.get('responsable'))[:120]},
            'resumen': f'{codigo} · {nombre}'}


def _fila_centro_costo(d, actualizar):
    from contabilidad.models import CentroBeneficio, CentroCosto
    codigo, nombre = _txt(d.get('codigo')).upper(), _txt(d.get('nombre'))
    if not codigo or not nombre:
        raise ErrorFila('Código y nombre son obligatorios.')
    tipo = _elegir(d.get('tipo'), _opciones(CentroCosto.TIPOS, {'SERVICIO': 'SERVICIO', 'APOYO': 'SERVICIO'}),
                   'Tipo')
    padre = _txt(d.get('padre')).upper()
    if padre and padre == codigo:
        raise ErrorFila('Un centro no puede depender de sí mismo.')
    if padre and not CentroCosto.objects.filter(codigo=padre).exists() and padre not in _CONTEXTO['codigos']:
        raise ErrorFila(f'Centro superior {padre} no existe (cárguelo antes, en este archivo o en otro).')
    cb = _txt(d.get('centro_beneficio')).upper()
    if cb and not CentroBeneficio.objects.filter(codigo=cb).exists():
        raise ErrorFila(f'Centro de beneficio {cb} no existe.')
    existe = CentroCosto.objects.filter(codigo=codigo).exists()
    if existe and not actualizar:
        raise ErrorFila(f'{codigo} ya existe (marque "Actualizar existentes").')
    return {'accion': 'Actualizar' if existe else 'Nuevo', 'codigo': codigo,
            'datos': {'nombre': nombre[:100], 'tipo': tipo, 'padre': padre, 'centro_beneficio': cb,
                      'responsable': _txt(d.get('responsable'))[:120]},
            'resumen': f'{codigo} · {nombre} · {dict(CentroCosto.TIPOS)[tipo]}' + (f' · depende de {padre}' if padre else '')}


def _fila_puesto(d, actualizar):
    from contabilidad.models import CentroCosto
    from produccion.models import CentroTrabajo
    codigo, nombre = _txt(d.get('codigo')).upper(), _txt(d.get('nombre'))
    if not codigo or not nombre:
        raise ErrorFila('Código y nombre son obligatorios.')
    cc = _txt(d.get('centro_costo')).upper()
    if cc and not CentroCosto.objects.filter(codigo=cc).exists():
        raise ErrorFila(f'Centro de costo {cc} no existe.')
    dias = _txt(d.get('dias_laborables')) or '123456'
    if not set(dias) <= set('1234567'):
        raise ErrorFila('Días laborables: dígitos del 1 (lunes) al 7 (domingo).')
    existe = CentroTrabajo.objects.filter(codigo=codigo).exists()
    if existe and not actualizar:
        raise ErrorFila(f'{codigo} ya existe (marque "Actualizar existentes").')
    eficiencia = _dec(d.get('eficiencia'), 'Eficiencia', minimo=Decimal('1')) or Decimal('100')
    return {'accion': 'Actualizar' if existe else 'Nuevo', 'codigo': codigo, 'datos': {
        'nombre': nombre[:100], 'tipo': _elegir(d.get('tipo'), _opciones(CentroTrabajo.TIPOS), 'Tipo',
                                                requerido=False, defecto='MAQUINA'),
        'costo_hora_mo': str(_dec(d.get('costo_hora_mo'), 'Tarifa mano de obra', True, D0)),
        'costo_hora_cif': str(_dec(d.get('costo_hora_cif'), 'Tarifa máquina y CIF', True, D0)),
        'centro_costo': cc, 'horas_turno': str(_dec(d.get('horas_turno'), 'Horas por turno', minimo=D0) or 8),
        'turnos': int(_dec(d.get('turnos'), 'Turnos', minimo=D0) or 1), 'dias_laborables': ''.join(sorted(set(dias))),
        'eficiencia': str(eficiencia)}, 'resumen': f'{codigo} · {nombre}'}


def _fila_receta(d, actualizar):
    from produccion.models import ListaMateriales
    producto = _producto(d.get('producto'), 'Producto')
    if producto.clase not in ('PRODUCTO_TERMINADO', 'SEMIELABORADO'):
        raise ErrorFila(f'{producto.codigo} es {producto.get_clase_display()}: solo se fabrican productos terminados '
                        'o semielaborados.')
    receta = _txt(d.get('receta')).upper()
    if not receta:
        raise ErrorFila('Código de la receta: obligatorio.')
    insumo = _producto(d.get('insumo'), 'Insumo')
    if insumo.pk == producto.pk:
        raise ErrorFila('Un producto no puede ser insumo de su propia receta.')
    if ListaMateriales.objects.filter(producto=producto, codigo=receta).exists():
        raise ErrorFila(f'La receta {receta} de {producto.codigo} ya existe: cree una nueva versión con otro código.')
    base = _dec(d.get('cantidad_base'), 'Rinde', True, Decimal('0.01'))
    cantidad = _dec(d.get('cantidad'), 'Cantidad del insumo', True, Decimal('0.0001'))
    merma = _dec(d.get('merma'), 'Merma', minimo=D0) or D0
    op = _dec(d.get('operacion'), 'Operación', minimo=D0)
    return {'accion': 'Nueva receta', 'codigo': f'{producto.codigo}|{receta}|{insumo.codigo}', 'datos': {
        'producto': producto.pk, 'receta': receta, 'base': str(base), 'insumo': insumo.pk, 'cantidad': str(cantidad),
        'merma': str(merma), 'operacion': int(op) if op else None},
        'resumen': f'{producto.codigo} {receta} · {insumo.nombre} {cantidad:g} por {base:g}'}


def _fila_hoja(d, actualizar):
    from produccion.models import CentroTrabajo, HojaRuta
    hoja, nombre = _txt(d.get('hoja')).upper(), _txt(d.get('nombre'))
    if not hoja or not nombre:
        raise ErrorFila('Código y nombre de la hoja son obligatorios.')
    if HojaRuta.objects.filter(codigo=hoja).exists():
        raise ErrorFila(f'La hoja {hoja} ya existe: cree otra con otro código.')
    puesto = CentroTrabajo.objects.filter(codigo=_txt(d.get('puesto')).upper()).first()
    if not puesto:
        raise ErrorFila(f'Puesto {_txt(d.get("puesto")) or "(vacío)"} no existe.')
    secuencia = _dec(d.get('secuencia'), 'Operación', True, Decimal('1'))
    prep = _dec(d.get('preparacion'), 'Preparación', minimo=D0) or D0
    unidad = _dec(d.get('por_unidad'), 'Ejecución', minimo=D0) or D0
    if not prep and not unidad:
        raise ErrorFila('Indique horas de preparación o de ejecución.')
    descripcion = _txt(d.get('descripcion'))
    if not descripcion:
        raise ErrorFila('Descripción: obligatoria.')
    return {'accion': 'Nueva hoja', 'codigo': f'{hoja}|{int(secuencia)}', 'datos': {
        'hoja': hoja, 'nombre': nombre[:120], 'secuencia': int(secuencia), 'puesto': puesto.pk,
        'descripcion': descripcion[:100], 'preparacion': str(prep), 'por_unidad': str(unidad),
        'espera': str(_dec(d.get('espera'), 'Espera', minimo=D0) or D0)},
        'resumen': f'{hoja} · {int(secuencia)} {descripcion} · {puesto.codigo}'}


def _fila_version(d, actualizar):
    from produccion.models import HojaRuta, ListaMateriales, VersionFabricacion
    from .utils import a_fecha
    producto = _producto(d.get('producto'), 'Producto')
    version = _txt(d.get('version')).upper()
    if not version:
        raise ErrorFila('Versión: obligatoria.')
    if VersionFabricacion.objects.filter(producto=producto, codigo=version).exists():
        raise ErrorFila(f'La versión {version} de {producto.codigo} ya existe.')
    lista = ListaMateriales.objects.filter(producto=producto, codigo=_txt(d.get('receta')).upper()).first()
    if not lista:
        raise ErrorFila(f'La receta {_txt(d.get("receta"))} de {producto.codigo} no existe.')
    hoja_cod = _txt(d.get('hoja')).upper()
    hoja = HojaRuta.objects.filter(codigo=hoja_cod).first() if hoja_cod else None
    if hoja_cod and not hoja:
        raise ErrorFila(f'La hoja de ruta {hoja_cod} no existe.')
    try:
        desde = a_fecha(d.get('vigente_desde')) if d.get('vigente_desde') not in (None, '') else date.today()
    except ValueError as exc:
        raise ErrorFila(str(exc))
    lote_min = _dec(d.get('lote_desde'), 'Lote desde', minimo=D0) or D0
    lote_max = _dec(d.get('lote_hasta'), 'Lote hasta', minimo=D0)
    if lote_max is not None and lote_max < lote_min:
        raise ErrorFila('Lote hasta debe ser mayor o igual a lote desde.')
    costeo = _dec(d.get('lote_costeo'), 'Lote de costeo', minimo=Decimal('0.01'))
    dias = _dec(d.get('dias_fabricacion'), 'Plazo', minimo=D0)
    return {'accion': 'Nueva versión', 'codigo': f'{producto.codigo}|{version}', 'datos': {
        'producto': producto.pk, 'version': version, 'lista': lista.pk, 'hoja': hoja.pk if hoja else None,
        'lote_min': str(lote_min), 'lote_max': str(lote_max) if lote_max is not None else None,
        'lote_costeo': str(costeo) if costeo else None, 'desde': desde.isoformat(),
        'dias': int(dias) if dias is not None else 1},
        'resumen': f'{producto.codigo} versión {version} · receta {lista.codigo}' + (f' · ruta {hoja.codigo}' if hoja else '')}


def _fila_trabajador(d, actualizar):
    from contabilidad.models import CentroCosto
    from planillas.models import AFP, Trabajador
    from .utils import a_fecha
    numero = _txt(d.get('numero_doc'))
    tipo_doc = {'': '01', 'DNI': '01', 'CE': '04', 'CARNE DE EXTRANJERIA': '04', 'PASAPORTE': '07'}.get(
        normalizar(d.get('tipo_doc')))
    if tipo_doc is None:
        raise ErrorFila('Tipo de documento: DNI, CE o Pasaporte.')
    if tipo_doc == '01':
        numero = numero.zfill(8)
        if len(numero) != 8 or not numero.isdigit():
            raise ErrorFila('El DNI debe tener 8 dígitos.')
    if not numero:
        raise ErrorFila('N° documento: obligatorio.')
    existente = Trabajador.objects.filter(numero_doc=numero).first()
    if existente and not actualizar:
        raise ErrorFila(f'El trabajador {numero} ya existe (marque "Actualizar existentes").')
    paterno, nombres = _txt(d.get('apellido_paterno')), _txt(d.get('nombres'))
    if not paterno or not nombres:
        raise ErrorFila('Apellido paterno y nombres son obligatorios.')
    try:
        ingreso = a_fecha(d.get('fecha_ingreso'))
    except ValueError as exc:
        raise ErrorFila(str(exc))
    sueldo = _dec(d.get('sueldo'), 'Remuneración básica', requerido=True, minimo=D0)
    pensiones = {'ONP': 'ONP', 'AFP': 'AFP', 'NINGUNO': 'NINGUNO', 'SIN': 'NINGUNO'}.get(
        normalizar(d.get('sistema_pensiones')))
    if not pensiones:
        raise ErrorFila('Sistema de pensiones: ONP, AFP o Ninguno.')
    afp = None
    if pensiones == 'AFP':
        afp = AFP.objects.filter(codigo__iexact=_txt(d.get('afp'))).first() or \
            AFP.objects.filter(nombre__icontains=_txt(d.get('afp')) or '-').first()
        if afp is None:
            raise ErrorFila(f'AFP "{_txt(d.get("afp"))}" no existe (use HABITAT, INTEGRA, PRIMA o PROFUTURO).')
    centro = _txt(d.get('centro_costo'))
    if centro and not CentroCosto.objects.filter(codigo=centro).exists():
        raise ErrorFila(f'Centro de costo {centro} no existe.')
    regimen = {'': 'GENERAL', 'GENERAL': 'GENERAL', 'PEQUENA': 'PEQUENA', 'PEQUENA EMPRESA': 'PEQUENA',
               'MICRO': 'MICRO', 'MICROEMPRESA': 'MICRO'}.get(normalizar(d.get('regimen')))
    if regimen is None:
        raise ErrorFila('Régimen laboral: General, Pequeña o Micro.')
    datos = {'tipo_doc': tipo_doc, 'apellido_paterno': paterno[:60], 'apellido_materno': _txt(
        d.get('apellido_materno'))[:60], 'nombres': nombres[:80], 'fecha_ingreso': ingreso.isoformat(),
        'cargo': _txt(d.get('cargo'))[:80], 'tipo': 'OBRERO' if normalizar(d.get('tipo')) == 'OBRERO' else 'EMPLEADO',
        'regimen': regimen, 'centro_costo': centro, 'sueldo': str(sueldo),
        'asignacion_familiar': normalizar(d.get('asignacion_familiar')) in ('SI', 'S', 'X', '1'),
        'sistema_pensiones': pensiones, 'afp': afp.pk if afp else None,
        'comision_afp': 'MIXTA' if normalizar(d.get('comision')) == 'MIXTA' else 'FLUJO',
        'cuspp': _txt(d.get('cuspp'))[:12], 'banco': _txt(d.get('banco'))[:30],
        'cuenta_sueldo': _txt(d.get('cuenta_sueldo'))[:30]}
    previa = _dec(d.get('quinta_remuneracion_previa'), 'Remuneraciones previas', minimo=D0) or D0
    retencion = _dec(d.get('quinta_retencion_previa'), 'Retenciones previas', minimo=D0) or D0
    if previa or retencion:
        datos.update(quinta_anio=date.today().year, quinta_remuneracion_previa=str(previa),
                     quinta_retencion_previa=str(retencion))
    return {'accion': 'Actualizar' if existente else 'Nuevo', 'codigo': numero, 'datos': datos,
            'resumen': f'{numero} · {paterno} {nombres} · S/ {sueldo:,.2f} · {pensiones}'}


VALIDADORES = {'trabajadores': _fila_trabajador, 'productos': _fila_producto, 'terceros': _fila_tercero, 'saldos': _fila_saldo,
               'saldos_cxc': lambda d, a: _fila_documento(d, True), 'saldos_cxp': lambda d, a: _fila_documento(d, False),
               'centros_beneficio': _fila_beneficio, 'centros_costo': _fila_centro_costo, 'puestos': _fila_puesto,
               'recetas': _fila_receta, 'hojas_ruta': _fila_hoja, 'versiones': _fila_version}


def _crear_documentos(filas, es_venta, sustentar=lambda obj: None):
    from compras.models import Compra, CompraItem
    from ventas.models import Venta, VentaItem
    modelo, item_modelo = (Venta, VentaItem) if es_venta else (Compra, CompraItem)
    total = Decimal('0')
    for f in filas:
        d = f['datos']
        saldo = Decimal(d['saldo'])
        extra = {'descontar_stock': False} if es_venta else {'ingresar_almacen': False, 'clasificacion': 'GASTO'}
        doc = modelo(tercero_id=d['tercero'], tipo_comprobante=d['tipo'], serie=d['serie'], numero=d['numero'],
                     fecha_emision=date.fromisoformat(d['emision']), fecha_vencimiento=date.fromisoformat(d['vence']),
                     moneda=d['moneda'], tipo_cambio=Decimal(d['tc']), tipo_operacion='INAFECTA',
                     forma_pago='CREDITO' if d['vence'] > d['emision'] else 'CONTADO', es_saldo_inicial=True,
                     glosa='Saldo inicial cargado desde Excel', **extra)
        if es_venta:
            doc.estado_sunat = 'NO_ENVIADO'
            doc.sunat_descripcion = 'Saldo inicial: no se envía a SUNAT.'
        doc.save()
        item_modelo.objects.create(documento=doc, descripcion='Saldo inicial pendiente', cantidad=1,
                                   precio_unitario=saldo)
        doc.calcular_totales()
        doc.retencion_monto = doc.percepcion_monto = doc.detraccion_monto = Decimal('0')  # el saldo ya es neto
        doc.save()
        sustentar(doc)
        total += doc.total_pen
    return total


def leer_archivo(archivo, tipo):
    """[{'n': fila excel, 'datos': {clave: valor}}] con encabezados normalizados (acepta clave o título)."""
    from openpyxl import load_workbook
    definicion = DEFINICIONES[tipo]
    alias = {}
    for clave, titulo, *_ in definicion['columnas']:
        alias[normalizar(clave)] = clave
        alias[normalizar(titulo)] = clave
    try:
        wb = load_workbook(archivo, data_only=True, read_only=True)
    except Exception as exc:
        raise ErrorFila(f'No se pudo leer el Excel (.xlsx): {exc}')
    ws = wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise ErrorFila('El archivo está vacío.')
    enc = [alias.get(normalizar(c)) for c in filas[0]]
    faltan = [titulo for clave, titulo, req, *_ in definicion['columnas'] if req and clave not in enc]
    if faltan:
        raise ErrorFila(f'Faltan las columnas: {", ".join(faltan)}. Use la plantilla.')
    salida = []
    for n, valores in enumerate(filas[1:], start=2):
        if not any(v not in (None, '') for v in valores):
            continue
        salida.append({'n': n, 'datos': {c: v for c, v in zip(enc, valores) if c}})
    if len(salida) > 5000:
        raise ErrorFila('Máximo 5,000 filas por archivo.')
    return salida


def validar(archivo, tipo, actualizar=False):
    """(filas, errores): cada fila con accion/resumen/datos o error. No graba nada."""
    validador = VALIDADORES[tipo]
    resultado, errores, vistos = [], 0, {}
    filas = leer_archivo(archivo, tipo)
    _CONTEXTO['codigos'] = {_txt(f['datos'].get('codigo')).upper() for f in filas}
    for fila in filas:
        try:
            r = validador(fila['datos'], actualizar)
            clave = r['codigo']
            if tipo != 'saldos' and clave and clave in vistos:
                raise ErrorFila(f'Repetido en el archivo (fila {vistos[clave]}).')
            if clave:
                vistos[clave] = fila['n']
            resultado.append({'n': fila['n'], **r})
        except ErrorFila as exc:
            errores += 1
            resultado.append({'n': fila['n'], 'error': str(exc), 'resumen': ' · '.join(
                _txt(v) for v in fila['datos'].values() if v not in (None, ''))[:120]})
    return resultado, errores


# ---------------------------------------------------------------- grabación
def cargar(tipo, filas, usuario, adjunto=None):
    """Graba las filas validadas (todo o nada). Devuelve un texto con el resumen.

    adjunto: (bytes, nombre) del Excel, que queda como sustento de los saldos iniciales creados."""
    from .sustentos import guardar_archivo, vincular
    archivo = None

    def sustentar(obj):
        nonlocal archivo
        if adjunto is None:
            return
        if archivo is None:
            archivo = guardar_archivo(None, usuario, datos=adjunto[0], nombre=adjunto[1],
                                      tipo='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        vincular(obj, archivo, usuario, 'Carga masiva desde Excel')

    with transaction.atomic():
        if tipo == 'productos':
            nuevos = actualizados = 0
            for f in filas:
                p = Producto.objects.filter(codigo=f['codigo']).first() if f['codigo'] else None
                if p is None:
                    p, nuevos = Producto(codigo=f['codigo']), nuevos + 1
                else:
                    actualizados += 1
                for campo, valor in f['datos'].items():
                    setattr(p, campo, Decimal(valor) if campo in ('precio_venta', 'precio_compra', 'stock_minimo',
                                                                   'punto_reorden', 'stock_maximo') else valor)
                p.save()
            return f'{nuevos} productos nuevos y {actualizados} actualizados.'
        if tipo == 'terceros':
            nuevos = actualizados = 0
            for f in filas:
                t = Tercero.objects.filter(numero_doc=f['codigo']).first()
                if t is None:
                    t, nuevos = Tercero(), nuevos + 1
                else:
                    actualizados += 1
                for campo, valor in f['datos'].items():
                    setattr(t, campo, valor)
                t.save()
            return f'{nuevos} clientes/proveedores nuevos y {actualizados} actualizados.'
        if tipo in ('saldos_cxc', 'saldos_cxp'):
            total = _crear_documentos(filas, tipo == 'saldos_cxc', sustentar)
            return (f'{len(filas)} documentos por {"cobrar" if tipo == "saldos_cxc" else "pagar"} cargados '
                    f'(S/ {total:,.2f}).')
        if tipo == 'saldos':
            from inventario import servicios
            from inventario.models import Operacion, TipoOperacion
            tipo_op = TipoOperacion.objects.get(codigo='SALDO_INI')
            grupos = OrderedDict()
            for f in filas:
                grupos.setdefault((f['datos']['almacen'], f['datos']['fecha']), []).append(f['datos'])
            numeros = []
            for (almacen, fecha), lineas in grupos.items():
                op = Operacion.objects.create(tipo=tipo_op, fecha=date.fromisoformat(fecha), almacen_destino_id=almacen,
                                              referencia='Carga masiva', creado_por=usuario,
                                              glosa='Saldos iniciales cargados desde Excel')
                for d in lineas:
                    op.items.create(producto_id=d['producto'], cantidad=Decimal(d['cantidad']),
                                    costo_unitario=Decimal(d['costo']))
                sustentar(op)
                servicios.confirmar(op, usuario)
                numeros.append(op.numero)
            return f'{len(filas)} saldos cargados en {len(numeros)} operación(es): {", ".join(numeros)}.'
        if tipo == 'trabajadores':
            from contabilidad.models import CentroCosto
            from planillas.models import Trabajador
            for f in filas:
                d = dict(f['datos'])
                d['fecha_ingreso'] = date.fromisoformat(d['fecha_ingreso'])
                d['sueldo'] = Decimal(d['sueldo'])
                for campo in ('quinta_remuneracion_previa', 'quinta_retencion_previa'):
                    if campo in d:
                        d[campo] = Decimal(d[campo])
                d['centro_costo'] = CentroCosto.objects.filter(codigo=d['centro_costo']).first() \
                    if d['centro_costo'] else None
                d['afp_id'] = d.pop('afp')
                Trabajador.objects.update_or_create(numero_doc=f['codigo'], defaults=d)
            return f'{len(filas)} trabajadores cargados.'
        if tipo in ('centros_beneficio', 'centros_costo', 'puestos', 'recetas', 'hojas_ruta', 'versiones'):
            return _cargar_maestro(tipo, filas)
    raise ValueError(tipo)


def _cargar_maestro(tipo, filas):
    from contabilidad.models import CentroBeneficio, CentroCosto
    from produccion.models import CentroTrabajo, HojaRuta, ListaMateriales, VersionFabricacion
    if tipo == 'centros_beneficio':
        for f in filas:
            CentroBeneficio.objects.update_or_create(codigo=f['codigo'], defaults=f['datos'])
        return f'{len(filas)} centros de beneficio cargados.'
    if tipo == 'centros_costo':
        for f in filas:  # en orden: el centro superior ya existe o vino antes en el archivo
            d = dict(f['datos'])
            d['padre'] = CentroCosto.objects.get(codigo=d['padre']) if d['padre'] else None
            d['centro_beneficio'] = CentroBeneficio.objects.get(codigo=d['centro_beneficio']) \
                if d['centro_beneficio'] else None
            CentroCosto.objects.update_or_create(codigo=f['codigo'], defaults=d)
        return f'{len(filas)} centros de costo cargados.'
    if tipo == 'puestos':
        for f in filas:
            d = dict(f['datos'])
            d['centro_costo'] = CentroCosto.objects.filter(codigo=d['centro_costo']).first() if d['centro_costo'] \
                else None
            for campo in ('costo_hora_mo', 'costo_hora_cif', 'horas_turno', 'eficiencia'):
                d[campo] = Decimal(d[campo])
            CentroTrabajo.objects.update_or_create(codigo=f['codigo'], defaults=d)
        return f'{len(filas)} puestos de trabajo cargados.'
    if tipo == 'recetas':
        listas = {}
        for f in filas:
            d = f['datos']
            clave = (d['producto'], d['receta'])
            if clave not in listas:
                listas[clave] = ListaMateriales.objects.create(producto_id=d['producto'], codigo=d['receta'],
                                                               cantidad_base=Decimal(d['base']), estado='APROBADA')
            listas[clave].componentes.create(producto_id=d['insumo'], cantidad=Decimal(d['cantidad']),
                                             merma=Decimal(d['merma']), operacion=d['operacion'])
        return f'{len(listas)} recetas cargadas ({len(filas)} insumos).'
    if tipo == 'hojas_ruta':
        hojas = {}
        for f in filas:
            d = f['datos']
            if d['hoja'] not in hojas:
                hojas[d['hoja']] = HojaRuta.objects.create(codigo=d['hoja'], nombre=d['nombre'], estado='APROBADA')
            hojas[d['hoja']].operaciones.create(
                secuencia=d['secuencia'], centro_id=d['puesto'], descripcion=d['descripcion'],
                horas_preparacion=Decimal(d['preparacion']), horas_unidad=Decimal(d['por_unidad']),
                horas_espera=Decimal(d['espera']))
        return f'{len(hojas)} hojas de ruta cargadas ({len(filas)} operaciones).'
    for f in filas:  # versiones
        d = f['datos']
        VersionFabricacion.objects.create(
            producto_id=d['producto'], codigo=d['version'], lista_id=d['lista'], hoja_id=d['hoja'],
            lote_min=Decimal(d['lote_min']), lote_max=Decimal(d['lote_max']) if d['lote_max'] else None,
            lote_costeo=Decimal(d['lote_costeo']) if d['lote_costeo'] else None,
            vigente_desde=date.fromisoformat(d['desde']), dias_fabricacion=d['dias'])
    return f'{len(filas)} versiones de fabricación cargadas.'


def plantilla(tipo):
    """Libro Excel con la hoja de datos (encabezados y ejemplo) e instrucciones."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    definicion = DEFINICIONES[tipo]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos'
    ws.append([c[0] for c in definicion['columnas']])
    ws.append([c[3] for c in definicion['columnas']])
    for celda, col in zip(ws[1], definicion['columnas']):
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor='0F3D2E' if col[2] else '1F7A52')
        ws.column_dimensions[celda.column_letter].width = max(14, len(col[1]) + 4)
    ayuda = wb.create_sheet('Instrucciones')
    ayuda.append([definicion['titulo']])
    ayuda['A1'].font = Font(bold=True, size=13)
    ayuda.append([definicion['descripcion']])
    ayuda.append(['Borre la fila de ejemplo. Las columnas en verde oscuro son obligatorias.'])
    ayuda.append([])
    ayuda.append(['Columna', 'Descripción', 'Obligatoria', 'Notas'])
    for c in ayuda[5]:
        c.font = Font(bold=True)
    for clave, titulo, req, _, nota in definicion['columnas']:
        ayuda.append([clave, titulo, 'Sí' if req else 'No', nota])
    for col, ancho in zip('ABCD', (18, 30, 12, 80)):
        ayuda.column_dimensions[col].width = ancho
    return wb
