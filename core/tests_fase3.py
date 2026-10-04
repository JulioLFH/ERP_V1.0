"""Pruebas de la Fase 3 (v1.8): costos ocultos según permiso, carga masiva desde Excel y cierre de kardex."""
import io
from datetime import date, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from compras.models import Compra, CompraItem
from core.forms import AjusteInventarioForm
from core.inventario import valor_inventario
from core.modulos import GRUPO_COSTOS, GRUPOS
from core.models import Almacen, Producto, Tercero
from core.usuarios import UsuarioForm
from inventario import cierre, servicios
from inventario.models import CierreKardex, Operacion, TipoOperacion

D = Decimal
HOY = date.today()
TC = (D('3.441'), D('3.450'))
FIN_MES_ANTERIOR = HOY.replace(day=1) - timedelta(days=1)


def excel(filas):
    wb = Workbook()
    for f in filas:
        wb.active.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile('carga.xlsx', buf.getvalue())


class Fase3Test(TestCase):
    @classmethod
    def setUpTestData(cls):
        with patch('core.tipo_cambio._consultar', return_value=TC):
            call_command('seed_demo', verbosity=0)
        cls.admin = User.objects.create_superuser('t', 't@t.com', 'x')
        cls.almacenero = User.objects.create_user('almacen', 'a@a.com', 'x')
        cls.almacenero.groups.add(Group.objects.get_or_create(name=GRUPOS['inventario'])[0],
                                  Group.objects.get_or_create(name=GRUPOS['compras'])[0])
        cls.p3 = Producto.objects.get(codigo='P003')

    def setUp(self):
        p = patch('core.tipo_cambio._consultar', return_value=TC)
        p.start()
        self.addCleanup(p.stop)
        self.client.force_login(self.admin)

    # ---------------------------------------------------------------- 11. costos según permiso
    def test_sin_permiso_no_ve_costos(self):
        self.client.force_login(self.almacenero)
        r = self.client.get(reverse('inv_stock'))
        self.assertNotContains(r, 'Costo prom.')
        self.assertNotContains(r, 'Inventario valorizado')
        r = self.client.get(reverse('inv_kardex') + f'?producto={self.p3.pk}')
        self.assertContains(r, 'Kardex de unidades')
        self.assertNotContains(r, 'C. prom.')
        self.assertEqual(self.client.get(reverse('inv_valorizacion')).status_code, 403)
        self.assertEqual(self.client.get(reverse('inventario:cierres')).status_code, 403)
        self.assertNotContains(self.client.get(reverse('productos')), 'Costo prom.')
        self.assertNotContains(self.client.get(reverse('compras:oc_nuevo')), 'costo_promedio')
        r = self.client.get(reverse('inv_kardex') + f'?producto={self.p3.pk}&formato=excel')
        hoja = load_workbook(io.BytesIO(r.content)).active
        self.assertNotIn('Costo promedio', [c.value for c in hoja[3]])
        self.assertNotIn('Valorización al cierre', r.content.decode(errors='ignore'))
        op = Operacion.objects.create(tipo=TipoOperacion.objects.get(codigo='CONS_INT'), fecha=HOY,
                                      almacen_origen=Almacen.principal())
        op.items.create(producto=self.p3, cantidad=1)
        self.assertNotContains(self.client.get(reverse('inventario:detalle', args=[op.pk])), 'Costo unit.')

    def test_permiso_ver_costos_desde_usuarios(self):
        form = UsuarioForm(data={'username': 'almacen', 'first_name': '', 'last_name': '', 'email': 'a@a.com',
                                 'is_active': 'on', 'modulos': ['inventario'], 'ver_costos': 'on'},
                           instance=self.almacenero, editor=self.admin)
        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.assertTrue(self.almacenero.groups.filter(name=GRUPO_COSTOS).exists())
        self.client.force_login(User.objects.get(pk=self.almacenero.pk))
        self.assertContains(self.client.get(reverse('inv_stock')), 'Costo prom.')
        self.assertEqual(self.client.get(reverse('inv_valorizacion')).status_code, 200)

    # ---------------------------------------------------------------- 13. cierre de kardex
    def test_cierre_bloquea_movimientos_y_guarda_valorizacion(self):
        prov = Tercero.objects.filter(tipo='PROVEEDOR').first()
        compra = Compra.objects.create(tercero=prov, serie='F777', numero='1', fecha_emision=FIN_MES_ANTERIOR)
        CompraItem.objects.create(documento=compra, producto=self.p3, descripcion='x', cantidad=5,
                                  precio_unitario=D('40'))
        compra.calcular_totales()
        compra.save()
        compra.aplicar_stock()
        periodo = FIN_MES_ANTERIOR.strftime('%Y%m')
        # un borrador en el periodo impide cerrar
        borrador = Operacion.objects.create(tipo=TipoOperacion.objects.get(codigo='AJ_SAL'), fecha=FIN_MES_ANTERIOR,
                                            almacen_origen=Almacen.principal())
        self.assertTrue(any('borrador' in e for e in cierre.errores_cierre(periodo)[0]))
        borrador.delete()
        r = self.client.post(reverse('inventario:cierres'), {'periodo': periodo, 'observaciones': 'Inventario físico'})
        c = CierreKardex.objects.get()
        self.assertRedirects(r, reverse('inventario:cierre', args=[c.pk]))
        self.assertEqual(c.fecha_corte, FIN_MES_ANTERIOR)
        self.assertEqual(c.valor_total, valor_inventario(FIN_MES_ANTERIOR)[1])
        self.assertEqual(sum(s.valor for s in c.saldos.all()), c.valor_total)
        # no se puede mover el almacén en el periodo cerrado
        with self.assertRaises(cierre.KardexCerrado):
            self.p3.mover_stock(1, 'prueba', fecha=FIN_MES_ANTERIOR)
        form = AjusteInventarioForm(data={'producto': self.p3.pk, 'almacen': Almacen.principal().pk,
                                          'tipo': 'ENTRADA', 'concepto': '', 'cantidad': '1', 'costo_unitario': '1',
                                          'fecha': FIN_MES_ANTERIOR, 'motivo': 'x'})
        self.assertFalse(form.is_valid())
        self.assertIn('fecha', form.errors)
        r = self.client.post(reverse('compras:anular', args=[compra.pk]), {'motivo': 'Error de registro'})
        compra.refresh_from_db()
        self.assertEqual(compra.estado, 'REGISTRADO')  # anularla revertiría stock de un periodo cerrado
        op = Operacion.objects.create(tipo=TipoOperacion.objects.get(codigo='AJ_ING'), fecha=FIN_MES_ANTERIOR,
                                      almacen_destino=Almacen.principal())
        op.items.create(producto=self.p3, cantidad=1, costo_unitario=D('40'))
        self.assertTrue(any('cerrado' in e for e in servicios.errores_confirmacion(op)))
        # el mes actual sigue abierto
        self.p3.mover_stock(1, 'prueba', costo=D('40'), fecha=HOY)
        # un administrador reabre el último cierre
        self.client.post(reverse('inventario:cierre_reabrir', args=[c.pk]), {'motivo': 'Compra fuera de plazo'})
        c.refresh_from_db()
        self.assertEqual(c.estado, 'REABIERTO')
        self.p3.mover_stock(1, 'prueba', costo=D('40'), fecha=FIN_MES_ANTERIOR)
        self.assertEqual(self.client.get(reverse('inventario:cierre', args=[c.pk]) + '?formato=excel').status_code, 200)

    def test_no_se_cierra_el_mes_en_curso(self):
        errores, _ = cierre.errores_cierre(HOY.strftime('%Y%m'))
        self.assertTrue(any('aún no termina' in e for e in errores))

    def test_cierre_vuelve_a_la_pantalla_con_mensaje(self):
        cierre.cerrar(FIN_MES_ANTERIOR.strftime('%Y%m'), self.admin)
        op = Operacion.objects.create(tipo=TipoOperacion.objects.get(codigo='AJ_ING'), fecha=FIN_MES_ANTERIOR,
                                      almacen_destino=Almacen.principal())
        op.items.create(producto=self.p3, cantidad=1, costo_unitario=D('40'))
        url = reverse('inventario:detalle', args=[op.pk])
        r = self.client.post(reverse('inventario:confirmar', args=[op.pk]), HTTP_REFERER=url)
        self.assertRedirects(r, url)
        op.refresh_from_db()
        self.assertEqual(op.estado, 'BORRADOR')

    # ---------------------------------------------------------------- 12. carga masiva
    def test_carga_masiva_de_productos(self):
        url = reverse('carga_masiva')
        r = self.client.get(url + '?tipo=productos&plantilla=1')
        self.assertEqual([c.value for c in load_workbook(io.BytesIO(r.content)).active[1]][:3],
                         ['codigo', 'nombre', 'tipo'])
        archivo = excel([['codigo', 'nombre', 'tipo', 'unidad', 'precio_compra'],
                         ['', 'Harina', 'Materia prima', 'KG', 3.5],
                         ['', 'Envase', 'Suministros', '', None],
                         ['', 'Mal', 'Juguete', '', None]])
        r = self.client.post(url, {'tipo': 'productos', 'accion': 'validar', 'archivo': archivo})
        self.assertContains(r, 'no es válido')
        self.assertContains(r, '1 con error')
        r = self.client.post(url, {'tipo': 'productos', 'accion': 'confirmar'})
        self.assertFalse(Producto.objects.filter(nombre='Harina').exists())  # con errores no se carga nada
        archivo = excel([['Código', 'Nombre', 'Tipo de producto', 'Unidad', 'Precio de compra sin IGV'],
                         ['', 'Harina', 'Materia prima', 'KG', 3.5],
                         ['', 'Envase', 'Suministros', '', None]])
        self.client.post(url, {'tipo': 'productos', 'accion': 'validar', 'archivo': archivo})
        self.client.post(url, {'tipo': 'productos', 'accion': 'confirmar'})
        harina = Producto.objects.get(nombre='Harina')
        self.assertEqual((harina.codigo[:2], harina.unidad, harina.precio_compra, harina.cuenta_existencias.codigo),
                         ('MP', 'KGM', D('3.5'), '2411'))
        # actualizar existentes
        filas = [['codigo', 'nombre', 'tipo', 'precio_compra'], [harina.codigo, 'Harina especial', 'Materia prima', 4]]
        r = self.client.post(url, {'tipo': 'productos', 'accion': 'validar', 'archivo': excel(filas)})
        self.assertContains(r, 'ya existe')
        self.client.post(url, {'tipo': 'productos', 'accion': 'validar', 'archivo': excel(filas), 'actualizar': '1'})
        self.client.post(url, {'tipo': 'productos', 'accion': 'confirmar'})
        harina.refresh_from_db()
        self.assertEqual((harina.nombre, harina.precio_compra), ('Harina especial', D('4')))

    def test_carga_masiva_de_terceros(self):
        url = reverse('carga_masiva')
        from core.forms import digito_ruc
        cliente = Tercero.objects.filter(tipo='CLIENTE', tipo_doc='6').first()
        ruc = f'2060000000{digito_ruc("2060000000")}'
        archivo = excel([['tipo', 'tipo_doc', 'numero_doc', 'nombre', 'ubigeo', 'dias_credito'],
                         ['Proveedor', 'RUC', '20555555557', 'RUC MALO SAC', '', ''],
                         ['Proveedor', 'RUC', int(ruc), 'NUEVO PROVEEDOR SAC', 150131, 30]])
        r = self.client.post(url, {'tipo': 'terceros', 'accion': 'validar', 'archivo': archivo})
        self.assertContains(r, '1 con error')
        self.assertContains(r, 'dígito verificador')
        archivo = excel([['tipo', 'tipo_doc', 'numero_doc', 'nombre', 'ubigeo', 'dias_credito'],
                         ['Proveedor', 'RUC', cliente.numero_doc, cliente.nombre, 150131, 30]])
        self.client.post(url, {'tipo': 'terceros', 'accion': 'validar', 'archivo': archivo, 'actualizar': '1'})
        self.client.post(url, {'tipo': 'terceros', 'accion': 'confirmar'})
        cliente.refresh_from_db()
        self.assertEqual((cliente.tipo, cliente.dias_credito, cliente.ubigeo), ('AMBOS', 30, '150131'))

    def test_carga_masiva_de_saldos_iniciales(self):
        nuevo = Producto.objects.create(nombre='Cable', clase='MERCADERIA')
        url = reverse('carga_masiva')
        archivo = excel([['codigo', 'almacen', 'cantidad', 'costo_unitario', 'fecha'],
                         [nuevo.codigo, 'ALM01', 100, 2.5, HOY.strftime('%d/%m/%Y')],
                         [nuevo.codigo, 'ALM02', 20, 2.5, HOY.strftime('%d/%m/%Y')],
                         ['NOEXISTE', '', 1, 1, None]])
        r = self.client.post(url, {'tipo': 'saldos', 'accion': 'validar', 'archivo': archivo})
        self.assertContains(r, 'no existe')
        archivo = excel([['codigo', 'almacen', 'cantidad', 'costo_unitario', 'fecha'],
                         [nuevo.codigo, 'ALM01', 100, 2.5, HOY.strftime('%d/%m/%Y')],
                         [nuevo.codigo, 'ALM02', 20, 2.5, HOY.strftime('%d/%m/%Y')]])
        self.client.post(url, {'tipo': 'saldos', 'accion': 'validar', 'archivo': archivo})
        self.client.post(url, {'tipo': 'saldos', 'accion': 'confirmar'})
        nuevo.refresh_from_db()
        self.assertEqual((nuevo.stock, nuevo.costo_promedio), (D('120'), D('2.5000')))
        ops = Operacion.objects.filter(tipo__codigo='SALDO_INI', estado='CONFIRMADO', referencia='Carga masiva')
        self.assertEqual(ops.count(), 2)  # una por almacén
        # sin permiso de costos no se ofrecen los saldos
        self.client.force_login(self.almacenero)
        r = self.client.get(url + '?tipo=saldos')
        self.assertEqual(r.context['tipo'], 'productos')
