"""SIRE por la API de SUNAT (v1.44): token con usuario SOL (la clave no se guarda), ticket de la propuesta,
descarga del ZIP, cruce con compras y ventas, creación de lo que falta y aceptación de la propuesta. SUNAT se simula
(no hay llamadas reales). Un escenario por clase."""
import io
import json
import urllib.parse
import zipfile
from datetime import date
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from compras.models import Compra, CompraItem
from ventas.models import Venta

from . import sire_api
from .models import Empresa, PerfilUsuario, PropuestaSIRE, Tercero
from .modulos import GRUPOS

D = Decimal
PERIODO = '202609'
CAB_RCE = ('RUC|Apellidos y Nombres o Razón social|Periodo|CAR SUNAT|Fecha de emisión|Fecha Vcto/Pago|'
           'Tipo CP/Doc.|Serie del CDP|Año|Nro CP o Doc. Nro Inicial (Rango)|Nro Final (Rango)|'
           'Tipo Doc Identidad|Nro Doc Identidad|Apellidos Nombres/ Razón Social|BI Gravado DG|IGV / IPM DG|'
           'Total CP|Moneda')
CAB_RVIE = ('Ruc|Apellidos y Nombres o Razón social|Periodo|CAR SUNAT|Fecha de emisión|Fecha Vcto/Pago|'
            'Tipo CP/Doc.|Serie del CDP|Nro CP o Doc. Nro Inicial (Rango)|Nro Final (Rango)|Tipo Doc Identidad|'
            'Nro Doc Identidad|Apellidos Nombres/ Razón Social|Valor Facturado Exportación|BI Gravada|Dscto BI|'
            'IGV / IPM|Dscto IGV / IPM|Mto Exonerado|Mto Inafecto|ISC|BI Grav IVAP|IVAP|ICBPER|Otros Tributos|'
            'Total CP|Moneda|Tipo Cambio')


def zip_txt(nombre, texto):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        z.writestr(nombre, texto.encode('utf-8'))
    return buf.getvalue()


class SunatSimulado:
    """Responde como la API del SIRE; anota cada llamada."""

    def __init__(self, contenido, terminado_tras=0):
        self.contenido, self.consultas_faltantes, self.llamadas = contenido, terminado_tras, []

    def __call__(self, metodo, url, datos=None, cabeceras=None, binario=False):
        self.llamadas.append((metodo, url, datos))
        if 'oauth2/token' in url:
            return {'access_token': 'TOKEN-SIMULADO', 'expires_in': 3600}
        if 'exportapropuesta' in url or 'exportacioncomprobantepropuesta' in url:
            return {'numTicket': '20260300000777'}
        if 'consultaestadotickets' in url:
            if self.consultas_faltantes:
                self.consultas_faltantes -= 1
                return {'registros': [{'numTicket': '20260300000777', 'codEstadoProceso': '02',
                                       'desEstadoProceso': 'En proceso'}]}
            return {'registros': [{'numTicket': '20260300000777', 'perTributario': PERIODO, 'codProceso': '10',
                                   'codEstadoProceso': '06', 'desEstadoProceso': 'Terminado',
                                   'detalleTicket': {'numTicket': '20260300000777'},
                                   'archivoReporte': [{'codTipoAchivoReporte': '01',
                                                       'nomArchivoReporte': 'propuesta.zip'}]}]}
        if 'archivoreporte' in url:
            return self.contenido
        if 'aceptapropuesta' in url or 'aceptarpropuesta' in url:
            return {'numTicket': '20260300000999'}
        raise AssertionError(f'URL no simulada: {url}')


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        with patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450'))):
            call_command('seed_demo', verbosity=0)
        cls.admin = User.objects.create_superuser('admin_sire', 'a@a.pe', 'Clave-segura-9')

    def setUp(self):
        p = patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450')))
        p.start()
        self.addCleanup(p.stop)
        cache.clear()
        Empresa.objects.update(sire_client_id='cid-sire', sire_client_secret='secreto-sire', sire_usuario_sol='CONTA01')
        self.prov = Tercero.objects.create(numero_doc='20600000022', nombre='ANDINA SAC', tipo='PROVEEDOR')
        self.client.force_login(self.admin)

    def compra(self, numero, precio):
        c = Compra.objects.create(tercero=self.prov, serie='F001', numero=numero, fecha_emision=date(2026, 9, 10))
        CompraItem.objects.create(documento=c, descripcion='Servicio', cantidad=1, precio_unitario=D(precio))
        c.calcular_totales()
        c.save()
        return c

    def rce(self, *filas):
        return '\n'.join([CAB_RCE] + [
            f'20100000001|MI EMPRESA|{PERIODO}|X|10/09/2026||01|F001||{n}||6|{ruc}|{nom}|{b}|{i}|{t}|PEN'
            for n, ruc, nom, b, i, t in filas])


class TraerComprasPorAPITest(Base):
    def test_trae_cruza_y_crea_lo_que_falta(self):
        ok = self.compra('101', '100')
        txt = self.rce(('101', '20600000022', 'ANDINA SAC', '100.00', '18.00', '118.00'),
                       ('205', '20611111111', 'PROVEEDOR NUEVO SAC', '50.00', '9.00', '59.00'))
        sunat = SunatSimulado(zip_txt('propuesta.txt', txt))
        with patch('core.sire_api._http', sunat):
            r = self.client.post(reverse('sire_propuestas'), {'libro': 'RCE', 'periodo': '2026-09',
                                                              'clave_sol': 'MiClaveSol'})
        p = PropuestaSIRE.objects.get()
        self.assertRedirects(r, reverse('sire_propuesta', args=[p.pk]), fetch_redirect_response=False)
        self.assertEqual((p.estado, p.ticket, len(p.registros)), ('LISTA', '20260300000777', 2))
        # token con RUC + usuario SOL; la clave viaja solo en esa llamada y no queda guardada
        token = urllib.parse.parse_qs(sunat.llamadas[0][2].decode())
        self.assertEqual(token['username'], [Empresa.actual().ruc + 'CONTA01'])
        self.assertEqual(token['scope'], ['https://api-sire.sunat.gob.pe'])
        self.assertNotIn('MiClaveSol', json.dumps(p.registros) + p.mensaje)
        self.assertIn('/rce/propuesta/web/propuesta/202609/exportacioncomprobantepropuesta', sunat.llamadas[1][1])
        self.assertIn('codLibro=080000', sunat.llamadas[-1][1])
        r = self.client.get(reverse('sire_propuesta', args=[p.pk]))
        self.assertEqual([x['compra'] for x in r.context['r']['coinciden']], [ok])
        falta = r.context['r']['solo_sunat'][0]
        self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'crear', 'falta': [falta['clave']]})
        nueva = Compra.objects.get(serie='F001', numero='205')
        self.assertEqual((nueva.total, nueva.igv, nueva.periodo, nueva.tercero.nombre),
                         (D('59.00'), D('9.00'), PERIODO, 'PROVEEDOR NUEVO SAC'))
        p.refresh_from_db()
        self.assertEqual(p.creados, 1)
        r = self.client.get(reverse('sire_propuesta', args=[p.pk]))
        self.assertEqual(len(r.context['r']['solo_sunat']), 0)
        resp = self.client.get(reverse('sire_propuesta', args=[p.pk]), {'formato': 'excel'})
        self.assertIn('spreadsheetml', resp['Content-Type'])
        # todas las columnas de la propuesta de SUNAT, leídas del archivo guardado
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(resp.content)).active
        filas = [list(f) for f in hoja.iter_rows(min_row=3, values_only=True)]
        columnas = CAB_RCE.split('|')
        self.assertEqual(filas[0][:len(columnas) + 1], ['Resultado del cruce'] + columnas)
        f205 = next(f for f in filas if f[columnas.index('Nro CP o Doc. Nro Inicial (Rango)') + 1] == '205')
        self.assertEqual((f205[0], f205[columnas.index('Periodo') + 1], f205[columnas.index('Total CP') + 1]),
                         ('Coincide', PERIODO, 59))


class TicketEnProcesoTest(Base):
    def test_queda_solicitada_y_se_consulta_despues(self):
        sunat = SunatSimulado(zip_txt('p.txt', self.rce(('300', '20600000022', 'ANDINA SAC', '10', '1.80', '11.80'))),
                              terminado_tras=1)
        with patch('core.sire_api._http', sunat):
            self.client.post(reverse('sire_propuestas'), {'libro': 'RCE', 'periodo': '2026-09', 'clave_sol': 'x'})
            p = PropuestaSIRE.objects.get()
            self.assertEqual(p.estado, 'SOLICITADA')
            self.assertContains(self.client.get(reverse('sire_propuesta', args=[p.pk])), 'Consultar a SUNAT de nuevo')
            self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'consultar'})  # token en caché
        p.refresh_from_db()
        self.assertEqual((p.estado, len(p.registros)), ('LISTA', 1))


class VentasPorArchivoTest(Base):
    def test_rvie_crea_ventas_ya_informadas(self):
        cliente = Tercero.objects.filter(tipo='CLIENTE').first()
        Venta.objects.create(tercero=cliente, tipo_comprobante='01', serie='F001', numero='501',
                             fecha_emision=date(2026, 9, 5), total=D('118'), igv=D('18'), base_imponible=D('100'))
        fila = lambda tipo, serie, num, td, doc, nom, b, i, t: (
            f'20100000001|MI EMPRESA|{PERIODO}|X|05/09/2026||{tipo}|{serie}|{num}||{td}|{doc}|{nom}|0|{b}|0|{i}|0|'
            f'0|0|0|0|0|0|0|{t}|PEN|')
        txt = '\n'.join([CAB_RVIE, fila('01', 'F001', '501', '6', cliente.numero_doc, cliente.nombre, '100.00',
                                        '18.00', '118.00'),
                         fila('03', 'B001', '77', '0', '-', 'CLIENTES VARIOS', '20.00', '3.60', '23.60')])
        self.client.post(reverse('sire_propuestas'), {
            'libro': 'RVIE', 'periodo': '2026-09', 'archivo': SimpleUploadedFile('rvie.txt', txt.encode())})
        p = PropuestaSIRE.objects.get()
        self.assertEqual((p.origen, p.estado, len(p.registros)), ('ARCHIVO', 'LISTA', 2))
        r = self.client.get(reverse('sire_propuesta', args=[p.pk]))
        self.assertEqual(len(r.context['r']['coinciden']), 1)
        falta = r.context['r']['solo_sunat'][0]
        self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'crear', 'falta': [falta['clave']]})
        v = Venta.objects.get(serie='B001', numero='77')
        self.assertEqual((v.estado_sunat, v.descontar_stock, v.total, v.tercero.numero_doc),
                         ('ACEPTADO', False, D('23.60'), '00000000'))
        self.assertNotContains(self.client.get(reverse('sire_propuesta', args=[p.pk])), 'Aceptar la propuesta')


class AceptarPropuestaTest(Base):
    def test_solo_con_permiso_de_cierre(self):
        sunat = SunatSimulado(zip_txt('p.txt', self.rce(('400', '20600000022', 'ANDINA SAC', '10', '1.80', '11.80'))))
        with patch('core.sire_api._http', sunat):
            self.client.post(reverse('sire_propuestas'), {'libro': 'RCE', 'periodo': '2026-09', 'clave_sol': 'x'})
            p = PropuestaSIRE.objects.get()
            comprador = User.objects.create_user('comprador_sire', password='Clave-segura-9')
            comprador.groups.add(Group.objects.get_or_create(name=GRUPOS['compras'])[0])
            PerfilUsuario.objects.create(usuario=comprador, acciones=['compras.registrar'])
            self.client.force_login(comprador)
            self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'aceptar'})
            p.refresh_from_db()
            self.assertIsNone(p.aceptada_en)
            self.client.force_login(self.admin)
            self.client.post(reverse('sire_propuesta', args=[p.pk]), {'accion': 'aceptar'})
        p.refresh_from_db()
        self.assertEqual(p.ticket_aceptacion, '20260300000999')
        self.assertIsNotNone(p.aceptada_en)
        metodo, url, _ = sunat.llamadas[-1]
        self.assertEqual(metodo, 'POST')
        self.assertTrue(url.endswith('/rce/propuesta/web/registroslibros/202609/aceptarpropuesta'))


class SinCredencialesTest(Base):
    def test_pide_configurar_o_subir_archivo(self):
        Empresa.objects.update(sire_client_id='', sire_client_secret='', sire_usuario_sol='', sunat_client_id='')
        r = self.client.get(reverse('sire_propuestas'))
        self.assertContains(r, 'Ajustes › Empresa')
        r = self.client.post(reverse('sire_propuestas'), {'libro': 'RCE', 'periodo': '2026-09'}, follow=True)
        self.assertContains(r, 'Suba el archivo de la propuesta')
        self.assertFalse(PropuestaSIRE.objects.exists())
        with self.assertRaises(sire_api.ErrorSIRE):
            sire_api.iniciar_sesion('x')
