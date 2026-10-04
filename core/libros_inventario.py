"""Registro del inventario permanente: formato 12.1 (unidades físicas) y 13.1 (valorizado) de SUNAT."""
import re
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render

from .models import D0, Empresa, Kardex, Producto, TIPO_COMPROBANTE, r2
from .utils import _fin_de_mes

# Tabla 5: tipo de existencia
TIPO_EXISTENCIA = {'MERCADERIA': '01', 'PRODUCTO_TERMINADO': '02', 'MATERIA_PRIMA': '03', 'SEMIELABORADO': '02',
                   'SUMINISTRO': '05'}
# Tabla 6: unidad de medida (los códigos del sistema ya son los de SUNAT)
_DOCS = sorted([(nombre, codigo) for codigo, nombre in TIPO_COMPROBANTE] +
               [('Guía de remisión remitente', '09'), ('Guía de remisión transportista', '31')],
               key=lambda x: -len(x[0]))
_SERIE_NUMERO = re.compile(r'\b([A-Z0-9]{4})-(\d{1,8})\b')


def documento(referencia):
    """(tipo tabla 10, serie, número) desde la referencia del kardex; '00' para documentos internos."""
    texto = re.sub(r'^(Reversión|Anulación)\s+', '', referencia or '')
    tipo = '00'
    for nombre, codigo in _DOCS:
        if nombre.lower() in texto.lower():
            tipo = codigo
            break
    m = _SERIE_NUMERO.search(texto)
    return tipo, (m.group(1) if m else ''), (m.group(2) if m else (referencia or '')[:20])


def registro(desde, hasta):
    """[{p, inicial, filas, totales}] de los productos con saldo o movimientos en el rango."""
    productos = {p.pk: p for p in Producto.objects.filter(tipo='BIEN')}
    previos = {}
    for k in Kardex.objects.filter(fecha__lt=desde).order_by('producto_id', 'fecha', 'id').only(
            'producto_id', 'saldo', 'costo_promedio'):
        previos[k.producto_id] = (k.saldo, k.costo_promedio)
    movimientos = {}
    for k in Kardex.objects.filter(fecha__range=[desde, hasta]).order_by('fecha', 'id'):
        movimientos.setdefault(k.producto_id, []).append(k)
    bloques = []
    for pid in sorted(set(previos) | set(movimientos), key=lambda x: productos[x].codigo if x in productos else ''):
        p = productos.get(pid)
        if p is None:
            continue
        cant, costo = previos.get(pid, (D0, D0))
        if not cant and pid not in movimientos:
            continue
        inicial = {'cantidad': cant, 'costo': costo, 'total': r2(cant * costo)}
        filas = []
        tot = {'ent_cant': D0, 'ent_total': D0, 'sal_cant': D0, 'sal_total': D0}
        for k in movimientos.get(pid, []):
            tipo, serie, numero = documento(k.referencia)
            total = r2(k.cantidad * k.costo_unitario)
            entrada = k.tipo == 'ENTRADA'
            fila = {'k': k, 'tipo_doc': tipo, 'serie': serie, 'numero': numero, 'operacion': k.codigo_sunat or '99',
                    'ent_cant': k.cantidad if entrada else D0, 'sal_cant': D0 if entrada else k.cantidad,
                    'ent_costo': k.costo_unitario if entrada else D0, 'sal_costo': D0 if entrada else k.costo_unitario,
                    'ent_total': total if entrada else D0, 'sal_total': D0 if entrada else total,
                    'saldo_cant': k.saldo, 'saldo_costo': k.costo_promedio,
                    'saldo_total': r2(k.saldo * k.costo_promedio)}
            for clave in tot:
                tot[clave] += fila[clave]
            filas.append(fila)
        final = filas[-1] if filas else None
        tot['saldo_cant'] = final['saldo_cant'] if final else cant
        tot['saldo_total'] = final['saldo_total'] if final else inicial['total']
        bloques.append({'p': p, 'tipo_existencia': TIPO_EXISTENCIA.get(p.clase, '99'), 'inicial': inicial,
                        'filas': filas, 'totales': tot})
    return bloques


def excel(bloques, desde, hasta, valorizado):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    empresa = Empresa.actual()
    formato = '13.1' if valorizado else '12.1'
    titulo = ('REGISTRO DE INVENTARIO PERMANENTE VALORIZADO - DETALLE DEL INVENTARIO VALORIZADO' if valorizado else
              'REGISTRO DEL INVENTARIO PERMANENTE EN UNIDADES FÍSICAS - DETALLE DEL INVENTARIO PERMANENTE EN '
              'UNIDADES FÍSICAS')
    wb = Workbook()
    ws = wb.active
    ws.title = f'Formato {formato}'
    negrita = Font(bold=True)
    encabezado = PatternFill('solid', fgColor='1F4E79')
    blanco = Font(bold=True, color='FFFFFF')
    borde = Border(*(Side(style='thin', color='BBBBBB'),) * 4)
    ws.append([f'FORMATO {formato}: {titulo}'])
    ws['A1'].font = Font(bold=True, size=12)
    periodo = (f'{desde:%m/%Y}' if (desde.year, desde.month) == (hasta.year, hasta.month)
               else f'{desde:%d/%m/%Y} al {hasta:%d/%m/%Y}')
    for etiqueta, valor in [('PERIODO:', periodo), ('RUC:', empresa.ruc),
                            ('APELLIDOS Y NOMBRES, DENOMINACIÓN O RAZÓN SOCIAL:', empresa.razon_social),
                            ('ESTABLECIMIENTO (1):', empresa.direccion or 'Todos los almacenes')]:
        ws.append([etiqueta, valor])
        ws.cell(ws.max_row, 1).font = negrita
    if valorizado:
        ws.append(['MÉTODO DE VALUACIÓN:', 'Promedio ponderado móvil'])
        ws.cell(ws.max_row, 1).font = negrita
    cols = (['Fecha', 'Tipo (tabla 10)', 'Serie', 'Número', 'Tipo de operación (tabla 12)'] +
            (['Entradas cant.', 'Entradas costo unit.', 'Entradas costo total', 'Salidas cant.', 'Salidas costo unit.',
              'Salidas costo total', 'Saldo cant.', 'Saldo costo unit.', 'Saldo costo total'] if valorizado else
             ['Entradas', 'Salidas', 'Saldo final']))
    for b in bloques:
        p = b['p']
        ws.append([])
        for etiqueta, valor in [('CÓDIGO DE LA EXISTENCIA:', p.codigo), ('TIPO (TABLA 5):', b['tipo_existencia']),
                                ('DESCRIPCIÓN:', p.nombre), ('CÓDIGO DE LA UNIDAD DE MEDIDA (TABLA 6):', p.unidad)]:
            ws.append([etiqueta, valor])
            ws.cell(ws.max_row, 1).font = negrita
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font, c.fill, c.border = blanco, encabezado, borde
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ini = b['inicial']
        ws.append(['', '00', '', '', '16', ini['cantidad'], ini['costo'], ini['total'], 0, 0, 0, ini['cantidad'],
                   ini['costo'], ini['total']] if valorizado else ['', '00', '', '', '16', ini['cantidad'], 0,
                                                                   ini['cantidad']])
        ws.cell(ws.max_row, 1).value = 'Saldo inicial'
        for f in b['filas']:
            base = [f['k'].fecha.strftime('%d/%m/%Y'), f['tipo_doc'], f['serie'], f['numero'], f['operacion']]
            if valorizado:
                ws.append(base + [f['ent_cant'], f['ent_costo'], f['ent_total'], f['sal_cant'], f['sal_costo'],
                                  f['sal_total'], f['saldo_cant'], f['saldo_costo'], f['saldo_total']])
            else:
                ws.append(base + [f['ent_cant'], f['sal_cant'], f['saldo_cant']])
        t = b['totales']
        if valorizado:
            ws.append(['', '', '', '', 'TOTALES', ini['cantidad'] + t['ent_cant'], '', ini['total'] + t['ent_total'],
                       t['sal_cant'], '', t['sal_total'], t['saldo_cant'], '', t['saldo_total']])
        else:
            ws.append(['', '', '', '', 'TOTALES', ini['cantidad'] + t['ent_cant'], t['sal_cant'], t['saldo_cant']])
        for c in ws[ws.max_row]:
            c.font = negrita
    for fila in ws.iter_rows(min_row=1):
        for c in fila:
            if isinstance(c.value, Decimal):
                c.value = float(c.value)
    for letra, ancho in zip('ABCDEFGHIJKLMN', [16, 12, 8, 12, 14] + [13] * 9):
        ws.column_dimensions[letra].width = ancho
    ws.column_dimensions['A'].width = 28
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="Formato_{formato}_{desde:%Y%m}.xlsx"'
    wb.save(resp)
    return resp


@login_required
def libro(request):
    from .inventario import _sin_permiso_costos
    from .modulos import puede_ver_costos
    valorizado = request.GET.get('formato_sunat', '12.1') == '13.1'
    if valorizado and not puede_ver_costos(request.user):
        return _sin_permiso_costos(request)
    hoy = date.today()
    anio = int(request.GET.get('anio') or hoy.year)
    mes = request.GET.get('mes', str(hoy.month))
    if mes == 'anual':
        desde, hasta = date(anio, 1, 1), date(anio, 12, 31)
    else:
        desde = date(anio, int(mes), 1)
        hasta = _fin_de_mes(desde)
    bloques = registro(desde, hasta)
    if request.GET.get('descargar'):
        return excel(bloques, desde, hasta, valorizado)
    return render(request, 'inventario/libro_sunat.html', {
        'bloques': bloques, 'valorizado': valorizado, 'anio': anio, 'mes': mes, 'desde': desde, 'hasta': hasta,
        'anios': range(hoy.year, 2015, -1), 'meses': [(str(i), f'{i:02d}') for i in range(1, 13)],
        'puede_valorizado': puede_ver_costos(request.user)})
