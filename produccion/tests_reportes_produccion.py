"""Reportes de producción (v1.30): kg/TM con el peso neto, plan contra real, merma, consumo, productividad, OEE,
paradas, atrasos, costo por kg, trazabilidad, calidad, producción-ventas-stock, tablero, Excel y filtros guardados.
Un escenario por clase (las clases grandes hacen caer al SQLite de Python 3.10.0)."""
from datetime import date, time, timedelta
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from contabilidad.models import CentroCosto
from core.models import Almacen, FiltroGuardado, Lote, Producto, Tercero
from planillas.models import Turno

from . import avanzado, servicios
from . import reportes_produccion as rp
from .models import (CentroTrabajo, Equipo, HojaRuta, InspeccionCalidad, ListaMateriales, MotivoParada,
                     OrdenMantenimiento, OrdenProduccion, Parada, VersionFabricacion)

D = Decimal
HOY = timezone.localdate()


def F(**k):
    datos = {'desde': HOY.replace(day=1).isoformat(), 'hasta': HOY.isoformat()}
    datos.update({c: str(v) for c, v in k.items()})
    return rp.Filtros(datos)


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        with patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450'))):
            call_command('seed_demo', verbosity=0)
        cls.user = User.objects.create_superuser('t', 't@t.com', 'x')

    def setUp(self):
        self.client.force_login(self.user)
        p = patch('core.tipo_cambio._consultar', return_value=(D('3.441'), D('3.450')))
        p.start()
        self.addCleanup(p.stop)
        self.almacen = Almacen.principal()
        prov = Tercero.objects.filter(tipo__in=['PROVEEDOR', 'AMBOS']).first()
        self.masa = Producto.objects.create(nombre='Masa', unidad='KGM', clase='MATERIA_PRIMA', proveedor=prov,
                                            familia='Insumos')
        self.empaque = Producto.objects.create(nombre='Empaque', clase='MATERIA_PRIMA', proveedor=prov)
        # paquete de 500 g
        self.galleta = Producto.objects.create(nombre='Galleta paquete', clase='PRODUCTO_TERMINADO', control='LOTE',
                                               precio_venta=D('8'), peso_neto=D('0.5'), familia='Galletas',
                                               vida_util_dias=180)  # v1.84: el lote vence (INV-28)
        cc = CentroCosto.objects.create(codigo='PL', nombre='Planta', tipo='PRODUCCION')
        self.horno = CentroTrabajo.objects.create(codigo='HOR', nombre='Horno', costo_hora_mo=D('20'),
                                                  costo_hora_cif=D('30'), centro_costo=cc, horas_turno=D('8'),
                                                  turnos=1, dias_laborables='1234567')
        self.equipo = Equipo.objects.create(codigo='H1', nombre='Horno 1', centro=self.horno)
        lg = ListaMateriales.objects.create(producto=self.galleta, codigo='G1', cantidad_base=D('1'))
        lg.componentes.create(producto=self.masa, cantidad=D('0.5'))
        lg.componentes.create(producto=self.empaque, cantidad=D('1'))
        rg = HojaRuta.objects.create(codigo='R-GAL', nombre='Horneado', estado='APROBADA')
        rg.operaciones.create(secuencia=10, centro=self.horno, descripcion='Hornear', horas_unidad=D('0.01'))
        self.version = VersionFabricacion.objects.create(producto=self.galleta, codigo='1', lista=lg, hoja=rg,
                                                         lote_costeo=D('100'), dias_fabricacion=2)
        for p, cant, costo in ((self.masa, '500', '6'), (self.empaque, '500', '0.10')):
            p.mover_stock(D(cant), 'Saldo de prueba', costo=D(costo), almacen=self.almacen)

    def orden(self, cantidad='100', fecha=HOY):
        o = OrdenProduccion.objects.create(producto=self.galleta, cantidad=D(cantidad), fecha=fecha,
                                           almacen_insumos=self.almacen, almacen_destino=self.almacen,
                                           lista=self.version.lista, version=self.version)
        servicios.explotar(o)
        servicios.confirmar(o, self.user)
        return o

    def terminar(self, o, producida='100', masa='55'):
        c = o.consumos.get(producto=self.masa)
        h = o.horas.get()
        servicios.terminar(o, self.user, D(producida), {c.pk: D(masa)}, {h.pk: D('2')})
        o.refresh_from_db()
        return o


class ConversionYProduccionTest(Base):
    def test_kg_tm_y_detalle_con_fecha_contable(self):
        self.assertEqual(self.galleta.kg_por_unidad, D('0.5'))
        self.assertEqual(self.masa.kg_por_unidad, D('1'))
        self.assertIsNone(self.empaque.kg_por_unidad)
        self.assertEqual(rp.en(self.galleta, D('100'), 'TM'), D('0.05'))
        o = self.terminar(self.orden())
        d = rp.produccion(F(agrupar='mes', por='producto'))
        self.assertEqual(d['filas'][0]['cantidad'], D('100'))
        self.assertEqual(d['filas'][0]['kg'], D('50.00'))
        self.assertEqual(d['total']['tm'], D('0.050'))
        det = rp.produccion(F(agrupar='orden'))['filas'][0]
        self.assertEqual((det['orden'], det['fecha_real'], det['fecha_contable']), (o.numero, HOY, o.operacion.fecha))
        fam = rp.produccion(F(por='familia', familia='Galletas'))['filas']
        self.assertEqual(fam[0]['dimension'], 'Galletas')
        self.assertEqual(rp.produccion(F(familia='Otra'))['filas'], [])


class PlanRealYAtrasoTest(Base):
    def test_cumplimiento_desvio_y_atrasadas(self):
        self.terminar(self.orden(), producida='90')
        self.orden(fecha=HOY - timedelta(days=5))  # debía terminar hace 4 días
        d = rp.plan_real(F(desde=(HOY - timedelta(days=10)).isoformat(), agrupar='orden'))
        terminada = next(x for x in d['filas'] if x['real'])
        self.assertEqual(terminada['cumplimiento'], D('90.00'))
        self.assertEqual(terminada['alerta'], 'Desvío')
        self.assertEqual(d['total']['plan'], D('100.00'))  # 2 × 100 paquetes × 0.5 kg
        ab = rp.abiertas(F())
        self.assertEqual(ab['filas'][0]['atraso'], 4)
        self.assertEqual(ab['filas'][0]['pendiente'], D('100'))


class MermaConsumoCostoTest(Base):
    def test_merma_ranking_consumo_y_costo_por_kg(self):
        o = self.terminar(self.orden(), masa='55')
        d = rp.merma(F())
        x = d['filas'][0]
        self.assertEqual((x['plan'], x['real'], x['merma'], x['ranking']), (D('50.00'), D('55.00'), D('5.00'), 1))
        self.assertEqual(x['merma_pct'], D('10.00'))
        self.assertEqual(x['soles'], D('30.00'))
        self.assertNotIn('soles', [c for c, _, _ in rp.merma(F(), ver_costos=False)['columnas']])
        c = rp.consumo_mp(F())['filas'][0]
        self.assertEqual((c['real_tm'], c['plan_tm'], c['variacion_kg']), (D('0.055'), D('0.050'), D('5.00')))
        k = rp.costo_kg(F(por='orden'))['filas'][0]
        self.assertEqual(k['kg'], D('50.00'))
        self.assertEqual(k['total'], o.costo_real_final.quantize(D('0.01')))
        self.assertEqual(k['total_tm'], (o.costo_real_final / 50 * 1000).quantize(D('0.01')))


class PlantaTurnoParadasTest(Base):
    def test_productividad_turno_oee_y_paradas(self):
        manana = Turno.objects.create(nombre='Mañana', hora_entrada=time(0, 0), hora_salida=time(23, 59))
        o = self.orden()
        h = o.horas.get()
        h.maquinistas, h.ayudantes = 1, 1
        h.save()
        avanzado.registrar_avance(o, h, D('80'), D('20'), D('2'), 'Ana', '', self.user)
        a = o.avances.get()
        self.assertEqual(a.turno, manana)
        motivo = MotivoParada.objects.get(codigo='AVERIA')
        p = avanzado.registrar_parada(o, h, motivo, 30, self.user)
        self.assertEqual((p.equipo, p.horas), (self.equipo, D('0.50')))
        prod = rp.productividad(F(por='operario'))['filas'][0]
        self.assertEqual((prod['dimension'], prod['kg'], prod['horas_hombre'], prod['kg_hh'], prod['kg_hm']),
                         ('Ana', D('40.00'), D('4.00'), D('10.00'), D('20.00')))
        to = rp.turno_operario(F(turno=manana.pk))['filas']
        self.assertEqual((to[0]['turno'], to[0]['buena']), (str(manana), D('80')))
        oee = rp.oee(F())['filas'][0]
        self.assertEqual(oee['disponibilidad'], D('80.00'))  # 2 h ÷ (2 + 0.5)
        self.assertEqual(oee['calidad'], D('80.00'))  # 80 ÷ 100
        self.assertIsNotNone(oee['oee'])
        OrdenMantenimiento.objects.create(equipo=self.equipo, tipo='CORRECTIVO', estado='CERRADA', fecha_fin=HOY,
                                          horas_parada=D('1.5'), descripcion='Resistencia')
        pm = rp.paradas_mantenimiento(F())['filas'][0]
        self.assertEqual((pm['fallas'], pm['parada'], pm['mttr'], pm['mtbf']),
                         (2, D('2.00'), D('1.00'), D('1.00')))  # 2 h de operación ÷ 2 fallas


class TableroCalidadTest(Base):
    def test_calidad_por_lote_y_tablero(self):
        o = self.terminar(self.orden())
        for resultado in ('APROBADO', 'REPROCESO', 'RECHAZADO', 'APROBADO'):
            i = InspeccionCalidad.objects.create(producto=self.galleta, lote=o.numero, orden=o, resultado=resultado,
                                                 tipo='TERMINADO')
            if resultado != 'APROBADO':
                i.resultados.create(caracteristica='Humedad', conforme=False)
        d = rp.calidad(F())
        x = d['filas'][0]
        self.assertEqual((x['aprobado'], x['reprocesado'], x['rechazado'], x['lotes']),
                         (D('50.00'), D('25.00'), D('25.00'), 1))
        self.assertEqual(d['causas'], [('Humedad', 2)])
        t = rp.tablero(rp.Filtros({'fecha': HOY.isoformat()}))
        self.assertEqual((t['tm_dia'], t['tm_mes']), (D('0.050'), D('0.050')))
        self.assertEqual(t['cumplimiento'], D('100.00'))
        self.assertEqual(t['merma_pct'], D('10.00'))


class TrazabilidadYStockTest(Base):
    def test_lote_hacia_atras_y_produccion_ventas_stock(self):
        o = self.terminar(self.orden())
        lote = Lote.objects.get(producto=self.galleta, codigo=o.numero)
        d = rp.trazabilidad(F(lote=lote.codigo))
        info = d['lotes'][0]
        self.assertEqual(info['orden'], o)
        self.assertEqual({i['producto'] for i in info['insumos']}, {self.masa, self.empaque})
        self.assertEqual(rp.trazabilidad(F(lote='NO-EXISTE'))['avisos'], ['No se encontró el lote.'])
        s = next(x for x in rp.produccion_ventas_stock(F())['filas'] if 'Galleta' in x['producto'])
        self.assertEqual((s['producido'], s['stock'], s['cobertura']), (D('0.050'), D('0.050'), None))


class ReporteDiarioTest(Base):
    def test_parte_diario_secciones_y_excel(self):
        o = self.terminar(self.orden())
        abierta = self.orden()
        h = abierta.horas.get()
        avanzado.registrar_avance(abierta, h, D('30'), D('2'), D('1'), 'Luis', '', self.user)
        avanzado.registrar_parada(abierta, h, MotivoParada.objects.get(codigo='AVERIA'), 30, self.user)
        InspeccionCalidad.objects.create(producto=self.galleta, lote=o.numero, orden=o, resultado='APROBADO',
                                         tipo='TERMINADO')
        d = rp.diario(rp.Filtros({'fecha': HOY.isoformat()}))
        s = {x['titulo']: x for x in d['secciones']}
        self.assertEqual(s['Producción terminada']['filas'][0]['orden'], o.numero)
        self.assertEqual(s['Producción terminada']['total']['tm'], D('0.050'))
        self.assertEqual(s['Avance de planta por turno y operario']['filas'][0]['operario'], 'Luis')
        masa = next(x for x in s['Insumos consumidos']['filas'] if 'Masa' in x['insumo'])
        self.assertEqual((masa['plan'], masa['real'], masa['variacion']), (D('50'), D('55'), D('5')))
        self.assertEqual(s['Paradas de línea']['total']['horas'], D('0.50'))
        self.assertEqual(s['Calidad']['filas'][0]['resultado'], 'Aprobado')
        self.assertEqual(s['Órdenes del día aún abiertas']['filas'][0]['producido'], D('30'))
        self.assertEqual(dict(d['resumen'])['Inspecciones aprobadas'], '1 de 1')
        # otro día: vacío
        ayer = rp.diario(rp.Filtros({'fecha': (HOY - timedelta(days=1)).isoformat()}))
        self.assertEqual(ayer['secciones'][0]['filas'], [])
        r = self.client.get(reverse('manufactura:diario'))
        self.assertContains(r, 'Reporte diario de producción')
        self.assertContains(r, o.numero)
        r = self.client.get(reverse('manufactura:diario'), {'formato': 'excel'})
        from io import BytesIO
        from openpyxl import load_workbook
        hojas = load_workbook(BytesIO(r.content)).sheetnames
        self.assertEqual(hojas[0], 'Resumen')
        self.assertIn('Paradas de línea', hojas)


class LogModificacionesTest(Base):
    def test_solo_admin_y_registra_cambios_de_manufactura(self):
        o = self.orden()
        h = o.horas.get()
        avanzado.registrar_parada(o, h, MotivoParada.objects.get(codigo='LIMPIEZA'), 10, self.user)
        self.equipo.nombre = 'Horno túnel'
        self.equipo.save()
        r = self.client.get(reverse('manufactura:modificaciones'))
        self.assertContains(r, 'Log de modificaciones')
        self.assertContains(r, 'Horno túnel')  # el cambio de nombre del equipo, con el valor anterior
        self.assertContains(r, 'Horno 1')
        self.assertContains(r, 'Parada HOR')
        r = self.client.get(reverse('manufactura:modificaciones'),
                            {'modelo': 'produccion.OrdenProduccion', 'objeto': o.pk, 'volver': '//malo.com'})
        self.assertContains(r, o.numero)
        self.assertNotContains(r, 'Horno túnel')
        self.assertNotContains(r, '//malo.com')
        self.assertIn('spreadsheet', self.client.get(reverse('manufactura:modificaciones'),
                                                     {'formato': 'excel'})['Content-Type'])
        # un usuario de manufactura que no es administrador no lo ve
        from django.contrib.auth.models import Group
        jefe = User.objects.create_user('jefe', password='x')
        jefe.groups.add(Group.objects.get_or_create(name='Manufactura')[0])
        self.client.force_login(jefe)
        self.assertEqual(self.client.get(reverse('manufactura:modificaciones')).status_code, 403)
        self.assertNotContains(self.client.get(reverse('manufactura:ordenes')), 'Log de modificaciones')
        self.assertContains(self.client.get(reverse('manufactura:ordenes')), 'Reporte diario de producción')


class OrdenTerminadaSinTabletTest(Base):
    """Piloto MAN-04 y diagnóstico n.º 16: OEE y productividad con órdenes terminadas desde su pantalla, y merma
    contra la receta de lo PRODUCIDO."""

    def test_oee_sin_tablet_y_merma_sobre_lo_producido(self):
        o = self.terminar(self.orden(), producida='50', masa='27.5')  # 100 planificadas, 50 producidas
        h = o.horas.get()
        self.assertFalse(o.avances.exists())
        Parada.objects.create(centro=self.horno, equipo=self.equipo, orden=o,
                              motivo=MotivoParada.objects.get(codigo='AVERIA'),
                              inicio=timezone.now() - timedelta(hours=1), fin=timezone.now() - timedelta(minutes=30))
        fila = rp.oee(F())['filas'][0]
        self.assertEqual(fila['funcionando'], D('2.00'))  # las horas reales de la orden, no 0
        self.assertEqual(fila['disponibilidad'], D('80.00'))  # 2 ÷ (2 + 0.5)
        prod = rp.productividad(F(por='linea'))['filas'][0]
        self.assertEqual((prod['kg'], prod['horas_maquina']), (D('25.00'), D('2.00')))  # 50 paquetes × 0.5 kg
        self.assertEqual(h.horas_real, D('2'))
        m = rp.merma(F())['filas'][0]
        # receta de 50 paquetes = 25 kg de masa (no 50 de lo planificado); consumo 27.5 → merma 2.5 kg (10 %)
        self.assertEqual((m['plan'], m['real'], m['merma'], m['merma_pct']),
                         (D('25.00'), D('27.50'), D('2.50'), D('10.00')))


class MotivosParadaOEETest(Base):
    """Piloto n.º 9: maestro de motivos de parada con su pérdida OEE, reglas en tablet y pantalla, OEE y Pareto."""

    def parada(self, codigo, minutos, centro=None, equipo=None, hace=60):
        fin = timezone.now() - timedelta(minutes=hace)
        return Parada.objects.create(centro=centro or self.horno, equipo=equipo, motivo=MotivoParada.objects.get(
            codigo=codigo), inicio=fin - timedelta(minutes=minutos), fin=fin)

    def test_maestro_reglas_oee_y_pareto(self):
        from .avanzado import ErrorAvanzado, registrar_parada
        from .models import ConfiguracionPlanta, OrdenMantenimiento
        averia = MotivoParada.objects.get(codigo='AVERIA')  # los 9 anteriores quedaron clasificados
        self.assertEqual((averia.perdida, averia.tipo, averia.es_falla, averia.exige_equipo,
                          averia.genera_mantenimiento), ('DISPONIBILIDAD', 'AVERIA', True, True, True))
        self.assertEqual(MotivoParada.objects.get(codigo='RE-MIC').perdida, 'RENDIMIENTO')
        self.assertEqual(MotivoParada.objects.get(codigo='NP-SPED').perdida, 'NO_PROGRAMADO')
        # tablet: la avería toma el único equipo del puesto y propone la orden de mantenimiento
        o = self.orden()
        h = o.horas.get()
        p = registrar_parada(o, h, averia, 30, self.user)
        self.assertEqual(p.equipo, self.equipo)
        ot = p.orden_mantenimiento
        self.assertEqual((ot.tipo, ot.estado, ot.equipo), ('CORRECTIVO', 'PROGRAMADA', self.equipo))
        Equipo.objects.create(codigo='H2', nombre='Horno 2', centro=self.horno)  # con dos equipos hay que elegir
        with self.assertRaises(ErrorAvanzado):
            registrar_parada(o, h, averia, 10, self.user)
        otro = CentroTrabajo.objects.create(codigo='ENV', nombre='Envasado', costo_hora_mo=D('10'))
        limpieza = MotivoParada.objects.get(codigo='LIMPIEZA')
        limpieza.puestos.set([otro])  # ya no aplica en el horno
        with self.assertRaises(ErrorAvanzado):
            registrar_parada(o, h, limpieza, 10, self.user)
        limpieza.puestos.clear()
        Parada.objects.all().delete()
        OrdenMantenimiento.objects.all().delete()
        # OEE: 2 h reportadas; avería 30 min (disponibilidad), limpieza 60 min (planificada: no resta por defecto),
        # sin pedidos 120 min (fuera del OEE), falta de material 3 min (microparada: rendimiento)
        avanzado.registrar_avance(o, h, D('100'), D('0'), D('2'), 'Ana', '', self.user)
        self.parada('AVERIA', 30, equipo=self.equipo, hace=300)
        self.parada('LIMPIEZA', 60, hace=200)
        self.parada('NP-SPED', 120, hace=100)
        self.parada('FALTA_MAT', 3, hace=10)
        fila = rp.oee(F())['filas'][0]
        self.assertEqual((fila['paradas_no_plan'], fila['paradas_rend'], fila['paradas_plan'], fila['no_programado']),
                         (D('0.50'), D('0.05'), D('1.00'), D('2.00')))
        self.assertEqual(fila['disponibilidad'], D('80.39'))  # 2.05 ÷ (2.05 + 0.5)
        config = ConfiguracionPlanta.actual()
        config.planificadas_restan = True  # ahora la limpieza también baja la disponibilidad
        config.save()
        self.assertEqual(rp.oee(F())['filas'][0]['disponibilidad'], D('57.75'))  # 2.05 ÷ (2.05 + 1.5)
        # Pareto: de mayor a menor, con acumulado
        par = rp.pareto_paradas(F(por='motivo'))
        self.assertEqual([x['grupo'].split()[0] for x in par['filas']], ['NP-SPED', 'LIMPIEZA', 'AVERIA', 'FALTA_MAT'])
        self.assertEqual(par['filas'][-1]['acumulado'], D('100.00'))
        self.assertEqual(rp.pareto_paradas(F(por='area'))['filas'][0]['grupo'], 'Planeamiento')

    def test_pantalla_carga_masiva_y_paradas(self):
        r = self.client.get(reverse('manufactura:motivos_parada'))
        self.assertContains(r, 'Motivos de parada (OEE)')
        self.assertContains(r, 'Avería mecánica')
        self.client.post(reverse('manufactura:motivo_parada_nuevo'), {
            'codigo': 'av-hid', 'nombre': 'Avería hidráulica', 'perdida': 'DISPONIBILIDAD', 'tipo': 'AVERIA',
            'area': 'MANTENIMIENTO', 'es_falla': 'on', 'exige_equipo': 'on', 'genera_mantenimiento': 'on',
            'activo': 'on'})
        self.assertTrue(MotivoParada.objects.filter(codigo='AV-HID', es_falla=True).exists())
        r = self.client.post(reverse('manufactura:motivo_parada_nuevo'), {
            'codigo': 'X1', 'nombre': 'Mal', 'perdida': 'RENDIMIENTO', 'tipo': 'AVERIA', 'area': 'PRODUCCION',
            'activo': 'on'})
        self.assertContains(r, 'microparadas o velocidad reducida')
        self.client.post(reverse('manufactura:motivos_parada'), {'planificadas_restan': 'on',
                                                                 'microparada_minutos': '3'})
        from .models import ConfiguracionPlanta
        self.assertEqual(ConfiguracionPlanta.actual().microparada_minutos, 3)
        # carga masiva
        from core.carga_masiva import ErrorFila, _fila_motivo_parada, cargar
        fila = _fila_motivo_parada({'codigo': 'ex-agu', 'nombre': 'Corte de agua', 'perdida': 'Disponibilidad',
                                    'tipo': 'Externa', 'area': 'Servicios', 'puestos': 'HOR'}, False)
        cargar('motivos_parada', [fila], self.user)
        m = MotivoParada.objects.get(codigo='EX-AGU')
        self.assertEqual((m.tipo, m.area, list(m.puestos.values_list('codigo', flat=True))),
                         ('EXTERNA', 'SERVICIOS', ['HOR']))
        with self.assertRaises(ErrorFila):
            _fila_motivo_parada({'codigo': 'Y', 'nombre': 'Y', 'perdida': 'Rendimiento', 'tipo': 'Avería'}, False)
        # pantalla de paradas: la avería exige equipo y propone la orden de mantenimiento
        inicio = timezone.localtime().replace(second=0, microsecond=0) - timedelta(hours=3)
        datos = {'centro': self.horno.pk, 'motivo': MotivoParada.objects.get(codigo='AVERIA').pk,
                 'inicio': inicio.strftime('%Y-%m-%dT%H:%M'),
                 'fin': (inicio + timedelta(hours=1)).strftime('%Y-%m-%dT%H:%M')}
        self.assertContains(self.client.post(reverse('manufactura:paradas'), datos), 'exige indicar el equipo')
        self.client.post(reverse('manufactura:paradas'), {**datos, 'equipo': self.equipo.pk})
        self.assertIsNotNone(Parada.objects.get().orden_mantenimiento)
        self.assertEqual(self.client.get(reverse('manufactura:reporte', args=['pareto_paradas'])).status_code, 200)


class VistasReportesTest(Base):
    def test_paginas_excel_filtros_y_tablet(self):
        o = self.terminar(self.orden())
        self.assertEqual(self.client.get(reverse('manufactura:reportes')).status_code, 200)
        for clave in rp.REPORTES:
            r = self.client.get(reverse('manufactura:reporte', args=[clave]), {'lote': o.numero})
            self.assertEqual(r.status_code, 200, clave)
        self.assertEqual(self.client.get(reverse('manufactura:tablero')).status_code, 200)
        r = self.client.get(reverse('manufactura:reporte', args=['produccion']), {'formato': 'excel', 'unidad': 'TM'})
        self.assertIn('spreadsheet', r['Content-Type'])
        url = reverse('manufactura:reporte', args=['produccion']) + '?familia=Galletas&agrupar=dia'
        self.client.post(url, {'accion': 'guardar_filtro', 'nombre': 'Galletas diario'})
        g = FiltroGuardado.objects.get(usuario=self.user, reporte='produccion')
        self.assertIn('familia=Galletas', g.parametros)
        self.client.post(url, {'accion': 'borrar_filtro', 'filtro': g.pk})
        self.assertFalse(FiltroGuardado.objects.exists())
        # parada desde la tablet y desde la página de paradas
        o2 = self.orden()
        h = o2.horas.get()
        motivo = MotivoParada.objects.get(codigo='LIMPIEZA')
        self.client.post(reverse('manufactura:planta_orden', args=[o2.pk]),
                         {'accion': 'parada', 'hora': h.pk, 'motivo': motivo.pk, 'minutos': '15'})
        self.assertEqual(Parada.objects.get().horas, D('0.25'))
        inicio = timezone.localtime().replace(second=0, microsecond=0) - timedelta(hours=2)
        r = self.client.post(reverse('manufactura:paradas'), {
            'centro': self.horno.pk, 'equipo': self.equipo.pk, 'motivo': motivo.pk,
            'inicio': inicio.strftime('%Y-%m-%dT%H:%M'), 'fin': (inicio + timedelta(hours=1)).strftime('%Y-%m-%dT%H:%M')})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(Parada.objects.count(), 2)
        self.assertEqual(self.client.get(reverse('manufactura:paradas')).status_code, 200)
