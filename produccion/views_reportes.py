"""Manufactura › Reportes de producción (con Excel, filtros guardados y unidad de reporte) y registro de paradas."""
from datetime import date, datetime, timedelta
from decimal import Decimal
from urllib.parse import urlencode

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import Http404
from django.shortcuts import redirect, render
from django.utils import timezone

from core.forms import BootstrapMixin
from core.models import Almacen, FiltroGuardado
from core.modulos import puede_ver_costos
from core.utils import excel_response

from . import reportes_produccion as rp
from .models import CentroTrabajo, Equipo, MotivoParada, OrdenProduccion, Parada

POR = {'produccion': [('producto', 'Producto'), ('familia', 'Familia'), ('linea', 'Línea / puesto'),
                      ('planta', 'Planta')],
       'productividad': [('linea', 'Línea / puesto'), ('turno', 'Turno'), ('operario', 'Operario')],
       'costo_kg': [('producto', 'Producto'), ('mes', 'Mes'), ('orden', 'Orden')]}
AGRUPAR = {'produccion': rp.AGRUPAR_TIEMPO + [('orden', 'Detalle por orden (fecha real y contable)')],
           'plan_real': [('orden', 'Por orden'), ('linea', 'Por línea'), ('semana', 'Por semana')]}
MAX_FILAS = 2000  # en pantalla; el Excel lleva todo
ESTADOS = {'plan_real': OrdenProduccion.ESTADOS[1:4], 'abiertas': OrdenProduccion.ESTADOS[:3]}


def _formato(valor, tipo):
    if valor is None or valor == '':
        return '—' if tipo != 'texto' else ''
    if isinstance(valor, (date, datetime)):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, Decimal):
        texto = f'{valor:,f}'
        return f'{texto} %' if tipo == 'pct' else texto
    if isinstance(valor, int) and not isinstance(valor, bool):
        return f'{valor:,}'
    return str(valor)


@login_required
def indice(request):
    return render(request, 'produccion/reportes.html', {
        'reportes': [(clave, titulo) for clave, (titulo, _, _) in rp.REPORTES.items()
                     if clave != 'costo_kg' or puede_ver_costos(request.user)]})


def _guardados(request, clave):
    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion == 'guardar_filtro' and request.POST.get('nombre', '').strip():
            parametros = urlencode({k: v for k, v in request.GET.items() if v and k != 'formato'})
            FiltroGuardado.objects.update_or_create(usuario=request.user, reporte=clave,
                                                    nombre=request.POST['nombre'].strip()[:60],
                                                    defaults={'parametros': parametros[:1000]})
            messages.success(request, 'Filtro guardado.')
        elif accion == 'borrar_filtro':
            FiltroGuardado.objects.filter(usuario=request.user, reporte=clave, pk=request.POST.get('filtro')).delete()
            messages.success(request, 'Filtro eliminado.')
        return True
    return False


@login_required
def reporte(request, clave):
    if clave not in rp.REPORTES:
        raise Http404
    titulo, funcion, filtros_usados = rp.REPORTES[clave]
    costos = puede_ver_costos(request.user)
    if clave == 'costo_kg' and not costos:
        raise Http404
    if _guardados(request, clave):
        return redirect(f'{request.path}?{request.GET.urlencode()}')
    f = rp.Filtros(request.GET)
    datos = funcion(f, ver_costos=costos) if clave in ('merma', 'consumo') else funcion(f)
    if request.GET.get('formato') == 'excel' and 'secciones' in datos:
        return _excel_secciones(f'reporte_{clave}_{f.fecha:%Y%m%d}', f'{titulo} {f.fecha:%d/%m/%Y}', datos)
    if request.GET.get('formato') == 'excel' and 'columnas' in datos:
        cols = datos['columnas']
        filas = [[x.get(c) for c, _, _ in cols] for x in datos['filas']]
        if datos.get('total'):
            filas.append([('TOTAL' if i == 0 else datos['total'].get(c)) for i, (c, _, _) in enumerate(cols)])
        return excel_response(f'reporte_{clave}', f'{titulo} {f.desde:%d/%m/%Y} - {f.hasta:%d/%m/%Y}',
                              [t for _, t, _ in cols], filas)
    contexto = {
        'clave': clave, 'titulo': titulo, 'd': datos, 'f': f, 'usados': filtros_usados,
        'guardados': FiltroGuardado.objects.filter(usuario=request.user, reporte=clave),
        'plantas': Almacen.objects.filter(activo=True), 'lineas': CentroTrabajo.objects.filter(activo=True),
        'equipos': Equipo.objects.filter(activo=True), 'por': POR.get(clave), 'agrupar': AGRUPAR.get(clave),
        'estados': ESTADOS.get(clave), 'unidades': rp.UNIDADES, 'turnos': _turnos(),
        'familias': sorted(set(rp.Producto.objects.exclude(familia='').values_list('familia', flat=True))),
        'producto_sel': rp.Producto.objects.filter(pk=f.producto).first() if f.producto else None,
        'insumo_sel': rp.Producto.objects.filter(pk=f.insumo).first() if f.insumo else None,
        'cliente_sel': _tercero(f.cliente)}
    if 'columnas' in datos:
        filas = datos['filas']
        if len(filas) > MAX_FILAS:
            datos['avisos'] = list(datos.get('avisos') or []) + [
                f'Se muestran {MAX_FILAS:,} de {len(filas):,} filas (los totales son de todas); exporte a Excel para '
                f'verlas todas o acote los filtros.']
            filas = filas[:MAX_FILAS]
        contexto['tabla'] = [[(_formato(x.get(c), t), t) for c, _, t in datos['columnas']] for x in filas]
        if datos.get('total'):
            contexto['fila_total'] = [(_formato(datos['total'].get(c), t) if c in datos['total'] else '', t)
                                      for c, _, t in datos['columnas']]
    if 'secciones' in datos:
        otros = urlencode({k: v for k, v in request.GET.items() if v and k not in ('fecha', 'formato')})
        contexto.update({'dia_anterior': f.fecha - timedelta(days=1), 'dia_siguiente': f.fecha + timedelta(days=1),
                         'consulta_dia': f'{otros}&' if otros else ''})
        contexto['secciones'] = [{
            'titulo': s['titulo'], 'encabezados': [(t, tipo) for _, t, tipo in s['columnas']], 'n': len(s['filas']),
            'tabla': [[(_formato(x.get(c), t), t) for c, _, t in s['columnas']] for x in s['filas']],
            'total': ([(_formato(s['total'].get(c), t) if c in s['total'] else '', t) for c, _, t in s['columnas']]
                      if s.get('total') and s['filas'] else None)} for s in datos['secciones']]
    if datos.get('grafico'):
        maximo = max((v for _, v in datos['grafico']), default=Decimal('0')) or Decimal('1')
        contexto['barras'] = [(e, v, int(v * 100 / maximo)) for e, v in datos['grafico']]
    plantilla = {'trazabilidad': 'produccion/reporte_trazabilidad.html', 'diario': 'produccion/reporte_diario.html',
                 'tablero': 'produccion/reporte_tablero.html'}.get(clave, 'produccion/reporte.html')
    return render(request, plantilla, contexto)


def _excel_secciones(nombre, titulo, datos):
    """Excel con el resumen en la primera hoja y una hoja por sección."""
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    ws.append([titulo])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    for etiqueta, valor in datos.get('resumen', []):
        ws.append([etiqueta, float(valor) if isinstance(valor, Decimal) else valor])
    ws.column_dimensions['A'].width = 30
    for s in datos['secciones']:
        hoja = wb.create_sheet(''.join(c for c in s['titulo'] if c not in '[]:*?/\\')[:30])
        hoja.append([s['titulo']])
        hoja['A1'].font = Font(bold=True, size=12)
        hoja.append([t for _, t, _ in s['columnas']])
        for celda in hoja[2]:
            celda.font = Font(bold=True, color='FFFFFF')
            celda.fill = PatternFill('solid', fgColor='1F4E79')
        filas = [[x.get(c) for c, _, _ in s['columnas']] for x in s['filas']]
        if s.get('total') and filas:
            filas.append([('TOTAL' if i == 0 else s['total'].get(c)) for i, (c, _, _) in enumerate(s['columnas'])])
        for fila in filas:
            hoja.append([float(v) if isinstance(v, Decimal) else v for v in fila])
        for col in hoja.columns:
            ancho = max((len(str(c.value or '')) for c in col[1:]), default=10)
            hoja.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 10), 45)
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{nombre}.xlsx"'
    wb.save(resp)
    return resp


def _tercero(pk):
    from core.models import Tercero
    return Tercero.objects.filter(pk=pk).first() if pk else None


def _turnos():
    from planillas.models import Turno
    return Turno.objects.all()


# ---------------------------------------------------------------- log de modificaciones
@login_required
def modificaciones(request):
    """Quién creó, cambió, anuló o eliminó cada registro de Manufactura (valor anterior y nuevo). Solo el
    administrador."""
    if not request.user.is_superuser:
        return render(request, 'core/sin_acceso.html',
                      {'motivo': 'Solo un administrador puede ver el log de modificaciones.'}, status=403)
    from core.views import consulta_bitacora
    volver = request.GET.get('volver', '')
    volver = volver if volver.startswith('/') and not volver.startswith('//') else None  # solo rutas internas
    return consulta_bitacora(request, app='produccion', titulo='Log de modificaciones', volver=volver)


# ---------------------------------------------------------------- paradas
class ParadaForm(BootstrapMixin, forms.ModelForm):
    class Meta:
        model = Parada
        fields = ['centro', 'equipo', 'orden', 'motivo', 'inicio', 'fin', 'operario', 'nota']
        widgets = {'inicio': forms.DateTimeInput(attrs={'type': 'datetime-local'}, format='%Y-%m-%dT%H:%M'),
                   'fin': forms.DateTimeInput(attrs={'type': 'datetime-local'}, format='%Y-%m-%dT%H:%M')}

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['centro'].queryset = CentroTrabajo.objects.filter(activo=True)
        self.fields['equipo'].queryset = Equipo.objects.filter(activo=True)
        desde = timezone.localdate().replace(day=1) - timedelta(days=62)
        self.fields['orden'].queryset = OrdenProduccion.objects.filter(
            estado__in=['CONFIRMADA', 'EN_PROCESO', 'TERMINADA'], fecha__gte=desde).order_by('-id')
        self.fields['motivo'].queryset = MotivoParada.objects.filter(activo=True)

    def clean(self):
        d = super().clean()
        if d.get('inicio') and d.get('fin') and d['fin'] <= d['inicio']:
            self.add_error('fin', 'Debe ser posterior al inicio.')
        if d.get('equipo') and d.get('centro') and d['equipo'].centro_id not in (None, d['centro'].pk):
            self.add_error('equipo', 'El equipo es de otro puesto.')
        return d


@login_required
def paradas(request):
    form = ParadaForm(request.POST or None, initial={'inicio': timezone.localtime().replace(second=0, microsecond=0)})
    if request.method == 'POST' and form.is_valid():
        p = form.save(commit=False)
        p.usuario = request.user
        p.save()
        messages.success(request, 'Parada registrada.')
        return redirect('manufactura:paradas')
    return render(request, 'produccion/paradas.html', {
        'form': form, 'paradas': Parada.objects.select_related('centro', 'equipo', 'motivo', 'orden')[:150]})
