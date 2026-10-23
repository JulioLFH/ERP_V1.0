"""Reparto de gastos entre centros de costo: reglas, cifras de reparto (m², trabajadores…), simulación, ejecución,
reversión y resultado por obra o centro con los gastos asignados."""
from decimal import Decimal, InvalidOperation

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, redirect, render

from core.forms import BootstrapMixin
from core.utils import excel_response

from . import reparto as rp
from .models import CentroCosto, CifraReparto, EjecucionReparto, PeriodoContable, ReglaReparto, ValorCifra
from .views import _mes_input, _periodo, al_dia


class ReglaRepartoForm(BootstrapMixin, forms.ModelForm):
    porcentajes_texto = forms.CharField(
        label='Porcentajes fijos', required=False, widget=forms.Textarea(attrs={'rows': 3}),
        help_text='Solo con la base «porcentaje fijo»: un centro por línea, «código = %» (ej. OB0001 = 60)')

    class Meta:
        model = ReglaReparto
        fields = ['orden', 'nombre', 'origenes', 'incluir_dependientes', 'cuentas', 'base', 'cifra', 'tipo_destino',
                  'centro_padre', 'destinos', 'porcentajes_texto', 'cuenta_asignacion', 'desde', 'hasta', 'activo']
        widgets = {'origenes': forms.SelectMultiple(attrs={'size': 6}),
                   'destinos': forms.SelectMultiple(attrs={'size': 6})}

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        activos = CentroCosto.objects.filter(activo=True)
        for campo in ('origenes', 'destinos', 'centro_padre'):
            self.fields[campo].queryset = activos
        self.fields['cifra'].queryset = CifraReparto.objects.filter(activo=True)
        if self.instance.pk and self.instance.porcentajes:
            codigos = dict(CentroCosto.objects.values_list('pk', 'codigo'))
            self.initial['porcentajes_texto'] = '\n'.join(
                f'{codigos.get(int(c), c)} = {p}' for c, p in self.instance.porcentajes.items())

    def clean(self):
        data = super().clean()
        base, destino = data.get('base'), data.get('tipo_destino')
        if base == 'CIFRA' and not data.get('cifra'):
            self.add_error('cifra', 'Elija la cifra de reparto (m², trabajadores…).')
        if destino == 'HIJOS' and not data.get('centro_padre'):
            self.add_error('centro_padre', 'Elija el centro del que dependen los destinos (ej. OBRAS).')
        if destino == 'LISTA' and not data.get('destinos'):
            self.add_error('destinos', 'Elija los centros de destino.')
        cuenta = data.get('cuenta_asignacion')
        if cuenta and not cuenta.codigo.startswith('6'):
            self.add_error('cuenta_asignacion', 'Use una cuenta de gasto (clase 6) para que el reparto en cascada no '
                                                'cuente dos veces.')
        for campo in ('desde', 'hasta'):
            valor = (data.get(campo) or '').replace('-', '').replace('/', '')
            if valor and not (len(valor) == 6 and valor.isdigit()):
                self.add_error(campo, 'Periodo AAAAMM (ej. 202611).')
            data[campo] = valor
        porcentajes = {}
        if base == 'FIJO':
            por_codigo = {c.codigo.upper(): c.pk for c in CentroCosto.objects.all()}
            for linea in (data.get('porcentajes_texto') or '').splitlines():
                if not linea.strip():
                    continue
                codigo, _, pct = linea.partition('=')
                pk = por_codigo.get(codigo.strip().upper())
                try:
                    valor = Decimal(pct.strip().rstrip('%'))
                except InvalidOperation:
                    valor = None
                if pk is None or valor is None or valor <= 0:
                    self.add_error('porcentajes_texto', f'Revise «{linea.strip()}»: código de centro = porcentaje.')
                    continue
                porcentajes[str(pk)] = str(valor)
            if not porcentajes:
                self.add_error('porcentajes_texto', 'Indique el % de cada centro de destino.')
            if destino != 'LISTA':
                self.add_error('tipo_destino', 'Con porcentajes fijos, elija «Los centros de una lista».')
        data['porcentajes'] = porcentajes
        return data

    def save(self, commit=True):
        self.instance.porcentajes = self.cleaned_data.get('porcentajes', {})
        return super().save(commit)


class CifraForm(BootstrapMixin, forms.ModelForm):
    class Meta:
        model = CifraReparto
        fields = ['codigo', 'nombre', 'unidad', 'activo']


def _bloqueado(periodo):
    return PeriodoContable.esta_cerrado(periodo)


@login_required
@al_dia
def reparto(request):
    """Reglas del reparto, simulación del periodo, ejecución y reversión."""
    periodo = _periodo(request, 'mes') if request.method == 'GET' else (request.POST.get('periodo') or '')[:6]
    if request.method == 'POST':
        accion = request.POST.get('accion')
        try:
            if accion == 'ejecutar':
                _, n = rp.ejecutar(periodo, request.user)
                messages.success(request, f'Reparto ejecutado: {n} asiento(s). Se recalcula al centralizar el periodo.')
            elif accion == 'revertir':
                n = rp.revertir(periodo, request.user, request.POST.get('motivo', ''))
                messages.success(request, f'Reparto revertido ({n} asiento(s) anulados). Corrija las bases y vuelva a '
                                          f'ejecutarlo.')
        except rp.ErrorReparto as exc:
            messages.error(request, str(exc))
        return redirect(f'{request.path}?mes={_mes_input(periodo)}')
    filas = rp.calcular(periodo)
    if request.GET.get('formato') == 'excel':
        datos = []
        for f in filas:
            for d in f['destinos']:
                datos.append([f['regla'].nombre, f['regla'].get_base_display(), d['centro'].codigo, d['centro'].nombre,
                              d['base'], d['pct'], d['monto']])
        return excel_response(f'Reparto_{periodo}', f'REPARTO DE GASTOS {periodo[4:]}/{periodo[:4]} (simulación)',
                              ['Regla', 'Base', 'Centro', 'Nombre', 'Base del centro', '%', 'Asignado S/'], datos)
    return render(request, 'contabilidad/reparto.html', {
        'periodo': periodo, 'mes': _mes_input(periodo), 'filas': filas, 'reglas': ReglaReparto.objects.all(),
        'ejecucion': EjecucionReparto.objects.filter(periodo=periodo).first(), 'cerrado': _bloqueado(periodo),
        'total': sum((f['total'] for f in filas if f['destinos']), Decimal('0'))})


@login_required
def regla(request, pk=None):
    obj = get_object_or_404(ReglaReparto, pk=pk) if pk else None
    form = ReglaRepartoForm(request.POST or None, instance=obj)
    if request.method == 'POST' and request.POST.get('eliminar') and obj:
        obj.delete()
        messages.success(request, 'Regla eliminada (los repartos ya ejecutados se recalculan sin ella al centralizar).')
        return redirect('contabilidad:reparto')
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, 'Regla de reparto guardada.')
        return redirect('contabilidad:reparto')
    return render(request, 'contabilidad/reparto_regla.html', {'form': form, 'obj': obj})


@login_required
def cifras(request):
    """Cifras de reparto y sus valores por periodo y centro (a mano o desde Excel: Centro, Valor)."""
    periodo = _periodo(request, 'mes') if request.method == 'GET' else (request.POST.get('periodo') or '')[:6]
    cifra = CifraReparto.objects.filter(pk=request.GET.get('cifra') or request.POST.get('cifra')).first() \
        or CifraReparto.objects.filter(activo=True).first()
    form = CifraForm(request.POST if request.POST.get('accion') == 'nueva' else None)
    centros = list(CentroCosto.objects.filter(activo=True).order_by('codigo'))
    if request.method == 'POST':
        accion = request.POST.get('accion')
        if accion == 'nueva':
            if form.is_valid():
                cifra = form.save()
                messages.success(request, f'Cifra {cifra} creada: cargue sus valores por centro.')
            else:
                messages.error(request, 'Revise el código y el nombre de la cifra.')
        elif cifra and accion in ('valores', 'excel'):
            valores = {}
            if accion == 'excel':
                valores = _valores_excel(request.FILES.get('archivo'), centros, request)
            else:
                for c in centros:
                    texto = (request.POST.get(f'valor_{c.pk}') or '').strip().replace(',', '')
                    if texto:
                        try:
                            valores[c.pk] = Decimal(texto)
                        except InvalidOperation:
                            messages.error(request, f'{c.codigo}: «{texto}» no es un número.')
            for pk, valor in valores.items():
                ValorCifra.objects.update_or_create(cifra=cifra, periodo=periodo, centro_id=pk,
                                                    defaults={'valor': valor})
            if valores:
                messages.success(request, f'{len(valores)} valor(es) de {cifra} para {periodo[4:]}/{periodo[:4]}.')
                if EjecucionReparto.objects.filter(periodo=periodo).exists():
                    PeriodoContable.objects.filter(periodo=periodo).update(pendiente=True)  # se recalcula
        return redirect(f'{request.path}?mes={_mes_input(periodo)}&cifra={cifra.pk if cifra else ""}')
    actuales = {}
    if cifra:
        for v in ValorCifra.objects.filter(cifra=cifra, periodo__lte=periodo).order_by('periodo'):
            actuales[v.centro_id] = (v.valor, v.periodo)
    filas = [{'centro': c, 'valor': actuales.get(c.pk, (None, ''))[0], 'desde': actuales.get(c.pk, (None, ''))[1]}
             for c in centros]
    return render(request, 'contabilidad/reparto_cifras.html', {
        'cifras': CifraReparto.objects.all(), 'cifra': cifra, 'form': form, 'filas': filas, 'periodo': periodo,
        'mes': _mes_input(periodo)})


def _valores_excel(archivo, centros, request):
    """{centro: valor} de un Excel con columnas Centro (código) y Valor."""
    if archivo is None:
        messages.error(request, 'Adjunte el Excel (columnas Centro y Valor).')
        return {}
    from openpyxl import load_workbook
    try:
        hoja = load_workbook(archivo, data_only=True, read_only=True).active
    except Exception:
        messages.error(request, 'No se pudo leer el Excel (.xlsx).')
        return {}
    por_codigo = {c.codigo.upper(): c.pk for c in centros}
    valores, errores = {}, []
    for n, fila in enumerate(hoja.iter_rows(values_only=True), 1):
        if not fila or fila[0] is None or n == 1 and str(fila[0]).strip().lower().startswith('centro'):
            continue
        pk = por_codigo.get(str(fila[0]).strip().upper())
        try:
            valor = Decimal(str(fila[1]).replace(',', '')) if len(fila) > 1 and fila[1] is not None else None
        except InvalidOperation:
            valor = None
        if pk is None or valor is None:
            errores.append(f'fila {n}')
            continue
        valores[pk] = valor
    if errores:
        messages.warning(request, f'Filas omitidas (centro desconocido o valor vacío): {", ".join(errores[:10])}.')
    return valores


@login_required
@al_dia
def resultado(request):
    hasta = _periodo(request, 'hasta')
    desde = _periodo(request, 'desde', defecto=f'{hasta[:4]}01')
    solo_obras = request.GET.get('obras') == '1' or getattr(request, 'modulo_actual', '') == 'obras'
    filas, totales = rp.resultado(desde, hasta, solo_obras=solo_obras)
    if request.GET.get('formato') == 'excel':
        datos = [[f['centro'].codigo, f['centro'].nombre, f['ventas'], f['propios'], f['asignados'], f['resultado'],
                  f['gg_presupuesto']] for f in filas]
        return excel_response(f'Resultado_centros_{desde}_{hasta}', 'RESULTADO POR CENTRO DE COSTO CON GASTOS ASIGNADOS',
                              ['Centro', 'Nombre', 'Ventas S/', 'Gastos propios S/', 'Gastos asignados S/',
                               'Resultado S/', 'GG del presupuesto S/'], datos)
    return render(request, 'contabilidad/reparto_resultado.html', {
        'filas': filas, 'totales': totales, 'desde': _mes_input(desde), 'hasta': _mes_input(hasta),
        'solo_obras': solo_obras})
