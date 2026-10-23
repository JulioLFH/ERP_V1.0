"""Importación de presupuestos: el Excel que exporta S10 (hoja de presupuesto + hoja de análisis), el reporte del visor
del S10 (Excel 97-2000 «solo datos», sin fila de títulos), el CSV de análisis del S10 (.chr) o la plantilla de Ceiba.

S10 · Presupuesto:  Ítem | Descripción | Und. | Metrado | Precio S/. | Parcial S/.  (títulos sin unidad ni metrado)
S10 · Reporte:      sin títulos de columnas y con columnas que cambian según el nivel: se reconoce por la forma del
                    ítem (01 / 01.01 / 01.01.01) y por las celdas con datos de cada fila. «Subpresupuesto 001 …» abre
                    un subpresupuesto; si un ítem se repite sin ese encabezado, empieza otro (cada uno numera desde 01)
S10 · Análisis:     bloques «Partida 01.01.01 TRAZO…» · «Rendimiento m2/DIA MO. 500.0000 EQ. 500.0000» ·
                    encabezado «Código | Descripción Recurso | Unidad | Cuadrilla | Cantidad | Precio | Parcial» ·
                    grupos «Mano de Obra», «Materiales», «Equipos», «Subcontratos», «Subpartidas»; los bloques
                    «Subpartida 9001… CONCRETO…» son sub-análisis
S10 · CSV de APU:   una línea por recurso con la cabecera de su partida repetida (ítem, descripción, unidad,
                    rendimiento…, código, recurso, unidad, cantidad, precio, parcial), sin cuadrilla
Plantilla Ceiba:    hoja Presupuesto (Subpresupuesto opcional, Ítem, Descripción, Unidad, Metrado, Precio), hoja APU
                    (Ítem, Rendimiento, Jornada, Tipo, Código, Recurso, Unidad, Cuadrilla, Cantidad, Precio) y hoja
                    Subanálisis (Sub-análisis, Descripción sub-análisis, Und. sub-análisis, Rendimiento, Jornada, Tipo,
                    Código, Recurso, Unidad, Cuadrilla, Cantidad, Precio).
Los recursos se buscan por su CÓDIGO: un código nuevo crea el recurso aunque haya otro con el mismo nombre.
"""
import csv
import io
import re
import unicodedata
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

ITEM = re.compile(r'^\d{1,3}(\.\d{1,4})*\.?$')
ITEM_S10 = re.compile(r'^\d{2,3}(\.\d{2,3})*$')  # sin títulos de columnas: ítems con ceros (01, 01.01)
CODIGO_RECURSO = re.compile(r'^\d{6,16}$')  # el S10 usa 10, 12 y 14 dígitos
GRUPOS = {'mano de obra': 'MO', 'materiales': 'MAT', 'material': 'MAT', 'equipos': 'EQ', 'equipo': 'EQ',
          'equipos y herramientas': 'EQ', 'subcontratos': 'SUB', 'subcontrato': 'SUB', 'servicios': 'SUB',
          'subpartidas': 'MAT', 'subpartida': 'MAT'}
TIPO_POR_CODIGO = {'01': 'MO', '02': 'MAT', '03': 'EQ', '04': 'SUB'}  # codificación habitual de S10
MONEDA = {'s/.', 's/', '$', 'us$', 'us$.'}
C2, C4 = Decimal('0.01'), Decimal('0.0001')


class ErrorImportacion(Exception):
    pass


def _norm(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', texto).strip(' .:')


def _num(valor):
    if valor is None or valor == '' or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace(',', '')
    if not texto or not re.fullmatch(r'-?\d+(\.\d+)?', texto):
        return None
    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _item(valor):
    t = _texto(valor).rstrip('.')
    return t if t and ITEM.match(t) else ''


def _r2(v):
    return Decimal(v).quantize(C2, rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------- lectura de archivos
def _contenido(archivo):
    archivo.seek(0)
    return archivo.read()


def _hojas(archivo):
    """[(nombre, filas)] del Excel: .xlsx (openpyxl) o .xls de Excel 97-2003 (xlrd), según su contenido."""
    datos = _contenido(archivo)
    if datos[:4] == b'\xd0\xcf\x11\xe0':  # Excel 97-2003 (.xls), aunque el archivo se llame .xlsx
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - xlrd está en requirements.txt
            raise ErrorImportacion('Falta el lector de Excel antiguo (xlrd): guarde el archivo como .xlsx.') from exc
        try:
            libro = xlrd.open_workbook(file_contents=datos)
        except Exception as exc:
            raise ErrorImportacion(f'No se pudo leer el Excel (.xls): {exc}') from exc
        vacias = (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_DATE, xlrd.XL_CELL_ERROR)
        return [(h.name, [[None if c.ctype in vacias else c.value for c in h.row(i)] for i in range(h.nrows)])
                for h in libro.sheets()]
    from openpyxl import load_workbook
    try:
        libro = load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    except Exception as exc:
        raise ErrorImportacion(f'No se pudo leer el Excel: {exc}') from exc
    return [(h.title, [list(f) for f in h.iter_rows(values_only=True)]) for h in libro.worksheets]


def _encabezado(filas, *requeridos, limite=40):
    """(índice de la fila, {campo: columna}) de la primera fila que tiene todas las columnas requeridas."""
    for i, fila in enumerate(filas[:limite]):
        celdas = [_norm(c) for c in fila]
        cols = {}
        for j, c in enumerate(celdas):
            for campo, pruebas in _CAMPOS.items():
                if campo not in cols and c and any(c == p or c.startswith(p) for p in pruebas):
                    cols[campo] = j
                    break
        if all(r in cols for r in requeridos):
            return i, cols
    return None, None


_CAMPOS = {
    'sub': ('subpresupuesto', 'sub presupuesto', 'subpres'),
    'item': ('item', 'it.'),
    'descripcion': ('descripcion', 'descripción', 'recurso', 'nombre', 'concepto'),
    'unidad': ('und', 'unidad', 'u.m', 'um'),
    'metrado': ('metrado', 'cantidad presupuestada'),
    'cantidad': ('cantidad', 'aporte'),
    'cuadrilla': ('cuadrilla', 'cuad'),
    'precio': ('precio', 'p.u', 'costo unitario', 'pu'),
    'parcial': ('parcial', 'total', 'importe'),
    'codigo': ('codigo', 'cod'),
    'rendimiento': ('rendimiento', 'rend'),
    'jornada': ('jornada',),
    'tipo': ('tipo',),
}


def _celda(fila, cols, campo):
    j = cols.get(campo)
    return fila[j] if j is not None and j < len(fila) else None


def _subpresupuesto(fila):
    """(código, nombre) si la fila es el encabezado de un subpresupuesto («Subpresupuesto 001 ESTRUCTURAS»)."""
    textos = [_texto(c) for c in fila if _texto(c)]
    for k, t in enumerate(textos):
        if not _norm(t).startswith(('subpresupuesto', 'sub presupuesto')):
            continue
        resto = re.sub(r'^\s*sub\s*-?\s*presupuesto\s*:?', '', t, flags=re.I).strip()
        partes = ([resto] if resto else []) + textos[k + 1:]
        unido = ' '.join(partes).strip()
        m = re.match(r'(\d{2,4})\s*[-–:]?\s*(.*)', unido)
        if m:
            return m.group(1), (m.group(2) or f'Subpresupuesto {m.group(1)}').strip()
        return None
    return None


class _Armador:
    """Junta las partidas con su subpresupuesto: el encabezado «Subpresupuesto …» lo abre y un ítem repetido sin
    encabezado empieza otro (S10 numera cada subpresupuesto desde 01)."""

    def __init__(self):
        self.subs, self.partidas, self.actual, self.vistos = [], [], '', set()

    def sub(self, codigo, nombre, auto=False):
        if not any(s['codigo'] == codigo for s in self.subs):
            self.subs.append({'codigo': codigo, 'nombre': nombre[:250], 'gg': 'gastos generales' in _norm(nombre),
                              'auto': auto})
        self.actual = codigo
        self.vistos = {p['item'] for p in self.partidas if p['sub'] == codigo}

    def partida(self, p):
        if p['item'] in self.vistos:
            if not self.subs:  # lo leído hasta aquí es el primer subpresupuesto
                self.subs.append({'codigo': '001', 'nombre': 'Subpresupuesto 001', 'gg': False, 'auto': True})
                for x in self.partidas:
                    x['sub'] = '001'
            n = len(self.subs) + 1
            while any(s['codigo'] == f'{n:03d}' for s in self.subs):
                n += 1
            self.sub(f'{n:03d}', f'Subpresupuesto {n:03d}', auto=True)
        p['sub'] = p.get('sub') or self.actual
        self.vistos.add(p['item'])
        self.partidas.append(p)

    def fin(self):
        if self.subs:
            for p in self.partidas:
                p['sub'] = p['sub'] or self.subs[0]['codigo']
        for p in self.partidas:
            p['clave'] = f"{p['sub']}|{p['item']}" if p['sub'] else p['item']
        # una partida con hijos es título aunque traiga unidad
        items = {(p['sub'], p['item']) for p in self.partidas}
        for p in self.partidas:
            if any(s == p['sub'] and i.startswith(p['item'] + '.') for s, i in items):
                p['titulo'] = True
        # subpresupuesto sin encabezado (el reporte del S10 no lo trae): toma el nombre de su primer título
        for s in self.subs:
            if s.get('auto'):
                titulo = next((p['descripcion'] for p in self.partidas if p['sub'] == s['codigo'] and p['titulo']
                               and '.' not in p['item']), '')
                if titulo:
                    s['nombre'], s['gg'] = titulo[:250], 'gastos generales' in _norm(titulo)
        return self.subs, self.partidas


def _partida(item, descripcion, unidad, metrado, precio, parcial, titulo, sub=''):
    return {'sub': sub, 'item': item, 'descripcion': descripcion[:400], 'unidad': unidad[:12],
            'metrado': str(metrado or 0), 'precio': str(precio or 0), 'parcial': str(parcial or 0),
            'titulo': titulo, 'precio_archivo': bool(precio)}


# ---------------------------------------------------------------- presupuesto (partidas)
def leer_presupuesto(filas):
    """Hoja con títulos de columnas (S10 o plantilla de Ceiba): (subpresupuestos, partidas) o None."""
    i, cols = _encabezado(filas, 'item', 'descripcion', 'unidad')
    if i is None or ('metrado' not in cols and 'cantidad' not in cols):
        return None
    metrado = 'metrado' if 'metrado' in cols else 'cantidad'
    armador = _Armador()
    for fila in filas[i + 1:]:
        sub = _subpresupuesto(fila)
        if sub:
            armador.sub(*sub)
            continue
        item = _item(_celda(fila, cols, 'item'))
        descripcion = _texto(_celda(fila, cols, 'descripcion'))
        if not item or not descripcion:
            continue
        codigo_sub = _texto(_celda(fila, cols, 'sub'))
        if codigo_sub and codigo_sub != armador.actual:
            armador.sub(codigo_sub, f'Subpresupuesto {codigo_sub}')
        unidad = _texto(_celda(fila, cols, 'unidad'))
        cant = _num(_celda(fila, cols, metrado))
        armador.partida(_partida(item, descripcion, unidad, cant, _num(_celda(fila, cols, 'precio')),
                                 _num(_celda(fila, cols, 'parcial')), not unidad and not cant))
    subs, partidas = armador.fin()
    return (subs, partidas) if partidas else None


def _fila_sin_encabezado(fila):
    """Partida o título de una fila del reporte del S10 (sin títulos de columnas), o None."""
    celdas = [(j, c) for j, c in enumerate(fila) if _texto(c)]
    if not celdas or isinstance(celdas[0][1], (int, float)):  # un número suelto (página, metrado) no es un ítem
        return None
    item = _texto(celdas[0][1]).rstrip('.')  # en el reporte el ítem es la primera celda con datos de la fila
    if not ITEM_S10.match(item):
        return None
    resto = celdas[1:]
    textos = [(j, _texto(c)) for j, c in resto if _num(c) is None and _norm(c) not in MONEDA]
    if not textos or len(textos[0][1]) < 2 or _norm(textos[0][1]).startswith(('pagina', 'fecha')):
        return None
    jd, descripcion = textos[0]
    despues = [(j, c) for j, c in resto if j > jd]
    unidad = next((_texto(c) for j, c in despues if _num(c) is None and _norm(c) not in MONEDA
                   and len(_texto(c)) <= 8), '')
    nums = [_num(c) for j, c in despues if _num(c) is not None]
    if unidad and len(nums) >= 2:
        metrado, parcial = nums[0], nums[-1]
        precio = nums[1] if len(nums) >= 3 else (_r2(parcial / metrado) if metrado else Decimal('0'))
        return _partida(item, descripcion, unidad, metrado, precio, parcial, False)
    return _partida(item, descripcion, '', 0, 0, nums[-1] if nums else 0, True)


def leer_presupuesto_reporte(filas):
    """Reporte del visor del S10 (sin títulos de columnas): (subpresupuestos, partidas) o None."""
    armador = _Armador()
    for fila in filas:
        sub = _subpresupuesto(fila)
        if sub:
            armador.sub(*sub)
            continue
        p = _fila_sin_encabezado(fila)
        if p:
            armador.partida(p)
    subs, partidas = armador.fin()
    if sum(1 for p in partidas if not p['titulo']) < 1:
        return None
    return subs, partidas


# ---------------------------------------------------------------- análisis de precios unitarios
def leer_apu_tabla(filas):
    """Plantilla de Ceiba: una fila por recurso con el ítem de su partida."""
    i, cols = _encabezado(filas, 'item', 'codigo', 'cantidad')
    if i is None:
        return None
    apus = {}
    for fila in filas[i + 1:]:
        item = _item(_celda(fila, cols, 'item'))
        nombre = _texto(_celda(fila, cols, 'descripcion'))
        if not item or not nombre:
            continue
        sub = _texto(_celda(fila, cols, 'sub'))
        a = apus.setdefault(f'{sub}|{item}' if sub else item,
                            {'rendimiento': '0', 'jornada': '8', 'unidad': '', 'recursos': []})
        if _num(_celda(fila, cols, 'rendimiento')):
            a['rendimiento'] = str(_num(_celda(fila, cols, 'rendimiento')))
        if _num(_celda(fila, cols, 'jornada')):
            a['jornada'] = str(_num(_celda(fila, cols, 'jornada')))
        codigo = _texto(_celda(fila, cols, 'codigo'))
        unidad = _texto(_celda(fila, cols, 'unidad'))
        a['recursos'].append(_recurso(codigo, nombre, unidad, _tipo(_texto(_celda(fila, cols, 'tipo')), codigo, unidad),
                                      _num(_celda(fila, cols, 'cuadrilla')), _num(_celda(fila, cols, 'cantidad')),
                                      _num(_celda(fila, cols, 'precio')), _num(_celda(fila, cols, 'parcial'))))
    return apus


_SUBANALISIS = {'sub-analisis': 'codigo_sub', 'subanalisis': 'codigo_sub', 'subpartida': 'codigo_sub',
                'descripcion sub-analisis': 'nombre_sub', 'descripcion subanalisis': 'nombre_sub',
                'und. sub-analisis': 'unidad_sub', 'und sub-analisis': 'unidad_sub', 'unidad sub-analisis': 'unidad_sub',
                'rendimiento': 'rendimiento', 'jornada': 'jornada', 'tipo': 'tipo', 'codigo': 'codigo',
                'recurso': 'nombre', 'unidad': 'unidad', 'cuadrilla': 'cuadrilla', 'cantidad': 'cantidad',
                'precio': 'precio', 'parcial': 'parcial'}


def leer_subanalisis_tabla(filas):
    """Hoja «Subanálisis» de la plantilla: una fila por componente con el código de su sub-análisis."""
    for i, fila in enumerate(filas[:20]):
        cols = {}
        for j, c in enumerate(fila):
            campo = _SUBANALISIS.get(_norm(c))
            if campo and campo not in cols:
                cols[campo] = j
        if 'codigo_sub' in cols and 'codigo' in cols and 'cantidad' in cols:
            break
    else:
        return None
    subs = {}
    for fila in filas[i + 1:]:
        padre, codigo, nombre = (_texto(_celda(fila, cols, c)) for c in ('codigo_sub', 'codigo', 'nombre'))
        if not padre or not codigo or not nombre:
            continue
        s = subs.setdefault(padre, {'nombre': _texto(_celda(fila, cols, 'nombre_sub')) or padre,
                                    'unidad': _texto(_celda(fila, cols, 'unidad_sub')) or 'und',
                                    'tipo': 'MAT', 'rendimiento': '0', 'jornada': '8', 'recursos': []})
        for campo in ('rendimiento', 'jornada'):
            if _num(_celda(fila, cols, campo)):
                s[campo] = str(_num(_celda(fila, cols, campo)))
        unidad = _texto(_celda(fila, cols, 'unidad'))
        s['recursos'].append(_recurso(codigo, nombre, unidad, _tipo(_texto(_celda(fila, cols, 'tipo')), codigo, unidad),
                                      _num(_celda(fila, cols, 'cuadrilla')), _num(_celda(fila, cols, 'cantidad')),
                                      _num(_celda(fila, cols, 'precio')), _num(_celda(fila, cols, 'parcial'))))
    return subs or None


def _tipo(texto, codigo, unidad):
    t = _norm(texto)
    if t in ('mo', 'mano de obra'):
        return 'MO'
    if t in ('mat', 'material', 'materiales'):
        return 'MAT'
    if t in ('eq', 'equipo', 'equipos', 'equipos y herramientas'):
        return 'EQ'
    if t in ('sub', 'subcontrato', 'subcontratos', 'servicios'):
        return 'SUB'
    u = _norm(unidad).replace(' ', '')
    if u == 'hh':
        return 'MO'
    if u in ('hm', '%mo', '%m.o'):
        return 'EQ'
    return TIPO_POR_CODIGO.get(codigo[:2], 'MAT')


def _recurso(codigo, nombre, unidad, tipo, cuadrilla, cantidad, precio, parcial):
    return {'codigo': codigo[:20], 'nombre': nombre[:250], 'unidad': unidad[:12] or 'und', 'tipo': tipo,
            'cuadrilla': str(cuadrilla) if cuadrilla is not None else None, 'cantidad': str(cantidad or 0),
            'precio': str(precio or 0), 'parcial': str(parcial or 0)}


def leer_apu_bloques(filas):
    """Formato de S10: bloques por partida (y por subpartida = sub-análisis). Devuelve (apus, subanálisis)."""
    apus, subanalisis = {}, {}
    actual, tipo, cols = None, None, None
    sub, vistos, n_auto = '', set(), 1
    for fila in filas:
        textos = [_texto(c) for c in fila]
        normas = [_norm(t) for t in textos]
        no_vacias = [(j, t) for j, t in enumerate(textos) if t]
        if not no_vacias:
            continue
        primera = normas[no_vacias[0][0]]
        encabezado_sub = _subpresupuesto(fila)
        if encabezado_sub:
            sub, vistos = encabezado_sub[0], set()
            continue
        if primera.startswith('subpartida') or primera.startswith('sub partida'):
            unido = ' '.join(t for _, t in no_vacias)
            m = re.search(r'(\d{4,12})\s*(.*)', re.sub(r'^\s*sub\s*partida\s*:?', '', unido, flags=re.I))
            if m:
                actual = subanalisis.setdefault(m.group(1), {'nombre': (m.group(2) or m.group(1))[:250], 'unidad': 'und',
                                                             'tipo': 'MAT', 'rendimiento': '0', 'jornada': '8',
                                                             'recursos': []})
                tipo, cols = None, None
            continue
        if primera.startswith('partida'):
            item = next((_item(t) for t in textos if _item(t)), '')
            if not item:  # «Partida 01.01.01 TRAZO…» en una sola celda
                m = re.search(r'(\d{1,3}(?:\.\d{1,4})+)', no_vacias[0][1])
                item = m.group(1) if m else ''
            if item:
                if not sub and item in vistos:  # sin encabezado de subpresupuesto: el ítem repetido abre otro
                    n_auto, vistos = n_auto + 1, set()
                vistos.add(item)
                clave = f'{sub}|{item}' if sub else (f'#{n_auto}|{item}' if n_auto > 1 else item)
                actual = apus.setdefault(clave, {'rendimiento': '0', 'jornada': '8', 'unidad': '', 'recursos': []})
                tipo, cols = None, None
            continue
        if actual is None:
            continue
        if primera.startswith('rendimiento'):
            linea = ' '.join(textos)
            m = re.search(r'M\.?O\.?\s*[:=]?\s*([\d.,]+)', linea, re.I)
            numeros = [_num(c) for c in fila if _num(c) is not None]
            rendimiento = _num(m.group(1)) if m else (numeros[0] if numeros else None)
            if rendimiento:
                actual['rendimiento'] = str(rendimiento)
            u = re.search(r'([A-Za-z0-9³²]+)\s*/\s*D[IÍ]A', linea, re.I)
            if u:
                actual['unidad'] = u.group(1)
            j = re.search(r'jornada\s*[:=]?\s*([\d.,]+)', linea, re.I)
            if j:
                actual['jornada'] = str(_num(j.group(1)))
            continue
        if any(n.startswith('codigo') for n in normas) and any(n.startswith('cantidad') for n in normas):
            _, cols = _encabezado([fila], 'codigo', 'cantidad', limite=1)
            continue
        grupo = next((GRUPOS[n] for n in normas if n in GRUPOS), None)
        if grupo and len(no_vacias) <= 2:
            tipo = grupo
            continue
        if cols is None:
            continue
        codigo = _texto(_celda(fila, cols, 'codigo'))
        nombre = _texto(_celda(fila, cols, 'descripcion'))
        cantidad = _num(_celda(fila, cols, 'cantidad'))
        if not codigo or not nombre or cantidad is None or not re.search(r'\d', codigo):
            continue
        unidad = _texto(_celda(fila, cols, 'unidad'))
        actual['recursos'].append(_recurso(
            codigo, nombre, unidad, tipo or _tipo('', codigo, unidad), _num(_celda(fila, cols, 'cuadrilla')), cantidad,
            _num(_celda(fila, cols, 'precio')), _num(_celda(fila, cols, 'parcial'))))
    return (apus or None), subanalisis


def _cantidad_desde_parcial(cantidad, precio, parcial, rendimiento=None):
    """El CSV del S10 trae las cantidades a 2 decimales (0.06): se recupera el aporte de 4 decimales que da el parcial
    (0.0625 × 70 = 4.38). Entre los que cumplen, gana el que sale de una cuadrilla redonda con el rendimiento
    (1 vibrador / 16 m3 por día = 0.0625) y luego el más cercano a parcial ÷ precio."""
    if not (precio and parcial) or precio <= 0:
        return cantidad
    if cantidad is not None and _r2(cantidad * precio) == parcial:  # la cantidad mostrada ya da el parcial (0.40 × 5.81)
        return cantidad
    desde = ((parcial - Decimal('0.005')) / precio).quantize(C4, rounding='ROUND_CEILING')
    hasta = ((parcial + Decimal('0.005')) / precio).quantize(C4, rounding='ROUND_FLOOR')
    if hasta < desde or (hasta - desde) / C4 > 5000:
        return cantidad
    candidatos = []
    v = desde
    while v <= hasta:
        if _r2(v * precio) == parcial and (cantidad is None or _r2(v) == _r2(cantidad)):
            candidatos.append(v)
        v += C4
    if not candidatos:
        return cantidad
    exacto = parcial / precio

    def redonda(v):  # la cuadrilla (aporte × rendimiento, o ÷ jornada de 8 h) con 2 decimales
        if not rendimiento:
            return 1
        return 0 if any(x == x.quantize(C2) for x in (v * rendimiento, v * rendimiento / 8)) else 1
    return min(candidatos, key=lambda v: (redonda(v), abs(v - exacto)))


def _fila_apu_csv(celdas):
    """(sub, ítem, cabecera de la partida, recurso) de una línea del CSV de análisis del S10, o None."""
    k = next((n for n, c in enumerate(celdas) if ITEM_S10.match(c.rstrip('.'))), None)
    if k is None:
        return None
    r = next((n for n in range(k + 1, len(celdas)) if CODIGO_RECURSO.match(celdas[n])), None)
    if r is None:
        return None
    sub = celdas[k - 1] if k > 0 and re.fullmatch(r'\d{3}', celdas[k - 1]) else ''
    partida = [c for c in celdas[k + 1:r] if c]
    textos = [c for c in partida if _num(c) is None]
    numeros = [_num(c) for c in partida if _num(c) is not None]
    recurso = [c for c in celdas[r + 1:] if c]
    nombres = [c for c in recurso if _num(c) is None]
    if not nombres:
        return None
    nombre = nombres[0]
    unidad = next((c for c in nombres[1:] if len(c) <= 8 and _norm(c) not in GRUPOS), 'und')
    nums = [_num(c) for c in recurso[recurso.index(nombre) + 1:] if _num(c) is not None]
    if len(nums) >= 3:
        cantidad, precio, parcial = nums[-3], nums[-2], nums[-1]
    elif len(nums) == 2:
        cantidad, precio = nums
        parcial = _r2(cantidad * precio)
    else:
        return None
    grupo = next((GRUPOS[_norm(c)] for c in celdas if _norm(c) in GRUPOS), None)
    codigo = celdas[r]
    tipo = grupo or _tipo('', codigo, unidad)
    if not _es_pct(unidad):  # en las herramientas (%mo) la cantidad es el porcentaje
        cantidad = _cantidad_desde_parcial(cantidad, precio, parcial)
    cabecera = {'descripcion': textos[0] if textos else '',
                'unidad': next((t for t in textos[1:] if len(t) <= 8), ''),
                'rendimiento': str(numeros[0]) if numeros else '0'}
    return sub, celdas[k].rstrip('.'), cabecera, _recurso(codigo, nombre, unidad, tipo, None, cantidad, precio, parcial)


GRUPOS_CSV = {'materiales': 'MAT', 'material': 'MAT', 'equipos': 'EQ', 'equipo': 'EQ', 'mano de obra': 'MO',
              'subcontratos': 'SUB', 'subcontrato': 'SUB', 'sub-analisis': 'SUBANALISIS', 'subanalisis': 'SUBANALISIS',
              'sub analisis': 'SUBANALISIS', 'subpartidas': 'SUBANALISIS', 'subpartida': 'SUBANALISIS'}


def _tras(celdas, normas, etiqueta):
    """El valor que sigue a la etiqueta en la fila («Rendimiento», 10.00)."""
    if etiqueta not in normas:
        return ''
    j = normas.index(etiqueta) + 1
    return celdas[j] if j < len(celdas) else ''


def _recurso_reporte(celdas, normas, r, grupo, rendimiento):
    """Recurso de una fila del reporte de APU del S10: [grupo,] código, recurso, recurso, unidad, cantidad, precio, '',
    parcial, [subtotales…]. Devuelve (recurso, grupo vigente)."""
    if r > 0 and normas[r - 1] in GRUPOS_CSV:  # el grupo solo viene en el primer recurso de cada grupo
        grupo = GRUPOS_CSV[normas[r - 1]]
    resto = celdas[r + 1:]
    textos = []
    for c in resto:
        if _num(c) is not None:
            break
        if c:
            textos.append(c)
    if not textos:
        return None, grupo
    nombre, unidad = textos[0], (textos[-1] if len(textos) > 1 else 'und')
    nums = [_num(c) for c in resto[resto.index(textos[-1]) + 1:] if _num(c) is not None]
    if len(nums) < 2:
        return None, grupo
    cantidad, precio = nums[0], nums[1]
    parcial = nums[2] if len(nums) > 2 else _r2(cantidad * precio)
    codigo = celdas[r]
    es_sub = grupo == 'SUBANALISIS'
    tipo = 'MAT' if es_sub else (grupo or _tipo('', codigo, unidad))
    if not _es_pct(unidad):  # en las herramientas (%mo) la cantidad es el porcentaje
        cantidad = _cantidad_desde_parcial(cantidad, precio, parcial, rendimiento)
    recurso = _recurso(codigo, nombre, unidad, tipo, None, cantidad, precio, parcial)
    if es_sub:
        recurso['subanalisis'] = True
    return recurso, grupo


def leer_apu_csv(archivo):
    """CSV de análisis de precios unitarios que exporta el S10 (.chr o .csv, separado por comas, en cp1252). Dos formas:
    el reporte («Partida», ítem, «Rendimiento», …, «Presupuesto», subpresupuesto, «Descripción», … en la cabecera de
    cada partida y luego sus recursos, con el grupo solo en el primero de cada grupo) o una línea por recurso con la
    cabecera de su partida repetida."""
    datos = _contenido(archivo)
    for codificacion in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            texto = datos.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialecto = csv.Sniffer().sniff(texto[:5000], delimiters=',;|\t')
    except csv.Error:
        dialecto = csv.excel
    apus, vistos, n_auto, anterior = {}, set(), 1, None
    actual, grupo, rendimiento = None, None, None
    for fila in csv.reader(io.StringIO(texto), dialecto):
        celdas = [c.strip() for c in fila]
        normas = [_norm(c) for c in celdas]
        # reporte: cabecera de la partida
        if 'partida' in normas and ITEM.match(_tras(celdas, normas, 'partida').rstrip('.') or 'x'):
            item = _tras(celdas, normas, 'partida').rstrip('.')
            sub = _tras(celdas, normas, 'presupuesto')
            rendimiento = _num(_tras(celdas, normas, 'rendimiento'))
            actual = apus.setdefault(f'@{sub}|{item}' if sub else item, {
                'rendimiento': str(rendimiento or 0), 'jornada': '8', 'unidad': _tras(celdas, normas, 'unidad'),
                'descripcion': _tras(celdas, normas, 'descripcion'), 'recursos': []})
            grupo = None
            continue
        r = next((n for n, c in enumerate(celdas) if CODIGO_RECURSO.match(c)), None)
        if actual is not None and r is not None and not any(ITEM_S10.match(c.rstrip('.')) for c in celdas[:r]):
            recurso, grupo = _recurso_reporte(celdas, normas, r, grupo, rendimiento)
            if recurso:
                actual['recursos'].append(recurso)
            costo = _num(_tras(celdas, normas, 'costo directo (en s/.)'))
            if costo is not None:
                actual['costo'] = str(costo)
            continue
        r = _fila_apu_csv(celdas)
        if r is None:
            continue
        sub, item, cabecera, recurso = r
        if not sub and item != anterior and item in vistos:  # el ítem vuelve a aparecer: otro subpresupuesto
            n_auto, vistos = n_auto + 1, set()
        vistos.add(item)
        anterior = item
        clave = f'{sub}|{item}' if sub else (f'#{n_auto}|{item}' if n_auto > 1 else item)
        a = apus.setdefault(clave, {'rendimiento': cabecera['rendimiento'], 'jornada': '8',
                                    'unidad': cabecera['unidad'], 'recursos': []})
        a['recursos'].append(recurso)
    apus = {k: a for k, a in apus.items() if a['recursos']}
    if not apus:
        raise ErrorImportacion('No se encontraron análisis en el CSV: se espera el reporte de análisis de precios '
                               'unitarios del S10 (cabeceras «Partida» con sus recursos: código de 6 a 16 dígitos, '
                               'unidad, cantidad, precio y parcial).')
    return apus


def _resolver_apu(apus, partidas, subs):
    """Une cada análisis con su partida: por subpresupuesto e ítem; sin subpresupuesto en el análisis, por el orden
    en que aparecen (el 2.° 01.01 es del 2.° subpresupuesto). Los del reporte del S10 traen el nombre del
    subpresupuesto («@ESTRUCTURAS|01.01.01»): se unen por la descripción de la partida, luego por ese nombre y si no, por
    el orden en que aparecen; un subpresupuesto que no tenía nombre toma el del análisis."""
    claves = {p['clave'] for p in partidas}
    codigos = [s['codigo'] for s in subs]
    por_codigo = {s['codigo']: s for s in subs}
    nombres_csv = []
    unidos = {}
    for clave, a in apus.items():
        if clave in claves:
            unidos[clave] = a
            continue
        sub, _, item = clave.rpartition('|')
        destino = None
        if sub.startswith('#'):
            n = int(sub[1:]) - 1
            destino = f'{codigos[n]}|{item}' if n < len(codigos) else None
        elif sub.startswith('@'):
            nombre = sub[1:]
            if nombre not in nombres_csv:
                nombres_csv.append(nombre)
            hojas = [p for p in partidas if p['item'] == item and not p['titulo']]
            descripcion = _norm(a.get('descripcion'))
            elegida = (next((p for p in hojas if descripcion and _norm(p['descripcion']) == descripcion), None)
                       or next((p for p in hojas if _norm(por_codigo.get(p['sub'], {}).get('nombre')) == _norm(nombre)),
                               None))
            if elegida is None and hojas:
                n = nombres_csv.index(nombre)
                elegida = next((p for p in hojas if codigos and p['sub'] == codigos[min(n, len(codigos) - 1)]),
                               hojas[0])
            if elegida is not None:
                destino = elegida['clave']
                s = por_codigo.get(elegida['sub'])
                if s is not None and s.get('auto') and nombre:  # el reporte no traía el nombre del subpresupuesto
                    s['nombre'], s['auto'] = nombre[:250], False
                    s['gg'] = s['gg'] or 'gastos generales' in _norm(nombre)
        elif not sub and codigos:
            destino = f'{codigos[0]}|{item}'
        if destino in claves:
            unidos[destino] = a
    return unidos


def leer(archivo, analisis=None):
    """{'subpresupuestos', 'partidas', 'apu': {clave: {...}}, 'subanalisis': {código: {...}}, 'avisos', 'formato'}."""
    leido, formato, apus, subanalisis, avisos = None, '', {}, {}, []
    hojas = _hojas(archivo)
    for nombre, filas in hojas:
        if leido is None:
            leido = leer_presupuesto(filas)
            if leido:
                formato = 'Excel con títulos de columnas'
                continue
        s = leer_subanalisis_tabla(filas)
        if s:
            subanalisis.update(s)
            continue
        a = leer_apu_tabla(filas)
        if not a:
            a, s = leer_apu_bloques(filas)
            subanalisis.update(s)
        if a:
            apus.update(a)
    if leido is None:  # reporte del visor del S10: sin títulos de columnas
        for nombre, filas in hojas:
            leido = leer_presupuesto_reporte(filas)
            if leido:
                formato = 'Reporte del S10 (sin títulos de columnas)'
                break
    if not leido:
        raise ErrorImportacion('No se encontró el presupuesto: ni una hoja con las columnas Ítem, Descripción, Unidad y '
                               'Metrado, ni filas con ítems 01 / 01.01 con su unidad, metrado, precio y parcial.')
    if analisis:
        apus.update(leer_apu_csv(analisis))
        formato += ' + análisis en CSV'
    subs, partidas = leido
    apus = _resolver_apu(apus, partidas, subs)
    claves = {p['clave'] for p in partidas}
    sin_apu = [p['item'] for p in partidas if not p['titulo'] and p['clave'] not in apus]
    if sin_apu:
        avisos.append(f'{len(sin_apu)} partida(s) sin análisis: se importan con su precio del archivo (sin APU).')
    if len(subs) > 1:
        avisos.append(f'{len(subs)} subpresupuestos: ' + '; '.join(f'{s["codigo"]} {s["nombre"]}' for s in subs))
    if subanalisis:
        avisos.append(f'{len(subanalisis)} sub-análisis (se explotan en la relación de insumos).')
    return {'subpresupuestos': subs, 'partidas': partidas, 'apu': {c: a for c, a in apus.items() if c in claves},
            'subanalisis': subanalisis, 'avisos': avisos, 'formato': formato}


# ---------------------------------------------------------------- grabar
def _es_pct(unidad):
    return _norm(unidad).replace(' ', '') in ('%mo', '%m.o')


def _codigos(datos):
    codigos = {r['codigo'] for a in datos['apu'].values() for r in a['recursos'] if r['codigo']}
    for codigo, s in datos.get('subanalisis', {}).items():
        codigos.add(codigo)
        codigos |= {r['codigo'] for r in s['recursos'] if r['codigo']}
    return codigos


def grabar(datos, obra, tipo, nombre, usuario=None, origen='Excel'):
    """Crea el presupuesto con sus subpresupuestos, partidas, análisis, sub-análisis, recursos (los nuevos) y precios;
    devuelve (presupuesto, diferencias [(ítem, precio del archivo, precio calculado)]). Los avisos de la grabación
    quedan en presupuesto.avisos_importacion."""
    from django.db import transaction

    from . import calculos
    from .models import ComponenteRecurso, Partida, PrecioPresupuesto, Presupuesto, Recurso, RecursoAPU, Subpresupuesto
    avisos = []
    with transaction.atomic():
        pres = Presupuesto.objects.create(obra=obra, tipo=tipo, nombre=nombre[:250], origen=f'Importado de {origen}'[:40],
                                          gastos_generales_pct=obra.gastos_generales_pct,
                                          utilidad_pct=obra.utilidad_pct)
        subs = {s['codigo']: Subpresupuesto.objects.create(presupuesto=pres, codigo=s['codigo'], nombre=s['nombre'],
                                                           es_gastos_generales=s['gg'])
                for s in datos.get('subpresupuestos', [])}
        por_codigo = {r.codigo: r for r in Recurso.objects.filter(codigo__in=_codigos(datos))}
        sin_codigo = {(_norm(r['nombre']), _norm(r['unidad'])) for a in datos['apu'].values() for r in a['recursos']
                      if not r['codigo']}
        por_nombre = {(_norm(r.nombre), _norm(r.unidad)): r for r in Recurso.objects.all()} if sin_codigo else {}
        catalogo, precios = {}, {}

        def recurso(r, **extra):
            """El recurso del catálogo: por código (si no existe se crea aunque otro tenga el mismo nombre); sin código,
            por nombre y unidad."""
            clave = r['codigo'] or f'{r["nombre"]}|{r["unidad"]}'
            if clave in catalogo:
                return catalogo[clave]
            obj = por_codigo.get(r['codigo']) if r['codigo'] else por_nombre.get((_norm(r['nombre']), _norm(r['unidad'])))
            if obj is None:
                obj = Recurso.objects.create(codigo=r['codigo'], nombre=r['nombre'], unidad=r['unidad'] or 'und',
                                             tipo=r.get('tipo') or 'MAT',
                                             precio=Decimal('0') if _es_pct(r['unidad']) else Decimal(r.get('precio', '0')),
                                             **extra)
                if r['codigo']:
                    por_codigo[r['codigo']] = obj
            if r.get('subanalisis') and not obj.es_subanalisis:  # grupo «Sub-Análisis» del reporte del S10
                obj.es_subanalisis = True
                obj.save(update_fields=['es_subanalisis'])
            catalogo[clave] = obj
            if not obj.es_porcentaje_mo and 'precio' in r:
                precios.setdefault(obj.pk, Decimal(r['precio']))
            return obj

        # sub-análisis: el recurso con su análisis (catálogo de la empresa)
        for codigo, s in datos.get('subanalisis', {}).items():
            padre = recurso({'codigo': codigo, 'nombre': s['nombre'], 'unidad': s['unidad'], 'tipo': s['tipo']})
            padre.es_subanalisis = True
            padre.rendimiento, padre.jornada = Decimal(s['rendimiento']), Decimal(s['jornada'])
            padre.save(update_fields=['es_subanalisis', 'rendimiento', 'jornada'])
            nuevos = [(recurso(r), r) for r in s['recursos']]
            antes = list(padre.componentes.values_list('recurso_id', 'cantidad'))
            despues = [(obj.pk, Decimal(r['cantidad'])) for obj, r in nuevos]
            if antes and sorted(antes) != sorted(despues):
                avisos.append(f'Se actualizó el análisis del sub-análisis {padre.codigo} {padre.nombre} con el del archivo.')
            padre.componentes.all().delete()
            ComponenteRecurso.objects.bulk_create([ComponenteRecurso(
                padre=padre, recurso=obj, cuadrilla=Decimal(r['cuadrilla']) if r['cuadrilla'] is not None else None,
                cantidad=Decimal(r['cantidad']), orden=n) for n, (obj, r) in enumerate(nuevos)])
        for a in datos['apu'].values():
            for r in a['recursos']:
                recurso(r)
        PrecioPresupuesto.objects.bulk_create([PrecioPresupuesto(presupuesto=pres, recurso_id=k, precio=v)
                                               for k, v in precios.items()])
        for p in datos['partidas']:
            a = datos['apu'].get(p['clave'])
            partida = Partida.objects.create(
                presupuesto=pres, subpresupuesto=subs.get(p['sub']), item=p['item'], descripcion=p['descripcion'],
                unidad=p['unidad'] or (a or {}).get('unidad', ''), metrado=Decimal(p['metrado']), es_titulo=p['titulo'],
                precio_fijo=not p['titulo'] and not (a and a['recursos']), precio=Decimal(p['precio']),
                rendimiento=Decimal(a['rendimiento']) if a else Decimal('0'),
                jornada=Decimal(a['jornada']) if a else Decimal('8'))
            if a and not p['titulo']:
                RecursoAPU.objects.bulk_create([RecursoAPU(
                    partida=partida, recurso=recurso(r),
                    cuadrilla=Decimal(r['cuadrilla']) if r['cuadrilla'] is not None else None,
                    cantidad=Decimal(r['cantidad']), orden=n) for n, r in enumerate(a['recursos'])])
        calculos.recalcular(pres)
    # solo se compara cuando el archivo trae el precio unitario de la partida
    archivo = {(p['sub'], p['item']): Decimal(p['precio']) for p in datos['partidas']
               if not p['titulo'] and p.get('precio_archivo', True) and Decimal(p['precio'])}
    codigo_sub = {s.pk: s.codigo for s in pres.subpresupuestos.all()}
    diferencias = []
    for x in pres.partidas.filter(es_titulo=False):
        clave = (codigo_sub.get(x.subpresupuesto_id, ''), x.item)
        if clave in archivo and abs(archivo[clave] - x.precio) > Decimal('0.01'):
            diferencias.append((x.item, archivo[clave], x.precio))
    pres.avisos_importacion = avisos
    return pres, diferencias


def resumen(datos):
    """Para la vista previa: conteos, el total del archivo y cómo se emparejan los recursos (por código)."""
    from .models import Recurso
    hojas = [p for p in datos['partidas'] if not p['titulo']]
    nombres = {}
    for a in datos['apu'].values():
        for r in a['recursos']:
            nombres.setdefault(r['codigo'] or r['nombre'], r)
    for codigo, s in datos.get('subanalisis', {}).items():
        nombres.setdefault(codigo, {'codigo': codigo, 'nombre': s['nombre'], 'unidad': s['unidad']})
        for r in s['recursos']:
            nombres.setdefault(r['codigo'] or r['nombre'], r)
    codigos = [c for c, r in nombres.items() if r['codigo']]
    existentes = set(Recurso.objects.filter(codigo__in=codigos).values_list('codigo', flat=True))
    nuevos = [r for c, r in nombres.items() if r['codigo'] and c not in existentes]
    mismo_nombre = []
    if nuevos:
        otros = {}
        for codigo, nombre in Recurso.objects.filter(nombre__in=[r['nombre'] for r in nuevos]).values_list(
                'codigo', 'nombre'):
            otros.setdefault(_norm(nombre), codigo)
        mismo_nombre = [(r['codigo'], r['nombre'], otros[_norm(r['nombre'])]) for r in nuevos
                        if _norm(r['nombre']) in otros]
    total = sum((Decimal(p['parcial']) for p in hojas), Decimal('0'))
    gg = {s['codigo'] for s in datos.get('subpresupuestos', []) if s['gg']}
    return {'titulos': len(datos['partidas']) - len(hojas), 'partidas': len(hojas), 'apu': len(datos['apu']),
            'recursos': len(nombres), 'existentes': len(existentes), 'nuevos': len(nuevos),
            'mismo_nombre': mismo_nombre, 'subanalisis': len(datos.get('subanalisis', {})),
            'total_archivo': sum((Decimal(p['parcial']) for p in hojas if p['sub'] not in gg), Decimal('0')),
            'gg_archivo': sum((Decimal(p['parcial']) for p in hojas if p['sub'] in gg), Decimal('0')) if gg else None,
            'total_hojas': total}


def plantilla():
    """Excel de ejemplo con el formato de la plantilla de Ceiba."""
    from openpyxl import Workbook
    libro = Workbook()
    h = libro.active
    h.title = 'Presupuesto'
    h.append(['Ítem', 'Descripción', 'Unidad', 'Metrado', 'Precio', 'Parcial'])
    for fila in [['01', 'OBRAS PROVISIONALES', '', '', '', ''],
                 ['01.01', 'TRAZO, NIVELES Y REPLANTEO', 'm2', 250, '', ''],
                 ['02', 'CONCRETO SIMPLE', '', '', '', ''],
                 ['02.01', 'CONCRETO f\'c=210 kg/cm2', 'm3', 35.5, '', '']]:
        h.append(fila)
    a = libro.create_sheet('APU')
    a.append(['Ítem', 'Rendimiento', 'Jornada', 'Tipo', 'Código', 'Recurso', 'Unidad', 'Cuadrilla', 'Cantidad',
              'Precio'])
    for fila in [['01.01', 500, 8, 'MO', '0147010004', 'PEON', 'hh', 2, '', 19.85],
                 ['01.01', 500, 8, 'MO', '0147010003', 'TOPOGRAFO', 'hh', 1, '', 25.50],
                 ['01.01', 500, 8, 'MAT', '0229060005', 'YESO BOLSA 28 kg', 'bls', '', 0.02, 9.50],
                 ['01.01', 500, 8, 'EQ', '0337010001', 'HERRAMIENTAS MANUALES', '%mo', '', 3, ''],
                 ['02.01', 25, 8, 'MO', '0147010002', 'OPERARIO', 'hh', 2, '', 26.10],
                 ['02.01', 25, 8, 'MO', '0147010004', 'PEON', 'hh', 8, '', 19.85],
                 ['02.01', 25, 8, 'MAT', '0221000001', 'CEMENTO PORTLAND TIPO I (42.5 kg)', 'bls', '', 9.73, 28.50],
                 ['02.01', 25, 8, 'MAT', '0238000000', 'HORMIGON', 'm3', '', 1.05, 65.00],
                 ['02.01', 25, 8, 'EQ', '0349100007', 'MEZCLADORA DE CONCRETO 9-11 p3', 'hm', 1, '', 15.00]]:
        a.append(fila)
    s = libro.create_sheet('Subanálisis')
    s.append(['Sub-análisis', 'Descripción sub-análisis', 'Und. sub-análisis', 'Rendimiento', 'Jornada', 'Tipo',
              'Código', 'Recurso', 'Unidad', 'Cuadrilla', 'Cantidad', 'Precio'])
    s.append(['(opcional: un sub-análisis se usa en el APU como un recurso más, con su código)'])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
